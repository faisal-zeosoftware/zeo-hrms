"""Asset Management screens that are not a plain register: scan, labels, counting, floor plans, journal export."""
from io import BytesIO

from django.db import transaction
from django.http import FileResponse, Http404, HttpResponse
from rest_framework.response import Response

from zeo.appkit.engine import AppError, In, detail, d, guard, jsonable, need, now, today, user_name
from zeo.appkit.views import Base

from . import services as SV
from .models import CONDITION, AssetCard, Audit, AuditLine, DepreciationRun, Location


def can(app, request, *roles):
    return app.is_admin(request) or app.any(request, *roles)


def card_summary(c):
    return {'id': c.pk, 'code': c.code, 'name': c.name, 'status': c.status, 'status_label': c.get_status_display(), 'holder': c.holder,
            'location': c.location.path() if c.location_id else '', 'category': c.category.name, 'serial_number': c.serial_number,
            'condition': c.get_condition_display()}


class ScanView(Base):
    """GET ?code=… – the asset of a scanned QR / barcode / RFID tag or a typed code, with what the user may do now."""

    @guard
    def get(self, request):
        r = self.app.regs['asset']
        code = request.query_params.get('code', '')
        c = SV.find(code)
        if c is None:
            raise AppError(f'No asset found for "{code[:60]}".', 'code', status=404)
        if not r.scope(request).filter(pk=c.pk).exists():
            raise AppError('This asset is not in your branches or you have no access to it.', status=403)
        if can(self.app, request, 'ast_manager', 'ast_coordinator', 'ast_auditor'):
            SV.log_move(c, 'scan', request.user, to=c.location.path() if c.location_id else '', note='scanned')
        return Response({'record': detail(r, c, request), 'summary': card_summary(c)})


# ------------------------------------------------------------------ labels (PDF with QR code and Code 128 barcode)
SIZES = {'small': (50, 25), 'medium': (70, 35), 'large': (100, 50)}


def label_pdf(cards, size, heading):
    from reportlab.graphics import renderPDF
    from reportlab.graphics.barcode import code128, qr
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    w, h = (x * mm for x in SIZES.get(size, SIZES['medium']))
    pw, ph = A4
    margin, gap = 8 * mm, 3 * mm
    cols = max(1, int((pw - 2 * margin + gap) // (w + gap)))
    rows = max(1, int((ph - 2 * margin + gap) // (h + gap)))
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    pdf.setTitle('Asset labels')
    for i, c in enumerate(cards):
        k = i % (cols * rows)
        if i and k == 0:
            pdf.showPage()
        x = margin + (k % cols) * (w + gap)
        y = ph - margin - (k // cols + 1) * h - (k // cols) * gap
        pdf.setLineWidth(0.3)
        pdf.setStrokeGray(0.75)
        pdf.roundRect(x, y, w, h, 2 * mm)
        pad = 2 * mm
        qs = h - 2 * pad
        q = qr.QrCodeWidget(f'ASSET:{c.code}|SN:{c.serial_number}|ID:{c.hr_asset_id}')
        b = q.getBounds()
        dr = Drawing(qs, qs, transform=[qs / (b[2] - b[0]), 0, 0, qs / (b[3] - b[1]), 0, 0])
        dr.add(q)
        renderPDF.draw(dr, pdf, x + pad, y + pad)
        tx = x + qs + 2 * pad
        tw = w - qs - 3 * pad
        fs = max(5, h / mm / 5)
        pdf.setFillGray(0.35)
        pdf.setFont('Helvetica', fs * 0.75)
        pdf.drawString(tx, y + h - pad - fs * 0.75, (heading or '')[:40])
        pdf.setFillGray(0)
        pdf.setFont('Helvetica-Bold', fs * 1.25)
        pdf.drawString(tx, y + h - pad - fs * 2.2, c.code[:24])
        pdf.setFont('Helvetica', fs * 0.85)
        name = c.name
        while name and pdf.stringWidth(name, 'Helvetica', fs * 0.85) > tw:
            name = name[:-1]
        pdf.drawString(tx, y + h - pad - fs * 3.3, name)
        bar = code128.Code128(c.code, barHeight=h * 0.28, barWidth=0.6)
        scale = min(1.0, tw / max(bar.width, 1))
        pdf.saveState()
        pdf.translate(tx, y + pad)
        pdf.scale(scale, 1)
        bar.drawOn(pdf, 0, 0)
        pdf.restoreState()
    pdf.save()
    return buf.getvalue()


def company_name():
    try:
        from django.db import connection
        from UserManagement.models import company
        t = company.objects.filter(schema_name=connection.schema_name).first()
        return t.name if t else ''
    except Exception:
        return ''


class LabelsView(Base):
    """GET ?ids=1,2,3 | ?location=ID | ?category=ID [&size=small|medium|large] – a PDF sheet of labels (A4)."""

    @guard
    def get(self, request):
        q = request.query_params
        qs = self.app.regs['asset'].scope(request).exclude(status='disposed').select_related('category')
        ids = [int(x) for x in (q.get('ids') or '').split(',') if x.strip().isdigit()]
        if ids:
            qs = qs.filter(pk__in=ids)
        elif q.get('location', '').isdigit():
            loc = Location.objects.filter(pk=int(q['location'])).first()
            qs = qs.filter(location_id__in=loc.descendants() if loc else [])
        elif q.get('category', '').isdigit():
            qs = qs.filter(category_id=int(q['category']))
        else:
            raise AppError('Choose the assets, a location or a category to print.')
        cards = list(qs.order_by('code')[:2000])
        if not cards:
            raise AppError('No assets to print.')
        s = self.app.settings()
        size = q.get('size') if q.get('size') in SIZES else s.label_size
        data = label_pdf(cards, size, s.label_title or company_name())
        r = HttpResponse(data, content_type='application/pdf')
        r['Content-Disposition'] = f'inline; filename="asset-labels-{today():%Y%m%d}.pdf"'
        return r


# ------------------------------------------------------------------ counting screen
def line_json(ln, show_expected=True):
    c = ln.card
    return {'id': ln.pk, 'card': c.pk, 'code': c.code, 'name': c.name, 'expected_location': (ln.expected_location.path() if ln.expected_location_id else '') if show_expected else '',
            'found_location': ln.found_location.path() if ln.found_location_id else '', 'result': ln.result, 'result_label': ln.get_result_display(),
            'condition': ln.condition, 'holder': c.holder, 'counted_at': d(ln.counted_at), 'self_confirmed': ln.self_confirmed}


class AuditScanView(Base):
    """GET – the count list; POST {code, location, condition, result, note} – count one asset (scan or manual)."""

    def audit(self, request, pk):
        o = Audit.objects.filter(pk=pk).first()
        if o is None:
            raise Http404
        me = self.app.my_emp(request)
        allowed = can(self.app, request, 'ast_manager', 'ast_coordinator', 'ast_auditor', 'ast_viewer', 'ast_finance') or \
            (me is not None and o.auditors.filter(pk=me.pk).exists())
        need(allowed, 'You are not an auditor of this count.')
        return o

    def counter(self, o, request):
        me = self.app.my_emp(request)
        return can(self.app, request, 'ast_manager', 'ast_coordinator', 'ast_auditor') or (me is not None and o.auditors.filter(pk=me.pk).exists())

    def body(self, o, request):
        coord = can(self.app, request, 'ast_manager', 'ast_coordinator')
        show = coord or not o.blind
        lines = o.lines.select_related('card', 'expected_location', 'found_location')
        if not show:
            lines = lines.exclude(result='pending')
        return jsonable({'audit': {'id': o.pk, 'number': o.number, 'name': o.name, 'status': o.status, 'status_label': o.get_status_display(),
                                   'location': o.location.path() if o.location_id else '', 'blind': o.blind, 'due_date': o.due_date,
                                   'expected': o.expected, 'found': o.found, 'missing': o.missing, 'misplaced': o.misplaced, 'unexpected': o.unexpected,
                                   'damaged': o.damaged, 'can_count': o.status == 'counting' and self.counter(o, request)},
                         'locations': [{'id': x.pk, 'label': x.path()} for x in Location.objects.filter(pk__in=o.location.descendants() if o.location_id else
                                                                                                       Location.objects.filter(is_active=True).values('pk'))[:500]],
                         'conditions': [{'value': k, 'label': v} for k, v in CONDITION],
                         'lines': [line_json(ln, show) for ln in lines[:3000]]})

    @guard
    def get(self, request, pk):
        o = self.audit(request, pk)
        return Response(self.body(o, request))

    @guard
    def post(self, request, pk):
        o = self.audit(request, pk)
        need(o.status == 'counting', 'This count is not open for counting.')
        need(self.counter(o, request), 'You are not an auditor of this count.')
        f = In(request)
        if f.i('line'):
            ln = o.lines.filter(pk=f.i('line')).select_related('card').first()
            if ln is None:
                raise AppError('Line not found.')
            c = ln.card
        else:
            c = SV.find(f.s('code'))
            if c is None:
                raise AppError(f'No asset found for "{f.s("code")[:60]}". Check the code or register the asset.', 'code', status=404)
            ln = o.lines.filter(card=c).first()
        loc = f.obj('location', Location, False, 'Location')
        cond = f.choice('condition', CONDITION, '') or ''
        res = f.s('result')
        with transaction.atomic():
            if ln is None:
                ln = AuditLine(audit=o, card=c, holder_id=c.holder_employee_id, result='unexpected')
                msg = f'{c.code} is not on this count list – recorded as found here.'
            else:
                if res in ('found', 'missing', 'damaged', 'misplaced', 'pending'):
                    ln.result = res
                elif cond in ('minor_damage', 'major_damage'):
                    ln.result = 'damaged'
                elif loc is not None and ln.expected_location_id and loc.pk not in Location.objects.get(pk=ln.expected_location_id).descendants():
                    ln.result = 'misplaced'
                else:
                    ln.result = 'found' if ln.result in ('pending', 'missing', 'found') else ln.result
                msg = f'{c.code} {c.name}: {ln.get_result_display().lower()}.' + (' (Counted before.)' if ln.counted_at and not res else '')
            ln.found_location = loc or ln.found_location
            ln.condition = cond or ln.condition
            ln.note = (f.s('note') or ln.note)[:255]
            ln.counted_at, ln.counted_by = now(), request.user
            ln.save()
            SV.audit_counts(o)
        out = self.body(o, request)
        out['message'] = msg
        out['line'] = line_json(AuditLine.objects.select_related('card', 'expected_location', 'found_location').get(pk=ln.pk))
        return Response(out)


# ------------------------------------------------------------------ floor plans
class FloorPlanView(Base):
    """GET – the plan of a location with its assets; POST multipart {image} – upload the plan; POST {card, x, y} – place an asset (x, y in %)."""

    def loc(self, request, pk):
        o = self.app.regs['location'].scope(request).filter(pk=pk).first()
        if o is None:
            raise Http404
        return o

    @guard
    def get(self, request, pk):
        o = self.loc(request, pk)
        cards = self.app.regs['asset'].scope(request).filter(location_id__in=o.descendants()).exclude(status='disposed').select_related('location', 'category')
        return Response(jsonable({
            'location': {'id': o.pk, 'code': o.code, 'name': o.name, 'path': o.path(), 'kind': o.get_kind_display()},
            'image': f'/zeo-assets/api/floorplan/{o.pk}/image/' if o.floor_plan else '',
            'can_edit': can(self.app, request, 'ast_manager', 'ast_coordinator'),
            'children': [{'id': x.pk, 'label': x.path(), 'has_plan': bool(x.floor_plan)} for x in Location.objects.filter(parent=o, is_active=True)],
            'parent': o.parent_id,
            'cards': [dict(card_summary(c), x=float(c.plan_x) if c.plan_x is not None else None, y=float(c.plan_y) if c.plan_y is not None else None,
                           here=c.location_id == o.pk) for c in cards[:1000]]}))

    @guard
    def post(self, request, pk):
        o = self.loc(request, pk)
        need(can(self.app, request, 'ast_manager', 'ast_coordinator'), 'Only the asset manager or coordinator can change the floor plan.')
        img = request.FILES.get('image')
        if img is not None:
            if not (img.content_type or '').startswith('image/') or img.size > 10 * 1024 * 1024:
                raise AppError('Upload a PNG or JPG image of up to 10 MB.', 'image')
            o.floor_plan = img
            o.save(update_fields=['floor_plan'])
            self.app.event('location', o.pk, f'Floor plan uploaded by {user_name(request.user)}', request.user)
            return self.get(request, pk)
        f = In(request)
        c = self.app.regs['asset'].scope(request).filter(pk=f.i('card', True, 'Asset')).first()
        if c is None:
            raise AppError('Asset not found.', 'card')
        x, y = f.n('x'), f.n('y')
        if f.b('remove'):
            AssetCard.objects.filter(pk=c.pk).update(plan_x=None, plan_y=None)
        else:
            if x is None or y is None or not (0 <= x <= 100 and 0 <= y <= 100):
                raise AppError('The position is outside the plan.')
            if c.location_id != o.pk:
                SV.move(c, request.user, o, 'placed on the floor plan')
            AssetCard.objects.filter(pk=c.pk).update(plan_x=round(x, 2), plan_y=round(y, 2))
        return self.get(request, pk)


class FloorPlanImageView(Base):
    @guard
    def get(self, request, pk):
        o = self.app.regs['location'].scope(request).filter(pk=pk).first()
        if o is None or not o.floor_plan:
            raise Http404
        return FileResponse(o.floor_plan.open('rb'))


# ------------------------------------------------------------------ journal export
class JournalView(Base):
    """GET – the journal of a depreciation run as an Excel file (summary by account + one line per asset)."""

    @guard
    def get(self, request, pk):
        need(can(self.app, request, 'ast_manager', 'ast_finance'), 'Only the fixed asset accountant or the asset manager can export the journal.')
        o = DepreciationRun.objects.filter(pk=pk).first()
        if o is None:
            raise Http404
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from .app import journal
        wb = Workbook()
        ws = wb.active
        ws.title = 'Journal'
        cur = self.app.settings().currency
        ws.append([f'Depreciation {o.number}', f'{o.period_from:%d %b %Y} – {o.period_to:%d %b %Y}', o.get_status_display()])
        ws.append([])
        ws.append(['Date', 'Account', 'Description', f'Debit ({cur})', f'Credit ({cur})'])
        for c in ws[3]:
            c.font = Font(bold=True)
        for j in journal(o):
            text = f'Depreciation {o.period_from:%b %Y}–{o.period_to:%b %Y} – {j["category"]}'
            ws.append([o.period_to, j['debit'], text, j['amount'], None])
            ws.append([o.period_to, j['credit'], text, None, j['amount']])
        ws.append(['', '', 'Total', float(o.total), float(o.total)])
        ws2 = wb.create_sheet('Assets')
        ws2.append(['Asset code', 'Asset', 'Category', 'Months', f'Opening ({cur})', f'Depreciation ({cur})', f'Closing ({cur})', 'Debit account', 'Credit account'])
        for c in ws2[1]:
            c.font = Font(bold=True)
        for ln in o.lines.select_related('card', 'category'):
            ws2.append([ln.card.code, ln.card.name, ln.category.name if ln.category else '', float(ln.months), float(ln.opening), float(ln.amount),
                        float(ln.closing), ln.gl_expense, ln.gl_accumulated])
        for sh in (ws, ws2):
            for col in sh.columns:
                sh.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or '')) for c in col) + 2))
        buf = BytesIO()
        wb.save(buf)
        r = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r['Content-Disposition'] = f'attachment; filename="{o.number}-journal.xlsx"'
        return r


class SyncView(Base):
    """POST – bring in HR assets that have no card yet and refresh every card from HR (manager)."""

    @guard
    def post(self, request):
        need(can(self.app, request, 'ast_manager'), 'Only the asset manager can do this.')
        n = SV.import_all()
        self.app.event('settings', 0, f'HR assets synchronised by {user_name(request.user)}: {n} new', request.user)
        return Response({'imported': n, 'message': f'{n} asset(s) brought in from HR; every asset refreshed.'})
