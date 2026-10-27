"""v1.61.0 – the member file (census) of a medical insurance plan for the insurer: one row per insured person – the employee
(principal) and the insured dependants – with the data insurers ask for in the UAE (name, gender, date of birth, nationality,
Emirates ID, passport, visa / UID, marital status, salary band, member number, effective date). Also the changes since a
date (additions and deletions) for the monthly endorsement.

Sources: People › Enrolments (active on the date) of the plan; the dependants are the employee's family members marked
"insured" (Employee › Family, with gender, Emirates ID, passport, visa) – or, when none is marked, the free-text dependants
of the enrolment ("name – relation – date of birth").

Insurer formats differ: the columns here are the common ones; check the insurer's template before the first upload and
rename / reorder columns there if needed.
"""
import csv
import io
import re
from datetime import date

from django.apps import apps
from django.db.models import Q

M = apps.get_model
COLUMNS = (
    ('change', 'Change'), ('member_type', 'Member type'), ('relation', 'Relation'), ('employee_code', 'Employee no.'),
    ('member_name', 'Member name'), ('arabic_name', 'Name (Arabic)'), ('gender', 'Gender'), ('date_of_birth', 'Date of birth'),
    ('age', 'Age'), ('marital_status', 'Marital status'), ('nationality', 'Nationality'), ('emirates_id', 'Emirates ID'),
    ('passport_no', 'Passport no.'), ('visa_uid', 'Visa / UID no.'), ('visa_file_no', 'Visa file no.'), ('mobile', 'Mobile'),
    ('email', 'E-mail'), ('work_location', 'Work location'), ('category', 'Category / grade'), ('salary_band', 'Salary band (monthly)'),
    ('member_no', 'Member / card no.'), ('policy_no', 'Policy no.'), ('effective_date', 'Effective date'), ('end_date', 'End date'),
)
SPOUSE = ('spouse', 'wife', 'husband', 'زوجة', 'زوج')
CHILD = ('son', 'daughter', 'child', 'kid', 'ابن', 'ابنة', 'بنت')
PARENT = ('father', 'mother', 'parent', 'أب', 'أم')


def _age(dob, on):
    if not dob:
        return ''
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def member_type(relation):
    r = str(relation or '').strip().lower()
    if any(w in r for w in SPOUSE):
        return 'Spouse'
    if any(w in r for w in CHILD):
        return 'Child'
    if any(w in r for w in PARENT):
        return 'Parent'
    return 'Dependant'


def salary_band(wage):
    if wage is None:
        return ''
    w = float(wage)
    if w <= 4000:
        return 'Up to 4,000'
    if w <= 12000:
        return '4,001 – 12,000'
    return 'Above 12,000'


def _gender(g):
    g = str(g or '').strip().upper()[:1]
    return {'M': 'Male', 'F': 'Female'}.get(g, '')


def _enrolments(plan, on=None, since=None):
    E = M('ZeoPeople', 'Enrolment')
    qs = E.objects.filter(plan=plan).exclude(status__in=('requested', 'rejected')).select_related(
        'employee', 'employee__emp_branch_id', 'employee__emp_nationality', 'employee__emp_ctgry_id')
    if on is not None:
        qs = qs.filter(start__lte=on).filter(Q(end__isnull=True) | Q(end__gte=on))
    return qs.order_by('employee__emp_code')


def _dependants(enr, emp):
    """[(name, relation, dob, extra or None, family row or None)]"""
    Fam = M('EmpManagement', 'emp_family')
    DX = M('EmployeeProfile', 'DependentExtra')
    fams = {f.pk: f for f in Fam.objects.filter(emp_id=emp, is_active=True)}
    extras = {x.family_id: x for x in DX.objects.filter(family_id__in=list(fams), insured=True)}
    out = [(f.ef_member_name, f.emp_relation, f.ef_date_of_birth, extras[f.pk], f) for pk, f in fams.items() if pk in extras]
    if out:
        return out
    for line in enr.dependants or []:
        parts = [p.strip() for p in re.split(r'\s[–-]\s|,', str(line)) if p.strip()]
        if not parts:
            continue
        dob = None
        if len(parts) > 2:
            try:
                dob = date.fromisoformat(parts[2][:10])
            except ValueError:
                dob = None
        out.append((parts[0], parts[1] if len(parts) > 1 else '', dob, None, None))
    return out


def rows(plan, on, since=None, only=None):
    """The member rows of the plan on `on`; with `since`: only the additions and deletions from `since` to `on`.
    only: enrolment ids the user may see (None = all)."""
    from DashboardManagement.reports import wage_by_employee
    ID = M('EmployeeProfile', 'EmployeeIdentity')
    N = M('Core', 'Nationality')
    if since is None:
        enrs = list(_enrolments(plan, on))
        changes = {e.pk: '' for e in enrs}
    else:
        allq = list(_enrolments(plan))
        enrs, changes = [], {}
        for e in allq:
            if since <= e.start <= on:
                enrs.append(e); changes[e.pk] = 'Addition'
            elif e.end and since <= e.end <= on:
                enrs.append(e); changes[e.pk] = 'Deletion'
    if only is not None:
        enrs = [e for e in enrs if e.pk in only]
    emp_ids = [e.employee_id for e in enrs]
    ids = {x.employee_id: x for x in ID.objects.filter(employee_id__in=emp_ids)}
    wages = wage_by_employee(emp_ids)
    nat_names = dict(N.objects.values_list('id', 'N_name'))
    out = []
    for e in enrs:
        emp, idn = e.employee, ids.get(e.employee_id)
        base = {'employee_code': emp.emp_code, 'policy_no': plan.policy_no, 'effective_date': e.start, 'end_date': e.end or '',
                'work_location': emp.emp_branch_id.branch_name if emp.emp_branch_id_id else '',
                'category': emp.emp_ctgry_id.ctgry_title if getattr(emp, 'emp_ctgry_id_id', None) else '', 'change': changes.get(e.pk, '')}
        out.append({**base, 'member_type': 'Principal', 'relation': 'Self',
                    'member_name': ' '.join(x for x in (emp.emp_first_name, emp.emp_middle_name, emp.emp_last_name) if x),
                    'arabic_name': getattr(idn, 'arabic_name', '') if idn else '', 'gender': _gender(emp.emp_gender),
                    'date_of_birth': emp.emp_date_of_birth or '', 'age': _age(emp.emp_date_of_birth, on),
                    'marital_status': (emp.emp_marital_status or '').title(), 'nationality': emp.emp_nationality.N_name if emp.emp_nationality_id else '',
                    'emirates_id': getattr(idn, 'emirates_id', '') if idn else '', 'passport_no': getattr(idn, 'passport_no', '') if idn else '',
                    'visa_uid': getattr(idn, 'visa_uid', '') if idn else '', 'visa_file_no': getattr(idn, 'visa_file_no', '') if idn else '',
                    'mobile': emp.emp_mobile_number_1 or '', 'email': emp.emp_company_email or emp.emp_personal_email or '',
                    'salary_band': salary_band(wages.get(emp.pk)), 'member_no': e.member_no})
        if not plan.covers_dependants:
            continue
        deps = _dependants(e, emp)
        if plan.max_dependants:
            deps = deps[:plan.max_dependants]
        for name, rel, dob, x, f in deps:
            out.append({**base, 'member_type': member_type(rel), 'relation': (rel or '').title(), 'member_name': name, 'arabic_name': '',
                        'gender': _gender(getattr(x, 'gender', '')) if x else '', 'date_of_birth': dob or '', 'age': _age(dob, on),
                        'marital_status': 'Married' if member_type(rel) == 'Spouse' else '',
                        'nationality': nat_names.get(getattr(x, 'nationality_id', None), '') if x else '',
                        'emirates_id': getattr(x, 'emirates_id', '') if x else '', 'passport_no': getattr(x, 'passport_no', '') if x else '',
                        'visa_uid': getattr(x, 'visa_number', '') if x else '', 'visa_file_no': '', 'mobile': '', 'email': '',
                        'salary_band': '', 'member_no': ''})
    return out


def problems(rs):
    """Rows the insurer will refuse: what is missing per member."""
    out = []
    for r in rs:
        miss = [lbl for k, lbl in (('gender', 'gender'), ('date_of_birth', 'date of birth'), ('nationality', 'nationality')) if not r.get(k)]
        if not r.get('emirates_id') and not r.get('passport_no'):
            miss.append('Emirates ID or passport')
        if miss:
            out.append({'member': r['member_name'], 'employee_code': r['employee_code'], 'missing': miss})
    return out


def _cell(v):
    return v.isoformat() if isinstance(v, date) else ('' if v is None else v)


def as_csv(rs):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([lbl for _, lbl in COLUMNS])
    for r in rs:
        w.writerow([_cell(r.get(k, '')) for k, _ in COLUMNS])
    return ('﻿' + buf.getvalue()).encode('utf-8')


def as_xlsx(rs, plan, on, since=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Members' if since is None else 'Changes'
    ws.append([lbl for _, lbl in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='5B4FE0')
    for r in rs:
        ws.append([_cell(r.get(k, '')) for k, _ in COLUMNS])
    for i, (k, lbl) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(32, len(lbl) + 4))
    ws.freeze_panes = 'A2'
    info = wb.create_sheet('About')
    for row in (('Plan', plan.name), ('Insurer', plan.provider), ('Policy no.', plan.policy_no), ('Members on', on.isoformat()),
                ('Changes since', since.isoformat() if since else '–'), ('Members', len(rs)), ('Made by', 'ZEO HRMS')):
        info.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
