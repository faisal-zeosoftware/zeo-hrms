"""v1.52.0 – scheduled report delivery: a report of Reports e-mailed on a schedule, as Excel (xlsx) or CSV, to a list.

The report runs with the rights of its owner (the user who made the schedule – branches and rights apply), for the period
chosen (last month, this month so far, last week, yesterday, as of the day). Beat 'zeoadmin-schedules' runs every 15 minutes
and sends what is due; *Send now* sends at once.
"""
import calendar
import csv
import io
import logging
import re
from datetime import date, datetime, timedelta

from django.utils import timezone

log = logging.getLogger(__name__)
EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def period_dates(period, today=None):
    t = today or timezone.localdate()
    if period == 'last_month':
        end = t.replace(day=1) - timedelta(days=1)
        return end.replace(day=1), end
    if period == 'this_month':
        return t.replace(day=1), t
    if period == 'last_week':
        start = t - timedelta(days=t.weekday() + 7)
        return start, start + timedelta(days=6)
    if period == 'yesterday':
        y = t - timedelta(days=1)
        return y, y
    return None, t


def next_run(s, after=None):
    """The next time the schedule is due (company time)."""
    now = timezone.localtime(after or timezone.now())
    hour = min(max(int(s.hour or 0), 0), 23)
    base = now.replace(hour=hour, minute=0, second=0, microsecond=0)
    for i in range(0, 70):
        d = base + timedelta(days=i)
        if d <= now:
            continue
        if s.frequency == 'daily':
            return d
        if s.frequency == 'weekly' and d.weekday() == (s.weekday or 0):
            return d
        if s.frequency == 'monthly':
            want = min(max(s.day or 1, 1), calendar.monthrange(d.year, d.month)[1])
            if d.day == want:
                return d
    return base + timedelta(days=1)


def check(s):
    from DashboardManagement.reports import REPORTS
    if s.report not in REPORTS:
        return 'Choose a report.'
    bad = [r for r in s.recipients or [] if not EMAIL_RE.match(r)]
    if not s.recipients:
        return 'Give at least one e-mail address.'
    if bad:
        return f'Not an e-mail address: {", ".join(bad[:3])}.'
    if len(s.recipients) > 50:
        return 'At most 50 addresses.'
    return ''


def run_report(s):
    """The report data (as the Reports screen gets it) run as the owner."""
    from rest_framework.test import APIRequestFactory, force_authenticate
    from django.db import connection
    from DashboardManagement.reports import ReportView
    a, b = period_dates(s.period)
    q = {}
    if a:
        q['from'] = a.isoformat()
    if b:
        q['to'] = b.isoformat()
        q['as_of'] = b.isoformat()
    for k, v in (s.params or {}).items():
        if v not in (None, ''):
            q[k] = ','.join(str(x) for x in v) if isinstance(v, list) else str(v)
    rf = APIRequestFactory().get(f'/dashboard/api/reports/{s.report}/', q)
    rf.tenant = getattr(connection, 'tenant', None)
    force_authenticate(rf, user=s.owner)
    r = ReportView.as_view()(rf, key=s.report)
    if r.status_code != 200:
        d = getattr(r, 'data', {}) or {}
        raise ValueError(d.get('detail') if isinstance(d, dict) else str(d))
    return r.data


def to_file(data, fmt):
    cols = data['columns']
    if fmt == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([c['label'] for c in cols])
        for row in data['rows']:
            w.writerow(['' if row.get(c['key']) is None else row.get(c['key']) for c in cols])
        return ('﻿' + buf.getvalue()).encode(), 'text/csv', 'csv'
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = (data.get('title') or 'Report')[:31].replace('/', '-')
    ws.append([data.get('title') or ''])
    ws['A1'].font = Font(bold=True, size=13)
    per = data.get('period')
    ws.append([f"{per['from']} – {per['to']}" if per else f"As of {timezone.localdate():%d/%m/%Y}"])
    ws.append([])
    ws.append([c['label'] for c in cols])
    for i in range(1, len(cols) + 1):
        cell = ws.cell(row=4, column=i)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F2433')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in data['rows']:
        out = []
        for c in cols:
            v = row.get(c['key'])
            if isinstance(v, (list, dict)):
                v = str(v)
            if c.get('type') == 'date' and isinstance(v, str) and re.match(r'^\d{4}-\d{2}-\d{2}', v):
                try:
                    v = datetime.strptime(v[:10], '%Y-%m-%d').date()
                except ValueError:
                    pass
            out.append(v)
        ws.append(out)
    if data.get('totals'):
        ws.append([('Total' if i == 0 else data['totals'].get(c['key'], '')) for i, c in enumerate(cols)])
        for i in range(1, len(cols) + 1):
            ws.cell(row=ws.max_row, column=i).font = Font(bold=True)
    for i, c in enumerate(cols, start=1):
        if c.get('type') == 'date':
            for r in range(5, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = 'DD/MM/YYYY'
        elif c.get('type') == 'number':
            for r in range(5, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = '#,##0.00'
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(c['label']) + 2, 12), 40)
    ws.freeze_panes = 'A5'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'


def mail_connection():
    from django.apps import apps
    from django.conf import settings
    from django.core.mail import get_connection
    try:
        cfg = apps.get_model('EmpManagement', 'EmailConfiguration').objects.filter(is_active=True).first()
    except LookupError:
        cfg = None
    if cfg:
        return get_connection(host=cfg.email_host, port=cfg.email_port, username=cfg.email_host_user, password=cfg.email_host_password,
                              use_tls=cfg.email_use_tls), cfg.email_host_user
    return get_connection(), getattr(settings, 'DEFAULT_FROM_EMAIL', None) or getattr(settings, 'EMAIL_HOST_USER', None)


def send(s, now=None):
    """Run and e-mail one schedule. Returns the status text."""
    from django.core.mail import EmailMessage
    now = now or timezone.now()
    try:
        data = run_report(s)
        content, mime, ext = to_file(data, s.fmt)
        per = data.get('period')
        subject = f"{data['title']}" + (f" ({per['from']} – {per['to']})" if per else f' ({timezone.localdate():%d/%m/%Y})')
        body = (f"{data['title']}\n{data.get('description') or ''}\n\n{data['count']} rows – the report is attached ({ext.upper()}).\n"
                + ''.join(f"\n{c['label']}: {data['totals'][c['key']]:,}" for c in data['columns'] if c['key'] in (data.get('totals') or {}))
                + f"\n\nScheduled report \"{s.name}\" from ZEO HRMS ({s.get_frequency_display().lower()}).")
        conn, sender = mail_connection()
        msg = EmailMessage(subject=subject, body=body, from_email=sender, to=list(s.recipients), connection=conn)
        msg.attach(f"{s.report}-{timezone.localdate().isoformat()}.{ext}", content, mime)
        msg.send()
        status = f"sent to {len(s.recipients)} – {data['count']} rows"
    except Exception as exc:
        log.warning('scheduled report %s failed: %s', s.pk, exc)
        status = f'failed: {exc}'[:255]
    s.last_run, s.last_status, s.next_run = now, status[:255], next_run(s, now)
    s.save(update_fields=['last_run', 'last_status', 'next_run'])
    return status


def due(now=None):
    from .models import ReportSchedule
    now = now or timezone.now()
    out = []
    for s in ReportSchedule.objects.filter(active=True, next_run__lte=now).select_related('owner')[:100]:
        out.append((s.pk, send(s, now)))
    return out


def json_row(s):
    from DashboardManagement.reports import REPORTS
    spec = REPORTS.get(s.report)
    return {'id': s.pk, 'name': s.name, 'report': s.report, 'report_title': spec[0] if spec else s.report, 'params': s.params, 'period': s.period,
            'frequency': s.frequency, 'weekday': s.weekday, 'day': s.day, 'hour': s.hour, 'fmt': s.fmt, 'recipients': s.recipients,
            'owner': s.owner.get_full_name() or s.owner.username, 'owner_id': s.owner_id, 'active': s.active, 'last_run': s.last_run,
            'last_status': s.last_status, 'next_run': s.next_run}
