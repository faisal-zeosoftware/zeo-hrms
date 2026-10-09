"""AttendancePlus services (v1.12.0): punch checks (method, IP, GPS, geofence, selfie), QR / PIN,
kiosk and device intake, bulk import and attendance corrections."""
import csv
import hashlib
import hmac
import io
import ipaddress
import logging
import time as _time
from datetime import datetime, timedelta

from django.conf import settings
from django.contrib.auth.hashers import check_password, make_password
from django.db import connection, transaction
from django.utils import timezone

from . import engine as E
from . import models as M

log = logging.getLogger(__name__)


class Refused(Exception):
    """A punch / request that must be refused – message is shown to the user."""

    def __init__(self, message, code=400):
        super().__init__(message)
        self.code = code


# ----------------------------------------------------------------------------- network / location
def client_ip(request):
    fwd = request.META.get('HTTP_X_FORWARDED_FOR', '')
    if fwd:
        return fwd.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', '') or ''


def parse_cidrs(text):
    nets, bad = [], []
    for line in str(text or '').replace(',', '\n').splitlines():
        s = line.strip()
        if not s:
            continue
        try:
            nets.append(ipaddress.ip_network(s, strict=False))
        except ValueError:
            bad.append(s)
    return nets, bad


def ip_allowed(employee, ip):
    rows = M.IPRestriction.objects.filter(is_active=True)
    rows = [r for r in rows if r.branch_id in (None, employee.emp_branch_id_id)]
    if not rows:
        return False, 'No allowed office networks are set up for your branch. Ask HR to add them under Attendance Plus → IP restrictions.'
    try:
        addr = ipaddress.ip_address(ip)
    except ValueError:
        return False, 'Your network address could not be read.'
    for r in rows:
        nets, _ = parse_cidrs(r.cidrs)
        if any(addr in n for n in nets):
            return True, ''
    return False, f'You can only punch from the office network (your address {ip} is not allowed).'


def validation_policy(employee):
    try:
        from calendars.utils import get_employee_attendance_validation_policy
        return get_employee_attendance_validation_policy(employee)
    except Exception:
        return None


def check_punch(employee, rule, method, *, ip='', lat=None, lng=None, photo=None, skip_ip=False):
    """Raise Refused when the rule (or the existing validation policy) does not allow this punch.
    Returns geofence_ok (True / False / None when not checked)."""
    if not E.method_allowed(rule, method):
        names = dict(M.METHODS)
        allowed = ', '.join(names.get(m, m) for m in rule.allowed_methods)
        raise Refused(f'Punching by {names.get(method, method)} is not allowed for you. Allowed: {allowed}.', 403)
    pol = validation_policy(employee)
    need_geo = rule.require_geofence or bool(pol and pol.enable_geofencing)
    need_gps = rule.require_gps or need_geo
    need_photo = rule.require_selfie or bool(pol and pol.enable_photo_capture)
    if method in ('kiosk', 'device', 'import', 'correction', 'manual'):
        need_geo = need_gps = need_photo = False
        skip_ip = True
    if rule.require_ip and not skip_ip:
        ok, msg = ip_allowed(employee, ip)
        if not ok:
            raise Refused(msg, 403)
    if need_gps and (lat in (None, '') or lng in (None, '')):
        raise Refused('Turn on location (GPS) and try again – your location is required to punch.')
    geo = None
    if lat not in (None, '') and lng not in (None, ''):
        try:
            from calendars.utils import validate_employee_geofence
            geo = bool(validate_employee_geofence(employee, lat, lng))
        except Exception:
            geo = None
    if need_geo and geo is False:
        raise Refused('You are outside the allowed work location (geofence). Move inside the site and try again.', 403)
    if need_photo and not photo:
        raise Refused('Take a selfie photo to punch – a photo is required.')
    return geo


# ----------------------------------------------------------------------------- QR / PIN
def _sig(text, n=16):
    return hmac.new(settings.SECRET_KEY.encode(), text.encode(), hashlib.sha256).hexdigest()[:n]


def employee_key(emp_id):
    k, _ = M.EmployeeKey.objects.get_or_create(employee_id=emp_id)
    return k


def employee_qr(employee):
    k = employee_key(employee.id)
    body = f'{connection.schema_name}:{employee.id}:{k.qr_version}'
    return f'ZQR1.{employee.id}.{k.qr_version}.{_sig(body)}'


def employee_from_qr(code):
    """Employee badge QR (ZQR1...) or the existing barcode number."""
    from EmpManagement.models import emp_master
    code = (code or '').strip()
    if code.startswith('ZQR1.'):
        try:
            _, emp_id, ver, sig = code.split('.')
            emp_id, ver = int(emp_id), int(ver)
        except ValueError:
            raise Refused('This QR code is not valid.')
        body = f'{connection.schema_name}:{emp_id}:{ver}'
        if not hmac.compare_digest(sig, _sig(body)):
            raise Refused('This QR code is not valid for this company.')
        k = M.EmployeeKey.objects.filter(employee_id=emp_id).first()
        if not k or k.qr_version != ver:
            raise Refused('This QR badge was cancelled. Ask HR for a new one.')
        e = emp_master.objects.filter(pk=emp_id, is_active=True).first()
        if not e:
            raise Refused('Employee not found or inactive.')
        return e
    e = emp_master.objects.filter(barcode_number=code, is_active=True).first() if code else None
    if not e:
        raise Refused('Badge / barcode not recognised.')
    return e


def site_qr(kiosk, at=None):
    step = max(10, kiosk.site_qr_seconds or 30)
    win = int((at or _time.time()) // step)
    body = f'{connection.schema_name}:site:{kiosk.id}:{win}'
    return f'ZSITE1.{kiosk.id}.{win}.{_sig(body)}', step - int((at or _time.time()) % step)


def kiosk_from_site_qr(code):
    try:
        _, kid, win, sig = (code or '').strip().split('.')
        kid, win = int(kid), int(win)
    except ValueError:
        raise Refused('This site QR code is not valid. Scan the code shown on the kiosk screen.')
    k = M.Kiosk.objects.filter(pk=kid, is_active=True, show_site_qr=True).first()
    if not k:
        raise Refused('This site QR code is not active.')
    step = max(10, k.site_qr_seconds or 30)
    cur = int(_time.time() // step)
    if win not in (cur, cur - 1):
        raise Refused('The site QR code has expired. Scan the code on the screen again.')
    body = f'{connection.schema_name}:site:{k.id}:{win}'
    if not hmac.compare_digest(sig, _sig(body)):
        raise Refused('This site QR code is not valid.')
    return k


def set_pin(emp_id, pin):
    pin = str(pin or '').strip()
    if not pin.isdigit() or not 4 <= len(pin) <= 8:
        raise Refused('The PIN must be 4 to 8 digits.')
    k = employee_key(emp_id)
    k.pin_hash = make_password(pin)
    k.save(update_fields=['pin_hash', 'updated_at'])


def employee_from_pin(code, pin):
    from EmpManagement.models import emp_master
    e = emp_master.objects.filter(emp_code__iexact=str(code or '').strip(), is_active=True).first()
    k = M.EmployeeKey.objects.filter(employee_id=e.id).first() if e else None
    if not e or not k or not k.pin_hash or not check_password(str(pin or ''), k.pin_hash):
        raise Refused('Employee code or PIN is wrong.', 403)
    return e


# ----------------------------------------------------------------------------- devices
ZK_STATUS = {'0': 'in', '1': 'out', '2': 'break_out', '3': 'break_in', '4': 'in', '5': 'out'}


def employee_for_pin(pin):
    from calendars.models import EmployeeMachineMapping
    from EmpManagement.models import emp_master
    m = EmployeeMachineMapping.objects.filter(machine_code=str(pin).strip()).select_related('employee').first()
    if m:
        return m.employee
    return emp_master.objects.filter(emp_code__iexact=str(pin).strip()).first()


def device_seen(device, ip=''):
    M.Device.objects.filter(pk=device.pk).update(last_seen=timezone.now(), last_ip=ip or device.last_ip)


def device_punches(device, rows, ip=''):
    """rows: [(pin, local datetime text / datetime, status_key or kind)] → result counts. Idempotent."""
    out = {'received': 0, 'saved': 0, 'duplicates': 0, 'unknown_employee': 0, 'errors': []}
    touched = set()
    for pin, when, status in rows:
        out['received'] += 1
        e = employee_for_pin(pin)
        if not e:
            out['unknown_employee'] += 1
            continue
        try:
            ts = when if isinstance(when, datetime) else datetime.strptime(str(when).strip()[:19], '%Y-%m-%d %H:%M:%S')
        except ValueError:
            out['errors'].append(f'bad time "{when}" for {pin}')
            continue
        ts = E.as_aware(ts)
        kind = 'auto'
        s = str(status).strip() if status is not None else ''
        if s in ('in', 'out', 'break_out', 'break_in', 'lunch_out', 'lunch_in'):
            kind = s
        elif device.use_status_keys and s in ZK_STATUS:
            kind = ZK_STATUS[s]
        try:
            p, created = E.add_punch(e, ts, kind, 'device', device=device, device_ref=device.serial_number, ip=ip,
                                     verified_by='device', check=False, recompute=False)
        except Exception as exc:
            out['errors'].append(f'{pin}: {exc}')
            continue
        if created:
            out['saved'] += 1
            touched.add((e.id, p.work_date))
        else:
            out['duplicates'] += 1
    for emp_id, day in touched:
        E.recompute_day(emp_id, day)
    if out['saved']:
        M.Device.objects.filter(pk=device.pk).update(punches_received=device.punches_received + out['saved'])
    device_seen(device, ip)
    return out


def parse_attlog(text):
    """ZKTeco ADMS ATTLOG body: PIN \t YYYY-MM-DD HH:MM:SS \t status \t verify \t workcode ..."""
    rows = []
    for line in (text or '').splitlines():
        line = line.strip()
        if not line:
            continue
        parts = line.split('\t') if '\t' in line else line.split()
        if len(parts) >= 3 and len(parts[1]) == 10 and ':' in parts[2]:   # date and time split by a space
            parts = [parts[0], f'{parts[1]} {parts[2]}'] + parts[3:]
        if len(parts) < 2:
            continue
        rows.append((parts[0], parts[1], parts[2] if len(parts) > 2 else ''))
    return rows


def find_schema_for_serial(sn):
    """ADMS devices cannot send ?schema= – find the company that registered this serial number."""
    from django_tenants.utils import get_tenant_model, schema_context
    from django.core.cache import cache
    key = f'attplus-sn-{sn}'
    hit = cache.get(key)
    if hit:
        return hit
    for schema in get_tenant_model().objects.exclude(schema_name='public').values_list('schema_name', flat=True):
        try:
            with schema_context(schema):
                if M.Device.objects.filter(serial_number=sn).exists():
                    cache.set(key, schema, 3600)
                    return schema
        except Exception:
            continue
    return None


# ----------------------------------------------------------------------------- bulk import
IMPORT_HEADERS = ['employee_code', 'date', 'time', 'type']


def read_rows(upload):
    name = (getattr(upload, 'name', '') or '').lower()
    data = upload.read()
    if name.endswith(('.xlsx', '.xls')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        head = [str(h or '').strip().lower().replace(' ', '_') for h in next(it, [])]
        return [dict(zip(head, r)) for r in it if any(v not in (None, '') for v in r)]
    text = data.decode('utf-8-sig', errors='replace')
    rdr = csv.DictReader(io.StringIO(text))
    return [{(k or '').strip().lower().replace(' ', '_'): v for k, v in r.items()} for r in rdr]


def import_punches(rows, commit=False, user_id=None, allowed_ids=None):
    from EmpManagement.models import emp_master
    out, seen = [], set()
    kinds = {'in': 'in', 'clock in': 'in', 'out': 'out', 'clock out': 'out', 'break out': 'break_out', 'break in': 'break_in',
             'break_out': 'break_out', 'break_in': 'break_in', 'lunch out': 'lunch_out', 'lunch in': 'lunch_in',
             'lunch_out': 'lunch_out', 'lunch_in': 'lunch_in', '': 'auto', 'auto': 'auto'}
    touched = set()
    for i, r in enumerate(rows, start=2):
        code = str(r.get('employee_code') or r.get('emp_code') or '').strip()
        res = {'row': i, 'employee_code': code, 'date': str(r.get('date') or '')[:10], 'time': str(r.get('time') or ''),
               'type': str(r.get('type') or '').strip().lower(), 'errors': [], 'status': 'ok'}
        e = emp_master.objects.filter(emp_code__iexact=code).first() if code else None
        if not e:
            res['errors'].append('Employee code not found.')
        elif allowed_ids is not None and e.id not in allowed_ids:
            res['errors'].append('This employee is not in your branches.')
        ts = None
        try:
            d = r.get('date')
            d = d.date() if isinstance(d, datetime) else d
            d = d if hasattr(d, 'year') else datetime.strptime(str(d).strip()[:10], '%Y-%m-%d').date()
            t = r.get('time')
            if hasattr(t, 'hour'):
                t = t.strftime('%H:%M:%S')
            t = str(t).strip()
            t = datetime.strptime(t, '%H:%M:%S' if t.count(':') == 2 else '%H:%M').time()
            ts = E.aware(d, t)
            if ts > E.now():
                res['errors'].append('The time is in the future.')
        except (ValueError, TypeError):
            res['errors'].append('Date must be YYYY-MM-DD and time HH:MM.')
        kind = kinds.get(res['type'])
        if kind is None:
            res['errors'].append('Type must be in, out, break out, break in, lunch out, lunch in or empty.')
        if e and ts:
            key = (e.id, ts)
            if key in seen:
                res['errors'].append('Same employee and time twice in the file.')
            seen.add(key)
            if not res['errors'] and M.Punch.objects.filter(employee_id=e.id, ts=ts).exists():
                res['status'] = 'duplicate'
        if res['errors']:
            res['status'] = 'error'
        elif commit and res['status'] == 'ok':
            p, created = E.add_punch(e, ts, kind, 'import', user_id=user_id, check=False, recompute=False)
            touched.add((e.id, p.work_date))
            res['status'] = 'imported' if created else 'duplicate'
        out.append(res)
    for emp_id, day in touched:
        E.recompute_day(emp_id, day)
    return out


# ----------------------------------------------------------------------------- corrections
def _parse_dt(v, day=None):
    if v in (None, ''):
        return None
    if isinstance(v, datetime):
        return E.as_aware(v)
    s = str(v).strip().replace('T', ' ')
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M'):
        try:
            return E.as_aware(datetime.strptime(s[:19] if fmt.endswith('%S') else s[:16], fmt))
        except ValueError:
            pass
    if day is not None:
        for fmt in ('%H:%M:%S', '%H:%M'):
            try:
                return E.aware(day, datetime.strptime(s, fmt).time())
            except ValueError:
                pass
    raise Refused(f'"{v}" is not a valid date and time (use YYYY-MM-DD HH:MM).')


def new_correction(employee, data, user, by_hr=False, attachment=None):
    rule = E.rule_for(employee)
    try:
        day = E.as_date(data.get('date'))
    except Exception:
        raise Refused('Pick the date to correct (YYYY-MM-DD).')
    if day > E.today():
        raise Refused('You cannot correct a future date.')
    if not by_hr and rule.correction_window_days and (E.today() - day).days > rule.correction_window_days:
        raise Refused(f'Corrections must be requested within {rule.correction_window_days} days. Ask HR to correct older dates.')
    kind = data.get('kind') or 'missing_punch'
    if kind not in dict(M.CorrectionRequest.KINDS):
        raise Refused('Pick a correction type.')
    pin = _parse_dt(data.get('proposed_in'), day)
    pout = _parse_dt(data.get('proposed_out'), day)
    if pin and pout and pout <= pin:
        pout += timedelta(days=1)
        if pout - pin > timedelta(hours=20):
            raise Refused('The clock-out time must be after the clock-in time.')
    if kind in ('missing_punch', 'wrong_time', 'forgot') and not pin and not pout:
        raise Refused('Give the correct clock-in and / or clock-out time.')
    reason = str(data.get('reason') or '').strip()
    if len(reason) < 3:
        raise Refused('Write the reason for the correction.')
    if M.CorrectionRequest.objects.filter(employee_id=employee.id, date=day, status__in=('manager', 'hr')).exists():
        raise Refused('There is already a correction waiting for approval for this date.')
    if not by_hr and rule.correction_max_per_month:
        n = M.CorrectionRequest.objects.filter(employee_id=employee.id, date__year=day.year, date__month=day.month) \
            .exclude(status='cancelled').count()
        if n >= rule.correction_max_per_month:
            raise Refused(f'You can request at most {rule.correction_max_per_month} corrections a month.')
    has_mgr = bool(employee.emp_reporting_manager_id) and employee.emp_reporting_manager_id != getattr(user, 'id', None)
    status = 'manager' if has_mgr else 'hr'
    cr = M.CorrectionRequest.objects.create(
        employee_id=employee.id, branch_id=employee.emp_branch_id_id, date=day, kind=kind, proposed_in=pin,
        proposed_out=pout, reason=reason, attachment=attachment, waive_penalty=str(data.get('waive_penalty')).lower() in ('1', 'true', 'yes'),
        status=status, created_by_id=getattr(user, 'id', None))
    if status == 'manager':
        _notify_user(employee.emp_reporting_manager, 'Attendance correction to approve',
                     f'{_name(employee)} asked to correct attendance on {day:%d %b %Y}.')
    elif not rule.correction_needs_hr and not has_mgr:
        pass
    return cr


def _name(e):
    return ' '.join(x for x in [e.emp_first_name, e.emp_last_name] if x) or e.emp_code


def _notify_user(user, title, message):
    if not user:
        return
    try:
        from zeo.module_helpers import notify
        notify(user=user, title=title[:100], message=message, notification_type='attendance')
    except Exception:
        pass


def manager_decide(cr, user, approve, note=''):
    from EmpManagement.models import emp_master
    if cr.status != 'manager':
        raise Refused('This correction is not waiting for the manager.')
    e = emp_master.objects.get(pk=cr.employee_id)
    rule = E.rule_for(e)
    cr.manager_user_id = user.id
    cr.manager_at = timezone.now()
    cr.manager_note = note or ''
    if not approve:
        cr.status = 'rejected'
        cr.save()
        _notify_user(e.users, 'Attendance correction rejected', f'Your manager rejected the correction for {cr.date:%d %b %Y}. {note}')
        return cr
    if rule.correction_needs_hr:
        cr.status = 'hr'
        cr.save()
        return cr
    cr.save()
    return apply_correction(cr, user)


def hr_decide(cr, user, approve, note=''):
    from EmpManagement.models import emp_master
    if cr.status != 'hr':
        raise Refused('This correction is not waiting for HR.')
    e = emp_master.objects.get(pk=cr.employee_id)
    cr.hr_user_id = user.id
    cr.hr_at = timezone.now()
    cr.hr_note = note or ''
    if not approve:
        cr.status = 'rejected'
        cr.save()
        _notify_user(e.users, 'Attendance correction rejected', f'HR rejected the correction for {cr.date:%d %b %Y}. {note}')
        return cr
    cr.save()
    return apply_correction(cr, user)


def _snapshot(emp_id, day):
    ps = M.Punch.objects.filter(employee_id=emp_id, work_date=day, is_void=False).order_by('ts')
    d = M.AttendanceDay.objects.filter(employee_id=emp_id, date=day).first()
    return {
        'punches': [{'id': p.id, 'ts': p.ts.isoformat(), 'kind': p.kind, 'source': p.source} for p in ps],
        'day': ({'status': d.status, 'first_in': d.first_in and d.first_in.isoformat(), 'last_out': d.last_out and d.last_out.isoformat(),
                 'worked_minutes': d.worked_minutes, 'late_minutes': d.late_minutes, 'early_minutes': d.early_minutes} if d else None),
    }


def _put_punch(emp_id, ts, kind, cr, user):
    old = M.Punch.objects.filter(employee_id=emp_id, ts=ts).first()
    if old:
        old.is_void, old.void_reason, old.kind, old.work_date, old.correction = False, '', kind, cr.date, cr
        old.save(update_fields=['is_void', 'void_reason', 'kind', 'work_date', 'correction'])
        return old
    return M.Punch.objects.create(employee_id=emp_id, ts=ts, kind=kind, resolved_kind=kind, source='correction',
                                  work_date=cr.date, correction=cr, verified_by='approval', created_by_id=user.id,
                                  note=f'Correction #{cr.id}')


@transaction.atomic
def apply_correction(cr, user):
    """Final approval: write the punches (source = correction), keep the old ones as void history,
    recompute the day and waive the late / early penalty when asked."""
    from EmpManagement.models import emp_master
    e = emp_master.objects.get(pk=cr.employee_id)
    rule = E.rule_for(e)
    cr.original = _snapshot(e.id, cr.date)
    current = E.compute(e, cr.date, rule)
    sh = E.shift_info(e, cr.date, rule)
    new_in = cr.proposed_in or current.get('first_in') or (sh or {}).get('start')
    new_out = cr.proposed_out or current.get('last_out') or ((sh or {}).get('end') if cr.kind in ('on_duty', 'wfh') else None)
    if not new_in:
        raise Refused('There is no clock-in time to use – edit the request with the correct times.')
    M.Punch.objects.filter(employee_id=e.id, work_date=cr.date, is_void=False, kind__in=('in', 'out', 'auto')) \
        .update(is_void=True, void_reason=f'Replaced by correction #{cr.id}')
    _put_punch(e.id, new_in, 'in', cr, user)
    if new_out:
        _put_punch(e.id, new_out, 'out', cr, user)
    cr.status = 'approved'
    cr.save()
    rec = E.recompute_day(e, cr.date, rule)
    if cr.waive_penalty:
        rec.penalty_waived = True
        rec.save(update_fields=['penalty_waived'])
    try:
        from Chatter.models import Message
        Message.objects.create(model='AttendancePlus.correctionrequest', object_id=str(cr.id), author=user,
                               body=f'Approved. Attendance on {cr.date} now {rec.status} '
                                    f'({rec.worked_minutes // 60}h {rec.worked_minutes % 60}m). Old punches kept as void history.')
    except Exception:
        pass
    _notify_user(e.users, 'Attendance correction approved', f'Your attendance for {cr.date:%d %b %Y} was corrected.')
    return cr
