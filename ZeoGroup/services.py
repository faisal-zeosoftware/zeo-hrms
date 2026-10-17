"""v1.27.0 – group dashboard, group reports, one approval inbox, one notification feed.

Every figure is read company by company (`schema_context`), with the rights of a company administrator for the
summaries and reports (see scope.acting_user) and with the person's own rights for the approval inbox. Own-database
companies work the same way: the connection follows the company (ZeoCloud router).
"""
import datetime as dt
import logging
from decimal import Decimal

from django.core.cache import cache
from django.db.models import Q, Sum
from django.utils import timezone
from django_tenants.utils import schema_context

log = logging.getLogger(__name__)
SUMMARY_TTL = 15 * 60

GROUP_REPORTS = [
    # key, title (the company reports of the same key run in each company and are put together)
    ('headcount', 'Headcount'),
    ('joiners-leavers', 'Joiners and leavers'),
    ('payroll-register', 'Payroll register'),
    ('payroll-cost', 'Payroll cost'),
    ('leave-balance', 'Leave balances'),
    ('leave-liability', 'Leave liability'),
    ('gratuity', 'Gratuity provision'),
    ('documents', 'Document expiry'),
    ('requests-pending', 'Requests waiting for approval'),
]


def drf_request(user, params=None, lang=None):
    """A request object for the company views and reports (they read user, query_params, headers)."""
    from django.test import RequestFactory
    from rest_framework.request import Request
    headers = {'HTTP_X_LANGUAGE': lang} if lang else {}
    r = Request(RequestFactory().get('/dashboard/api/group/', data=params or {}, **headers))
    r.user = user
    r._zeo_ctx = None
    return r


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


# ------------------------------------------------------------------ dashboard
def company_summary(comp, user, fresh=False):
    key = f'zeo:group:sum:{comp.schema_name}'
    if not fresh:
        hit = cache.get(key)
        if hit:
            return dict(hit, cached=True)
    today = timezone.localdate()
    m_from = today.replace(day=1)
    out = {'company_id': comp.id, 'company': comp.name, 'schema_name': comp.schema_name, 'error': '',
           'employees': 0, 'joiners': 0, 'leavers': 0, 'on_leave': 0, 'payroll_month': '', 'payroll_gross': 0.0,
           'payroll_net': 0.0, 'payslips': 0, 'documents_30': 0, 'approvals': 0, 'at': timezone.now().isoformat()}
    try:
        with schema_context(comp.schema_name):
            from django.apps import apps
            E = apps.get_model('EmpManagement', 'emp_master')
            act = E.objects.filter(is_active=True)
            out['employees'] = act.count()
            out['joiners'] = E.objects.filter(emp_joined_date__gte=m_from, emp_joined_date__lte=today).count()
            R = apps.get_model('EmpManagement', 'EmployeeResignation')
            out['leavers'] = R.objects.exclude(status__iexact='rejected').filter(
                Q(last_working_date__gte=m_from, last_working_date__lte=today) |
                Q(last_working_date__isnull=True, resigned_on__gte=m_from, resigned_on__lte=today)).count()
            L = apps.get_model('calendars', 'employee_leave_request')
            out['on_leave'] = L.objects.filter(status__iexact='approved', start_date__lte=today, end_date__gte=today,
                                               employee__is_active=True).values('employee').distinct().count()
            Run = apps.get_model('PayrollManagement', 'PayrollRun')
            Slip = apps.get_model('PayrollManagement', 'Payslip')
            run = Run.objects.filter(payslips__isnull=False).order_by('-year', '-month').distinct().first()
            if run is not None:
                same = Run.objects.filter(year=run.year, month=run.month)
                agg = Slip.objects.filter(payroll_run__in=same).aggregate(g=Sum('gross_salary'), n=Sum('net_salary'))
                out['payroll_month'] = f'{run.year}-{int(run.month):02d}'
                out['payroll_gross'] = _num(agg['g'])
                out['payroll_net'] = _num(agg['n'])
                out['payslips'] = Slip.objects.filter(payroll_run__in=same).count()
            req = drf_request(user, {'from': today.isoformat(), 'to': (today + dt.timedelta(days=30)).isoformat()})
            out['documents_30'] = _report_count(req, 'documents')
            out['approvals'] = _report_count(drf_request(user), 'requests-pending')
    except Exception as e:
        log.exception('group summary of %s', comp.schema_name)
        out['error'] = str(e)[:200]
    cache.set(key, out, SUMMARY_TTL)
    return out


def _report_count(req, key):
    try:
        from DashboardManagement.reports import REPORTS, Run
        spec = REPORTS.get(key)
        if not spec:
            return 0
        res = spec[3](Run(req, {'period': spec[4]}))
        return len(res[1] or [])
    except Exception:
        log.exception('group report count %s', key)
        return 0


def _one(args):
    from django.db import close_old_connections
    comp, user, fresh = args
    try:
        return company_summary(comp, user, fresh)
    finally:
        close_old_connections()     # each worker thread has its own connection


def dashboard(group, user, fresh=False):
    comps = group['companies']
    if len(comps) > 2:     # four companies at a time (each in its own connection / database)
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=4) as ex:
            rows = list(ex.map(_one, [(c, user, fresh) for c in comps]))
    else:
        rows = [company_summary(c, user, fresh) for c in comps]
    keys = ['employees', 'joiners', 'leavers', 'on_leave', 'payroll_gross', 'payroll_net', 'payslips', 'documents_30', 'approvals']
    total = {k: sum(_num(r.get(k)) for r in rows) for k in keys}
    return {'companies': rows, 'total': total, 'today': timezone.localdate().isoformat()}


# ------------------------------------------------------------------ reports
def report_list():
    from DashboardManagement.reports import REPORTS
    return [{'key': k, 'title': REPORTS[k][0] if k in REPORTS else t, 'period': REPORTS[k][4] if k in REPORTS else None}
            for k, t in GROUP_REPORTS if k in REPORTS]


def run_report(group, user, key, params, lang=None):
    """The company report `key` in every company of the group, one table with a Company column first."""
    from DashboardManagement.reports import REPORTS, ReportView
    if key not in dict(GROUP_REPORTS) or key not in REPORTS:
        raise ValueError('Unknown group report.')
    cols, rows, totals, notes, errors, per_company = None, [], {}, [], [], []
    title = REPORTS[key][0]
    for comp in group['companies']:
        try:
            with schema_context(comp.schema_name):
                req = drf_request(user, {k: v for k, v in (params or {}).items() if k in ('from', 'to', 'as_of')}, lang)
                r = ReportView().get(req, key)
                if r.status_code != 200:
                    errors.append({'company': comp.name, 'error': str((r.data or {}).get('detail') or r.status_code)})
                    continue
                d = r.data
        except Exception as e:
            log.exception('group report %s in %s', key, comp.schema_name)
            errors.append({'company': comp.name, 'error': str(e)[:200]})
            continue
        title = d.get('title') or title
        if cols is None:
            cols = [{'key': '_company', 'label': 'Company', 'type': 'text'}] + list(d.get('columns') or [])
        for row in d.get('rows') or []:
            rows.append(dict(row, _company=comp.name, _schema=comp.schema_name))
        for k, v in (d.get('totals') or {}).items():
            totals[k] = _num(totals.get(k)) + _num(v)
        per_company.append({'company': comp.name, 'count': d.get('count', len(d.get('rows') or [])), 'totals': d.get('totals') or {}})
        for n in d.get('notes') or []:
            if n not in notes:
                notes.append(n)
    return {'key': key, 'title': title, 'columns': cols or [{'key': '_company', 'label': 'Company', 'type': 'text'}],
            'rows': rows, 'totals': totals, 'count': len(rows), 'per_company': per_company, 'notes': notes, 'errors': errors,
            'from': (params or {}).get('from'), 'to': (params or {}).get('to')}


def report_xlsx(data):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = (data.get('title') or 'Report')[:31]
    cols = data['columns']
    ws.append([c.get('label') or c['key'] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data['rows']:
        ws.append([_cell(r.get(c['key'])) for c in cols])
    if data.get('totals'):
        ws.append([('Total' if i == 0 else (data['totals'].get(c['key']) if c['key'] in data['totals'] else None)) for i, c in enumerate(cols)])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max(len(str(c.value or '')) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, (dict, list)):
        return str(v)
    if isinstance(v, Decimal):
        return float(v)
    return v


# ------------------------------------------------------------------ approval inbox and notifications
def inbox(group, user):
    """The person's own approvals in every company of the group they work in (their own rights in each)."""
    from DashboardManagement.services import pending_approvals
    items, errors = [], []
    mine = set(user.tenants.values_list('schema_name', flat=True))
    for comp in group['companies']:
        if comp.schema_name not in mine and not user.is_superuser:
            continue
        try:
            with schema_context(comp.schema_name):
                for i in pending_approvals(user, drf_request(user)):
                    d = {k: (v.isoformat() if hasattr(v, 'isoformat') else v) for k, v in i.items()}
                    d.update(company=comp.name, schema_name=comp.schema_name, company_id=comp.id)
                    items.append(d)
        except Exception as e:
            log.exception('group inbox %s', comp.schema_name)
            errors.append({'company': comp.name, 'error': str(e)[:200]})
    items.sort(key=lambda x: (x.get('date') or '9999'))
    return {'items': items, 'count': len(items), 'errors': errors}


def feed(group, user, limit=60):
    from UserManagement.models import UserNotificationInbox as N
    names = {c.schema_name: c.name for c in group['companies']}
    qs = N.objects.filter(user=user, schema_name__in=list(names)).order_by('-created_at')
    rows = [{'id': n.id, 'title': n.title, 'message': n.message, 'type': n.notification_type, 'company': names.get(n.schema_name, n.schema_name),
             'schema_name': n.schema_name, 'branch': n.branch_name, 'is_read': n.is_read, 'created_at': n.created_at.isoformat()}
            for n in qs[:limit]]
    return {'items': rows, 'unread': qs.filter(is_read=False).count()}


def feed_read(group, user, ids=None):
    from UserManagement.models import UserNotificationInbox as N
    qs = N.objects.filter(user=user, schema_name__in=[c.schema_name for c in group['companies']], is_read=False)
    if ids:
        qs = qs.filter(pk__in=[int(i) for i in ids if str(i).isdigit()])
    return qs.update(is_read=True)
