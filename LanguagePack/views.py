"""v1.15.0 – Language packs API (base /language/api/; `?schema=` like every call – languages/ and pack/ also work without it).

GET   languages/            anyone (also before login): active languages [{code, name, english_name, dir, date_format, version}]
GET   pack/<lang>/          logged in: the whole pack; not logged in: only the login-page terms. ETag / If-None-Match → 304
GET   me/  PUT me/          logged in: the user's language ({lang: 'ar'} / {lang: null} = follow the company default)
GET   settings/ PATCH       company admin: default language + languages the company uses
GET   terms/                company admin or `change_term`: ?lang=ar&q=&missing=1&reviewed=0|1&module=&page=&page_size=
PATCH terms/<id>/           {value, is_reviewed}
POST  terms/review/         {ids: [...], is_reviewed: true}
POST  terms/import/         multipart file (xlsx / csv: English, Translation[, Reviewed]) + lang [+ dry_run=1 → preview only]
GET   terms/export/?lang=   xlsx (English, Translation, Module, Reviewed, Where)
"""
import csv
import io

from django.db import connection
from django.db.models import Count, Q
from django.http import HttpResponse
from rest_framework import status
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.authentication import JWTAuthentication

from . import services as S
from .models import CompanyLanguageSettings, Language, Term

t = S.t
MAX_PAGE = 500
MAX_IMPORT_ROWS = 50000
MAX_IMPORT_BYTES = 10 * 1024 * 1024


class LenientJWT(JWTAuthentication):
    """Public endpoints: a missing / expired token = anonymous instead of 401 (the login page may still hold an old one)."""
    def authenticate(self, request):
        try:
            return super().authenticate(request)
        except Exception:
            return None


def _schema():
    s = getattr(connection, 'schema_name', 'public') or 'public'
    return None if s == 'public' else s


def _ctx(request):
    try:
        from AccessControl.access import ctx
        return ctx(request)
    except Exception:
        return None


def is_admin(request):
    u = request.user
    if getattr(u, 'is_superuser', False):
        return True
    if not _schema():
        return False
    c = _ctx(request)
    return bool(c and c.admin)


def can_translate(request):
    if is_admin(request):
        return True
    if not _schema():
        return False
    c = _ctx(request)
    return bool(c and 'change_term' in c.codes)


def lang_row(r):
    return {'code': r['code'], 'name': r['name'], 'english_name': r['english_name'], 'dir': r['dir'],
            'date_format': r['date_format'], 'version': r['version']}


def _languages(schema):
    allc = S._languages_all()
    return [lang_row(allc[c]) for c in S.company_languages(schema) if c in allc]


# ------------------------------------------------------------------------------------------------------------------ public
class LanguagesView(APIView):
    zeo_access = False
    authentication_classes = [LenientJWT]
    permission_classes = [AllowAny]

    def get(self, request):
        return Response(_languages(_schema()))


class PackView(APIView):
    zeo_access = False
    authentication_classes = [LenientJWT]
    permission_classes = [AllowAny]

    def get(self, request, lang):
        code = S.base_lang(lang)
        L = S.language(code)
        if not L or not L['active']:
            return Response({'detail': t('Unknown language "{0}".', None, [lang])}, status=status.HTTP_404_NOT_FOUND)
        public_only = not (request.user and request.user.is_authenticated)
        etag = f'W/"{code}-{L["version"]}-{"p" if public_only else "f"}"'
        if request.META.get('HTTP_IF_NONE_MATCH') == etag:
            r = HttpResponse(status=304)
            r['ETag'] = etag
            return r
        r = Response(S.pack_payload(code, public_only=public_only))
        r['ETag'] = etag
        r['Cache-Control'] = 'private, max-age=0, must-revalidate'
        return r


# --------------------------------------------------------------------------------------------------------------------- me
def me_payload(request):
    schema = _schema()
    uid = request.user.pk
    own = S.user_pref(uid)
    allowed = S.company_languages(schema)
    lang, src = S.resolve(None, uid, schema)
    L = S.language(lang) or {}
    return {'lang': lang, 'source': src, 'user_lang': own if own in allowed else None,
            'company_default': S.company_default(schema), 'dir': L.get('dir', 'ltr'), 'version': L.get('version', 1),
            'date_format': L.get('date_format', 'dd/MM/yyyy'), 'languages': _languages(schema)}


class MeView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(me_payload(request))

    def put(self, request):
        data = request.data if isinstance(request.data, dict) else {}
        if 'lang' not in data:
            raise ValidationError({'lang': [t('Choose a language.')]})
        v = data.get('lang')
        code = S.base_lang(v) if v not in (None, '') else ''
        if code:
            allowed = S.company_languages(_schema())
            if code not in allowed:
                names = ', '.join((S.language(c) or {}).get('english_name', c) for c in allowed)
                raise ValidationError({'lang': [t('This language is not available. Choose one of: {0}.', None, [names])]})
        S.set_user_pref(request.user.pk, code)
        return Response(me_payload(request))

    patch = put


# --------------------------------------------------------------------------------------------------------------- settings
class _Admin(APIView):
    permission_classes = [IsAuthenticated]

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if _schema() is None:
            raise ValidationError({'detail': t('Choose a company first (?schema=).')})
        if not is_admin(request):
            raise PermissionDenied(t('Only the company administrator can change the language settings.'))


def settings_payload(schema):
    st = S.company_settings(schema)
    company = S.company_languages(schema)
    counts = {r['lang']: r for r in Term.objects.values('lang').annotate(
        total=Count('id', filter=Q(in_source=True)),
        missing=Count('id', filter=Q(in_source=True, value='')),
        reviewed=Count('id', filter=Q(in_source=True, is_reviewed=True)))}
    langs = []
    for code, r in S._languages_all().items():
        if not r['active']:
            continue
        c = counts.get(code, {})
        langs.append({**lang_row(r), 'in_company': code in company, 'is_default': code == S.company_default(schema),
                      'is_source': code == S.SOURCE_LANG, 'terms_total': c.get('total', 0), 'terms_missing': c.get('missing', 0),
                      'terms_reviewed': c.get('reviewed', 0)})
    return {'default_lang': S.company_default(schema), 'active_langs': company, 'own_list': bool(st['active_langs']),
            'languages': langs}


class SettingsView(_Admin):
    def get(self, request):
        return Response(settings_payload(_schema()))

    def patch(self, request):
        schema = _schema()
        data = request.data if isinstance(request.data, dict) else {}
        for k in data:
            if k not in ('default_lang', 'active_langs'):
                raise ValidationError({k: [t('Unknown setting "{0}". Use default_lang or active_langs.', None, [k])]})
        known = S.active_codes()
        row, _ = CompanyLanguageSettings.objects.get_or_create(schema_name=schema)
        active = list(row.active_langs or []) or list(known)
        if 'active_langs' in data:
            v = data.get('active_langs')
            if not isinstance(v, list):
                raise ValidationError({'active_langs': [t('Send a list of language codes, for example ["en", "ar"].')]})
            codes = []
            for c in v:
                cc = S.base_lang(c)
                if cc not in known:
                    raise ValidationError({'active_langs': [t('Unknown language "{0}".', None, [c])]})
                if cc not in codes:
                    codes.append(cc)
            if S.SOURCE_LANG not in codes:
                codes.insert(0, S.SOURCE_LANG)
            active = codes
        default = S.base_lang(data.get('default_lang')) if data.get('default_lang') else (row.default_lang or S.SOURCE_LANG)
        if default not in active:
            if 'default_lang' in data:
                raise ValidationError({'default_lang': [t('Turn the language on for the company before making it the default.')]})
            default = S.SOURCE_LANG
        row.active_langs = active
        row.default_lang = default
        row.updated_by_id = request.user.pk
        row.save()
        S.forget_company(schema)
        return Response(settings_payload(schema))


# ------------------------------------------------------------------------------------------------------------------ terms
class _Translator(APIView):
    permission_classes = [IsAuthenticated]

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not can_translate(request):
            raise PermissionDenied(t('Only the company administrator or users with the right "Can change translation" can edit translations.'))


def _lang_param(request, required=True):
    code = S.base_lang(request.query_params.get('lang') or (request.data.get('lang') if hasattr(request.data, 'get') else ''))
    if not code:
        if required:
            raise ValidationError({'lang': [t('Choose the language to translate into (lang=ar).')]})
        return ''
    if code == S.SOURCE_LANG:
        raise ValidationError({'lang': [t('English is the source language – choose the language to translate into.')]})
    if not S.language(code):
        raise ValidationError({'lang': [t('Unknown language "{0}".', None, [code])]})
    return code


def term_row(r):
    return {'id': r.id, 'lang': r.lang, 'source': r.source, 'value': r.value, 'module': r.module, 'where': r.where[:5],
            'where_count': len(r.where or []), 'in_source': r.in_source, 'is_reviewed': r.is_reviewed, 'edited': r.edited,
            'placeholders': S.placeholders(r.source), 'updated_by': r.updated_by_name or None,
            'updated_at': r.updated_at if r.edited else None}


def filtered(request, lang):
    qp = request.query_params
    qs = Term.objects.filter(lang=lang)
    if qp.get('all') not in ('1', 'true'):
        qs = qs.filter(in_source=True)
    q = (qp.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(source__icontains=q) | Q(value__icontains=q))
    if qp.get('missing') in ('1', 'true'):
        qs = qs.filter(value='')
    elif qp.get('missing') in ('0', 'false'):
        qs = qs.exclude(value='')
    if qp.get('reviewed') in ('1', 'true'):
        qs = qs.filter(is_reviewed=True)
    elif qp.get('reviewed') in ('0', 'false'):
        qs = qs.filter(is_reviewed=False)
    if qp.get('edited') in ('1', 'true'):
        qs = qs.filter(edited=True)
    if qp.get('module'):
        qs = qs.filter(module=qp.get('module'))
    return qs


class TermsView(_Translator):
    def get(self, request):
        lang = _lang_param(request)
        qs = filtered(request, lang)
        try:
            page = max(1, int(request.query_params.get('page') or 1))
            size = min(MAX_PAGE, max(1, int(request.query_params.get('page_size') or 50)))
        except ValueError:
            raise ValidationError({'page': [t('Page and page size must be whole numbers.')]})
        total = qs.count()
        rows = list(qs.order_by('source', 'id')[(page - 1) * size: page * size])
        base = Term.objects.filter(lang=lang, in_source=True)
        counts = base.aggregate(total=Count('id'), missing=Count('id', filter=Q(value='')),
                                not_reviewed=Count('id', filter=Q(is_reviewed=False) & ~Q(value='')),
                                reviewed=Count('id', filter=Q(is_reviewed=True)))
        modules = list(base.values('module').annotate(count=Count('id')).order_by('module'))
        return Response({'lang': lang, 'count': total, 'page': page, 'page_size': size,
                         'results': [term_row(r) for r in rows], 'counts': counts, 'modules': modules})


def _check_value(source, value):
    """A translation keeps the placeholders of the English text."""
    if value and S.placeholders(source) != S.placeholders(value):
        need = ', '.join('{%s}' % p for p in S.placeholders(source)) or t('none')
        return t('Keep the placeholders of the English text in the translation: {0}.', None, [need])
    return ''


def _who(request):
    u = request.user
    return (u.get_full_name() if hasattr(u, 'get_full_name') else '') or getattr(u, 'username', '') or ''


class TermView(_Translator):
    def patch(self, request, pk):
        row = Term.objects.filter(pk=pk).first()
        if not row:
            return Response({'detail': t('This translation no longer exists. Refresh the list.')}, status=404)
        data = request.data if isinstance(request.data, dict) else {}
        for k in data:
            if k not in ('value', 'is_reviewed'):
                raise ValidationError({k: [t('Only the translation (value) and "reviewed" can be changed.')]})
        if 'value' in data:
            v = data.get('value')
            if v is not None and not isinstance(v, str):
                raise ValidationError({'value': [t('The translation must be text.')]})
            v = (v or '').strip()
            if len(v) > 4000:
                raise ValidationError({'value': [t('The translation is too long – use up to 4000 characters.')]})
            err = _check_value(row.source, v)
            if err:
                raise ValidationError({'value': [err]})
            row.value = v
            if 'is_reviewed' not in data and not v:
                row.is_reviewed = False
        if 'is_reviewed' in data:
            rv = data.get('is_reviewed')
            if not isinstance(rv, bool):
                raise ValidationError({'is_reviewed': [t('Reviewed must be true or false.')]})
            if rv and not row.value:
                raise ValidationError({'is_reviewed': [t('Translate the text before marking it reviewed.')]})
            row.is_reviewed = rv
        row.edited = True
        row.updated_by = request.user.pk
        row.updated_by_name = _who(request)
        row.save()
        S.bump(row.lang)
        return Response(term_row(row))


class ReviewView(_Translator):
    def post(self, request):
        data = request.data if isinstance(request.data, dict) else {}
        ids = data.get('ids')
        rv = data.get('is_reviewed', True)
        if not isinstance(ids, list) or not ids or not all(isinstance(i, int) for i in ids):
            raise ValidationError({'ids': [t('Send the ids of the translations to mark, for example [12, 13].')]})
        if not isinstance(rv, bool):
            raise ValidationError({'is_reviewed': [t('Reviewed must be true or false.')]})
        qs = Term.objects.filter(pk__in=ids)
        if rv:
            qs = qs.exclude(value='')
        langs = set(qs.values_list('lang', flat=True))
        n = qs.update(is_reviewed=rv, edited=True, updated_by=request.user.pk, updated_by_name=_who(request))
        for lg in langs:
            S.bump(lg)
        return Response({'updated': n})


# ----------------------------------------------------------------------------------------------------------- excel in/out
HEAD_SOURCE = ('english', 'source', 'english text', 'text', 'en')
HEAD_VALUE = ('translation', 'value', 'translated', 'translated text')
HEAD_REVIEWED = ('reviewed', 'is_reviewed', 'is reviewed')


def read_table(f):
    name = (getattr(f, 'name', '') or '').lower()
    raw = f.read()
    if len(raw) > MAX_IMPORT_BYTES:
        raise ValidationError({'file': [t('The file is too large – use up to 10 MB.')]})
    if name.endswith('.csv') or name.endswith('.txt'):
        try:
            txt = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValidationError({'file': [t('Save the CSV file as UTF-8 and import it again.')]})
        return [list(r) for r in csv.reader(io.StringIO(txt))]
    if name.endswith('.xlsx') or raw[:2] == b'PK':
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [['' if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
            wb.close()
            return rows
        except Exception:
            raise ValidationError({'file': [t('This Excel file cannot be read. Use the file from "Export to Excel" or a CSV file.')]})
    raise ValidationError({'file': [t('Choose an Excel (.xlsx) or CSV file.')]})


def _yes(v):
    return str(v or '').strip().lower() in ('yes', 'y', 'true', '1', 'x', 'reviewed', 'نعم')


class ImportView(_Translator):
    def post(self, request):
        lang = _lang_param(request)
        f = request.FILES.get('file')
        if not f:
            raise ValidationError({'file': [t('Choose the Excel or CSV file to import.')]})
        dry = str(request.data.get('dry_run') or request.query_params.get('dry_run') or '').lower() in ('1', 'true', 'yes')
        rows = read_table(f)
        if not rows:
            raise ValidationError({'file': [t('The file is empty.')]})
        head = [str(h or '').strip().lower() for h in rows[0]]
        L = S.language(lang) or {}
        vnames = HEAD_VALUE + tuple(x.lower() for x in (L.get('english_name'), L.get('name'), lang) if x)
        si = next((i for i, h in enumerate(head) if h in HEAD_SOURCE), None)
        vi = next((i for i, h in enumerate(head) if h in vnames), None)
        ri = next((i for i, h in enumerate(head) if h in HEAD_REVIEWED), None)
        if si is None or vi is None:
            raise ValidationError({'file': [t('The first row must have the columns "English" and "Translation".')]})
        if len(rows) - 1 > MAX_IMPORT_ROWS:
            raise ValidationError({'file': [t('Import up to {0} rows at a time.', None, [MAX_IMPORT_ROWS])]})
        existing = {r.source_hash: r for r in Term.objects.filter(lang=lang)}
        from .models import source_hash
        res = {'lang': lang, 'dry_run': dry, 'rows': 0, 'new': 0, 'changed': 0, 'unchanged': 0, 'empty': 0, 'errors': [], 'samples': []}
        todo_new, todo_upd, seen = [], [], set()
        who, uid = _who(request), request.user.pk
        for n, r in enumerate(rows[1:], start=2):
            src = S.norm(r[si] if si < len(r) else '')
            val = str(r[vi] if vi < len(r) else '').strip()
            if not src and not val:
                continue
            res['rows'] += 1
            if not src:
                res['errors'].append({'row': n, 'english': '', 'message': t('The English text is empty.')})
                continue
            if not val:
                res['empty'] += 1
                continue
            err = _check_value(src, val)
            if err:
                res['errors'].append({'row': n, 'english': src[:120], 'message': err})
                continue
            h = source_hash(src)
            if h in seen:
                res['errors'].append({'row': n, 'english': src[:120], 'message': t('This English text is in the file twice – the first row is used.')})
                continue
            seen.add(h)
            cell = str(r[ri]).strip() if ri is not None and ri < len(r) and r[ri] is not None else ''
            rv = _yes(cell) if cell else None          # empty Reviewed cell = keep as it is
            row = existing.get(h)
            if row is None:
                res['new'] += 1
                todo_new.append(Term(lang=lang, source=src, source_hash=h, value=val, in_source=False, edited=True,
                                     is_reviewed=bool(rv), updated_by=uid, updated_by_name=who))
                if len(res['samples']) < 20:
                    res['samples'].append({'english': src, 'old': None, 'new': val})
            elif row.value != val or (rv is not None and rv != row.is_reviewed):
                res['changed'] += 1
                if len(res['samples']) < 20:
                    res['samples'].append({'english': src, 'old': row.value, 'new': val})
                row.value = val
                if rv is not None:
                    row.is_reviewed = rv
                row.edited, row.updated_by, row.updated_by_name = True, uid, who
                todo_upd.append(row)
            else:
                res['unchanged'] += 1
        res['errors'] = res['errors'][:200]
        if not dry and (todo_new or todo_upd):
            from django.db import transaction
            with transaction.atomic():
                Term.objects.bulk_create(todo_new, batch_size=1000)
                Term.objects.bulk_update(todo_upd, ['value', 'is_reviewed', 'edited', 'updated_by', 'updated_by_name'], batch_size=1000)
                Term.objects.filter(pk__in=[x.pk for x in todo_upd]).update(updated_at=_now())
            S.bump(lang)
        res['message'] = (t('Preview: {0} new, {1} changed, {2} unchanged, {3} without translation, {4} with errors. Nothing is saved yet.',
                            None, [res['new'], res['changed'], res['unchanged'], res['empty'], len(res['errors'])]) if dry else
                          t('Imported: {0} new, {1} changed, {2} unchanged, {3} without translation, {4} with errors.',
                            None, [res['new'], res['changed'], res['unchanged'], res['empty'], len(res['errors'])]))
        return Response(res)


def _now():
    from django.utils import timezone
    return timezone.now()


class ExportView(_Translator):
    def get(self, request):
        lang = _lang_param(request)
        qs = filtered(request, lang).order_by('module', 'source')
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        L = S.language(lang) or {}
        wb = Workbook()
        ws = wb.active
        ws.title = (L.get('english_name') or lang)[:30]
        ws.append(['English', 'Translation', 'Module', 'Reviewed', 'Where'])
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='4338CA')
        for r in qs.iterator(chunk_size=2000):
            ws.append([r.source, r.value, r.module, 'Yes' if r.is_reviewed else 'No', ', '.join((r.where or [])[:3])])
        for col, w in zip('ABCDE', (60, 60, 22, 10, 50)):
            ws.column_dimensions[col].width = w
        rtl = L.get('dir') == 'rtl'
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, horizontal='right' if rtl else 'left', readingOrder=2 if rtl else 1)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="zeo-translations-{lang}.xlsx"'
        return resp
