"""Kit API views. Each view is made per app with View.as_view(app=the_app)."""
from io import BytesIO

from django.db import models as M, transaction
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse
from rest_framework import status
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from . import packs as P
from .engine import (REGISTRY, AppError, CompanyMember, In, action_ok, d, detail, field_of, fmeta, guard, jsonable, link_options,
                     need, other_app, ser, save, tables, user_name)


class Base(APIView):
    app = None
    zeo_access = False   # rights are checked here
    zeo_scope = False
    permission_classes = [CompanyMember]
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def zeo_model_for(self, kwargs):
        """v1.22.0: the record panel (notes, activities, files, history, extra fields) finds the register's model."""
        r = self.app.regs.get((kwargs or {}).get('kind')) if self.app is not None else None
        return r.model if r is not None else None


def can_setup(app, request):
    return app.is_admin(request) or app.is_manager(request)


def links_state(app):
    s = app.settings()
    out = []
    for label, title, what in app.links:
        other = REGISTRY.get(label)
        installed = other is not None and other.settings_model._meta.db_table.lower() in tables()
        active = bool(installed and other.activated())
        off = label in (s.links_off or [])
        out.append({'app': label, 'title': title, 'what': what, 'installed': installed, 'activated': active, 'off': off,
                    'working': active and not off, 'page': other.page if other else ''})
    return out


def app_info(app, request):
    s = app.settings()
    code, label, sug = P.suggested(app)
    r = app.rights(request)
    return jsonable({
        'app': app.label, 'slug': app.slug, 'title': app.title, 'activated': bool(s.activated_at), 'activated_at': s.activated_at,
        'activated_by': user_name(s.activated_by), 'industry': s.industry, 'industry_label': app.packs.get(s.industry, {}).get('label', s.industry),
        'industry_source': s.industry_source, 'company_industry': code, 'company_industry_label': label, 'suggested': sug,
        'packs_installed': s.packs_installed, 'terms': s.terms or {}, 'require_review': s.require_review,
        'rights': r, 'can_setup': can_setup(app, request), 'links': links_state(app),
        'settings': {f: getattr(s, f) for f in app.settings_fields},
        'settings_meta': [_smeta(app, f) for f in app.settings_fields],
        'to_review': len(P.to_review(app)) if s.activated_at else 0,
        'roles': [{'code': k, 'label': v} for k, v in app.roles],
    })


def _smeta(app, name):
    from .engine import ftype, label_of
    f = app.settings_model._meta.get_field(name)
    m = {'name': name, 'label': label_of(f), 'type': ftype(f), 'help': str(f.help_text or '')}
    if f.choices:
        m['choices'] = [{'value': k, 'label': str(v)} for k, v in f.choices]
    return m


class AppInfoView(Base):
    needs_activation = False

    @guard
    def get(self, request):
        return Response(app_info(self.app, request))

    @guard
    def patch(self, request):
        """Manager: require_review, terms and the app's own settings."""
        a = self.app
        need(can_setup(a, request), f'Only the {a.title} manager or an administrator can change settings.')
        f = In(request)
        s = a.settings()
        changed = []
        if f.has('require_review'):
            s.require_review = f.b('require_review')
            changed.append('review rule')
        if f.has('terms'):
            t = f.j('terms') or {}
            if not isinstance(t, dict):
                raise AppError('Terms are not valid.', 'terms')
            s.terms = {str(k)[:40]: str(v)[:60] for k, v in t.items() if str(v).strip()}
            changed.append('wording')
        from .engine import ftype
        for n in a.settings_fields:
            if not f.has(n):
                continue
            fld = a.settings_model._meta.get_field(n)
            t = ftype(fld)
            if t == 'bool':
                v = f.b(n)
            elif t == 'decimal':
                v = f.n(n, fld.default)
            elif t == 'int':
                v = f.i(n)
                v = fld.default if v is None else v
            elif t == 'json':
                v = f.j(n)
            elif t == 'choice':
                v = f.choice(n, fld.choices, fld.default)
            else:
                v = f.s(n)
            setattr(s, n, v)
            changed.append(str(fld.verbose_name))
        s.save()
        if changed:
            a.event('settings', 0, 'Settings changed: ' + ', '.join(changed), request.user)
        return Response(app_info(a, request))


class ActivateView(Base):
    needs_activation = False

    @guard
    def get(self, request):
        a = self.app
        out = app_info(a, request)
        out['packs'] = [P.summary(a, c) for c in a.packs]
        out['can_activate'] = can_setup(a, request)
        return Response(out)

    @guard
    def post(self, request):
        a = self.app
        need(can_setup(a, request), f'Only the {a.title} manager or an administrator can activate {a.title}.')
        f = In(request)
        added = P.activate(a, f.s('industry') or 'general', request.user)
        r = self.get(request)
        r.data['added'] = added
        return r


class PackView(Base):
    @guard
    def get(self, request):
        a = self.app
        s = a.settings()
        code, label, sug = P.suggested(a)
        return Response({'industry': s.industry, 'industry_source': s.industry_source, 'company_industry_label': label, 'suggested': sug,
                         'installed': s.packs_installed, 'packs': [P.summary(a, c) for c in a.packs], 'to_review': P.to_review(a),
                         'can_setup': can_setup(a, request)})

    @guard
    def post(self, request):
        """{code, primary}: add a pack (only adds); primary also makes it this app's industry."""
        a = self.app
        need(can_setup(a, request), 'Only a manager or an administrator can install packs.')
        f = In(request)
        code = f.s('code')
        added = a.install(code, request.user, primary=f.b('primary'))   # v1.33.0: through the app (CRM also seeds the masters)
        if f.b('primary'):
            s = a.settings()
            _, _, sug = P.suggested(a)
            s.industry_source = 'company' if code == sug else 'chosen'
            s.save()
        r = self.get(request)
        r.data['added'] = added
        return r


class ReviewView(Base):
    @guard
    def post(self, request):
        """{items: [{key, id}], action: reviewed|off}: confirm pack rows or switch them off."""
        a = self.app
        need(can_setup(a, request), 'Only a manager or an administrator can review pack items.')
        f = In(request)
        items = f.j('items') or []
        act = f.s('action') or 'reviewed'
        n = 0
        for it in items:
            key = str(it.get('key', ''))
            if key not in a.pack_models:
                continue
            model, _ = a.pack_models[key]
            x = model.objects.filter(pk=it.get('id')).first()
            if x is None:
                continue
            x.review_status = 'reviewed'
            if act == 'off':
                x.is_active = False
            x.save()
            n += 1
        a.event('settings', 0, f'{n} pack item(s) {"reviewed" if act != "off" else "switched off"} by {user_name(request.user)}', request.user)
        return Response({'done': n, 'to_review': P.to_review(a)})


class LinksView(Base):
    @guard
    def get(self, request):
        return Response(links_state(self.app))

    @guard
    def post(self, request):
        """{app, on}: switch the link to another app on or off (in this app)."""
        a = self.app
        need(can_setup(a, request), 'Only a manager or an administrator can change links.')
        f = In(request)
        label = f.s('app')
        if label not in [x[0] for x in a.links]:
            raise AppError('Unknown app.')
        s = a.settings()
        off = [x for x in (s.links_off or []) if x != label]
        if not f.b('on'):
            off.append(label)
        s.links_off = off
        s.save()
        a.event('settings', 0, f'Link to {label} switched {"on" if f.b("on") else "off"} by {user_name(request.user)}', request.user)
        return Response(links_state(a))


class DashboardView(Base):
    @guard
    def get(self, request):
        # v1.22.0: the period (?from=&to=) and the Branch picked on the dashboard
        b = request.query_params.get('branch')
        if b and str(b).isdigit():
            from .engine import branch_ids
            allowed = branch_ids(request)
            request._zeo_dash_branch = int(b) if allowed is None or int(b) in allowed else -1
        from .engine import dash_period
        frm, to = dash_period(request)
        out = self.app.dashboard(request) if self.app.dashboard else {}
        out['period'] = {'from': frm, 'to': to}
        return Response(jsonable(out))


class MyWorkView(Base):
    @guard
    def get(self, request):
        return Response(jsonable(self.app.my_work(request) if self.app.my_work else []))


# ------------------------------------------------------------------ registers
def get_reg(app, kind, request):
    r = app.regs.get(kind)
    if r is None:
        raise Http404
    need(r.can_open(request), f'You have no access to {r.plural.lower()}.')
    return r


def reg_meta(r, request):
    return {'kind': r.kind, 'label': r.label, 'plural': r.plural, 'group': r.group, 'help': r.help, 'child_of': r.child_of,
            'cols': [fmeta(r, c) for c in r.cols if not r.link_off(c)], 'form': [fmeta(r, c) for c in r.form if not r.link_off(c)],
            'filters': [fmeta(r, c) for c in r.filters if not r.link_off(c)], 'can_create': r.can_create(request) and r.create != 'none',
            'manages': r.manages(request), 'files': r.files, 'review': r.review}


class RegistersView(Base):
    @guard
    def get(self, request):
        return Response([reg_meta(r, request) for r in self.app.regs.values() if r.can_open(request)])


def filtered(r, request):
    q = request.query_params
    qs = r.scope(request)
    if q.get('q'):
        cond = Q()
        for s_ in r.search:
            cond |= Q(**{s_ + '__icontains': q['q']})
        qs = qs.filter(cond)
    for flt in r.filters:
        if q.get(flt) not in (None, ''):
            fld = field_of(r.model, flt)
            v = q[flt]
            if isinstance(fld, M.BooleanField):
                v = v in ('1', 'true', 'yes')
            qs = qs.filter(**{flt: v})
    if q.get('parent'):
        fk, _, pid = q['parent'].partition(':')
        if fk in [n.name for n in r.model._meta.fields] and pid.isdigit():
            qs = qs.filter(**{fk: pid})
    if q.get('open') == '1' and hasattr(r.model, 'status'):
        qs = qs.exclude(status__in=('closed', 'cancelled', 'archived', 'lost', 'expired', 'terminated', 'completed', 'rejected', 'paid', 'void'))
    if q.get('review') and hasattr(r.model, 'review_status'):
        qs = qs.filter(review_status=q['review'])
    df = q.get('date_field')   # v1.22.0: dashboard figures open their records for the same period
    if df and df in {f.name for f in r.model._meta.fields}:
        from django.utils.dateparse import parse_date
        fld = r.model._meta.get_field(df)
        suffix = '__date' if isinstance(fld, M.DateTimeField) else ''
        if parse_date(str(q.get('from') or '')):
            qs = qs.filter(**{f'{df}{suffix}__gte': parse_date(q['from'])})
        if parse_date(str(q.get('to') or '')):
            qs = qs.filter(**{f'{df}{suffix}__lte': parse_date(q['to'])})
    return qs


class RegisterListView(Base):
    @guard
    def get(self, request, kind):
        r = get_reg(self.app, kind, request)
        q = request.query_params
        qs = filtered(r, request)
        limit = max(1, min(1000, int(q.get('limit') or 200)))
        offset = max(0, int(q.get('offset') or 0))
        return Response({'count': qs.count(), 'results': jsonable([ser(r, x, request) for x in qs[offset:offset + limit]])})

    @guard
    def post(self, request, kind):
        a = self.app
        r = get_reg(a, kind, request)
        f = In(request)
        allowed = r.can_create(request) and r.create != 'none'
        if not allowed and r.child_of and f.has(r.child_of):
            pf = field_of(r.model, r.child_of)
            p = pf.related_model.objects.filter(pk=f.i(r.child_of)).first()
            pr = next((x for x in a.regs.values() if x.model is pf.related_model), None)
            allowed = bool(p and pr and pr.can_edit(p, request))
        need(allowed, f'You cannot add {r.plural.lower()}.')
        obj = save(r, r.model(), f, request, True)
        return Response(detail(r, r.model.objects.get(pk=obj.pk), request), status=status.HTTP_201_CREATED)


class RegisterDetailView(Base):
    @guard
    def get(self, request, kind, pk):
        r = get_reg(self.app, kind, request)
        obj = r.scope(request).filter(pk=pk).first()
        if obj is None:
            raise Http404
        return Response(detail(r, obj, request))

    @guard
    def patch(self, request, kind, pk):
        r = get_reg(self.app, kind, request)
        obj = r.scope(request).filter(pk=pk).first()
        if obj is None:
            raise Http404
        if not r.can_edit(obj, request):
            raise AppError('This record can no longer be changed here – use its actions.' if r.manages(request) else 'You cannot change this record.', status=403)
        obj = save(r, obj, In(request), request, False)
        return Response(detail(r, r.model.objects.get(pk=obj.pk), request))

    put = patch

    @guard
    def delete(self, request, kind, pk):
        """Only set-up records that nothing uses yet (pack rows, categories…) may be deleted; documents are cancelled instead."""
        r = get_reg(self.app, kind, request)
        obj = r.scope(request).filter(pk=pk).first()
        if obj is None:
            raise Http404
        need(r.manages(request) and getattr(r, 'deletable', False), 'This record cannot be deleted – cancel or archive it instead.')
        from django.db.models.deletion import ProtectedError
        try:
            obj.delete()
        except ProtectedError:
            raise AppError('This record is in use and cannot be deleted. Switch it off instead.')
        self.app.event(r.kind, pk, f'Deleted by {user_name(request.user)}', request.user)
        return Response(status=status.HTTP_204_NO_CONTENT)


def created_ref(res):
    for a in REGISTRY.values():
        for r in a.regs.values():
            if r.model is res.__class__:
                return {'app': a.slug, 'kind': r.kind, 'id': res.pk, 'number': getattr(res, 'number', '')}
    return None


class RegisterActionView(Base):
    @guard
    def post(self, request, kind, pk, action):
        r = get_reg(self.app, kind, request)
        obj = r.scope(request).filter(pk=pk).first()
        if obj is None:
            raise Http404
        a = next((x for x in r.actions if x.name == action), None)
        if a is None:
            raise Http404
        if not action_ok(r, a, obj, request):
            raise AppError(f'"{a.label}" is not possible now.', status=400)
        f = In(request)
        for n, l, t, req in a.fields:
            if req and t not in ('bool', 'file', 'files') and f.d.get(n) in (None, ''):
                raise AppError(f'{l} is required.', n)
        with transaction.atomic():
            res = a.run(obj, request, f)
        if isinstance(res, HttpResponse):
            return res
        body = detail(r, r.model.objects.get(pk=obj.pk), request)
        if isinstance(res, dict):
            body['result'] = jsonable(res)
        elif res is not None and hasattr(res, '_meta') and (res.__class__ is not obj.__class__ or res.pk != obj.pk):
            body['created'] = created_ref(res)
        return Response(body)


class RegisterFileView(Base):
    @guard
    def post(self, request, kind, pk):
        a = self.app
        r = get_reg(a, kind, request)
        obj = r.scope(request).filter(pk=pk).first()
        if obj is None:
            raise Http404
        need(r.manages(request) or r.is_own(obj, request) or r.can_edit(obj, request), 'You cannot attach files here.')
        f = In(request)
        xs = f.filelist()
        if not xs:
            raise AppError('Choose a file.', 'files')
        for x in xs:
            if (x.size or 0) > 20 * 1024 * 1024:
                raise AppError(f'{x.name} is larger than 20 MB.', 'files')
            a.file_model.objects.create(model=r.kind, object_id=obj.pk, file=x, name=x.name[:255], size=x.size or 0, note=f.s('note')[:255],
                                        portal=f.b('portal'), uploaded_by=request.user)
        a.event(r.kind, obj.pk, f'{len(xs)} file(s) added: ' + ', '.join(x.name for x in xs), request.user)
        return Response(detail(r, obj, request), status=status.HTTP_201_CREATED)


class FileView(Base):
    @guard
    def get(self, request, pk):
        a = self.app
        x = a.file_model.objects.filter(pk=pk).first()
        if x is None:
            raise Http404
        r = a.regs.get(x.model)
        need(r is not None and r.can_open(request) and r.scope(request).filter(pk=x.object_id).exists(), 'You cannot open this file.')
        return FileResponse(x.file.open('rb'), as_attachment=True, filename=x.name)


class OptionsView(Base):
    @guard
    def get(self, request, name):
        q = request.query_params
        ids = [int(x) for x in (q.get('ids') or '').split(',') if x.strip().isdigit()]
        limit = int(q.get('limit') or 500)
        if name.startswith('link:'):
            return Response(link_options(name[5:], q.get('q') or '', ids or None, limit, request))
        if name == 'branches':   # v1.22.0: the user's branches
            from OrganisationManager.models import brnch_mstr
            from .engine import branch_ids
            qs = brnch_mstr.objects.all().order_by('branch_name')
            b = branch_ids(request)
            if b is not None:
                qs = qs.filter(pk__in=b)
            if ids:
                qs = brnch_mstr.objects.filter(pk__in=ids)
            return Response([{'id': x.pk, 'label': x.branch_name} for x in qs[:limit]])
        if name == 'employees':
            from EmpManagement.models import emp_master
            from .engine import emp_name
            qs = emp_master.objects.filter(is_active=True) if hasattr(emp_master, 'is_active') else emp_master.objects.all()
            from .engine import branch_ids
            b = branch_ids(request)
            if b is not None:   # v1.22.0: employees of the user's branches
                qs = qs.filter(emp_branch_id__in=b)
            if ids:
                qs = emp_master.objects.filter(pk__in=ids)
            elif q.get('q'):
                qs = qs.filter(Q(emp_first_name__icontains=q['q']) | Q(emp_last_name__icontains=q['q']) | Q(emp_code__icontains=q['q']))
            return Response([{'id': e.pk, 'label': f'{emp_name(e)} ({e.emp_code})'} for e in qs.order_by('emp_first_name')[:limit]])
        if name == 'users':
            from tenant_users.tenants.models import UserTenantPermissions
            xs = [u.profile for u in _company_perms().select_related('profile') if u.profile and u.profile.is_active]
            if ids:
                xs = [u for u in xs if u.pk in ids]
            return Response([{'id': u.pk, 'label': user_name(u)} for u in xs[:limit]])
        if name not in self.app.options:
            raise Http404
        qs_fn, lbl, search = self.app.options[name]
        from .engine import in_branches
        qs = in_branches(qs_fn(), request)   # v1.22.0
        if q.get('parent'):   # e.g. quotes of one work order: parent=work_order:12
            fk, _, pid = q['parent'].partition(':')
            if pid.isdigit() and fk in {f.name for f in qs.model._meta.fields}:
                qs = qs.filter(**{fk: pid})
        if q.get('q'):
            cond = Q()
            for s_ in search:
                cond |= Q(**{s_ + '__icontains': q['q']})
            qs = qs.filter(cond)
        if ids:
            qs = qs.model.objects.filter(pk__in=ids)
        return Response([{'id': x.pk, 'label': lbl(x)} for x in qs[:limit]])


class RegisterExportView(Base):
    """Excel export of what the user can see (same filters as the list)."""
    @guard
    def get(self, request, kind):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        r = get_reg(self.app, kind, request)
        qs = filtered(r, request)[:20000]
        cols = [fmeta(r, c) for c in tuple(dict.fromkeys(r.cols + getattr(r, 'export_extra', ()))) if not r.link_off(c)]
        wb = Workbook()
        ws = wb.active
        ws.title = r.plural[:30]
        ws.append([c['label'] for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='334155')
        for x in qs:
            row = ser(r, x, request, [c['name'] for c in cols])
            ws.append([row.get(c['name'] + '_label', row.get(c['name'])) if not isinstance(row.get(c['name']), list)
                       else row.get(c['name'] + '_label', '') for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(c['label']) + 4))
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{self.app.slug}-{r.kind}.xlsx"'
        return resp



def _company_perms():   # v1.22.0: only the users of the open company (see AccessControl.access.company_perms)
    from AccessControl.access import company_perms
    return company_perms()
