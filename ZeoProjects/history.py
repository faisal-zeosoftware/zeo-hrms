"""Work history (v1.15.0): everything done for a customer, project or AMC contract in a period – time on support
tickets (AMC covered / chargeable) and on project tasks, the tickets handled and the state of the project tasks.

* On screen: Projects → Reports → "Work history – AMC and projects" (one row per time entry, Excel / CSV / PDF).
* As files: GET  /reports/work-history/file/?file=pdf|xlsx&<filters>  (a full report with sections).
* By e-mail: POST /reports/work-history/send/ {<filters>, to: "a@x.com, b@y.com", message} – the PDF and the Excel
  file attached; nothing is sent without the user pressing Send.
Rows follow the user's rights (tickets they may see, projects they may see).
"""
import io
from collections import defaultdict
from decimal import Decimal

from . import access as A
from . import services as S
from .models import AMCContract, Customer, ProjectProfile, Ticket, WorkLog

D0 = Decimal('0')
KIND_LABEL = {'amc': 'AMC support', 'chargeable': 'Chargeable support', 'support': 'Support', 'project': 'Project'}


def _hours_of(ts, extra):
    if extra is not None and not extra.running:
        return Decimal(extra.hours or 0)
    try:
        from ProjectControl.services import hours_of
        return hours_of(ts.time_spent)
    except Exception:
        return D0


def collect(request, p):
    """p = reports.P filters (+ p.contract). Returns entries, tickets, tasks, summary and a title."""
    from ProjectManagement.models import Project, Task, TimeSheet
    entries = []
    # ---- support time
    tq = A.ticket_scope(request, Ticket.objects.all())
    if p.customer:
        tq = tq.filter(customer_id=p.customer)
    if p.contract:
        tq = tq.filter(contract_id=p.contract)
    if p.project:
        tq = tq.filter(project_id=p.project)
    logs = WorkLog.objects.filter(ticket__in=tq, date__gte=p.frm, date__lte=p.to).exclude(status='rejected') \
        .select_related('ticket', 'ticket__customer', 'ticket__contract', 'employee').order_by('date', 'id')
    if p.employee:
        logs = logs.filter(employee_id=p.employee)
    for w in logs:
        kind = 'amc' if w.counts_against_amc else ('chargeable' if w.billable else 'support')
        entries.append({'date': w.date.isoformat(), 'customer': w.ticket.customer.name if w.ticket.customer_id else 'Internal',
                        'kind': KIND_LABEL[kind], 'kind_key': kind,
                        'reference': w.ticket.number + (f' · {w.ticket.contract.number}' if w.ticket.contract_id else ''),
                        'item': w.ticket.subject, 'employee': S.emp_name(w.employee), 'hours': float(w.hours),
                        'onsite': 'Onsite' if w.onsite else 'Remote', 'description': w.description or '',
                        'status': w.get_status_display(), 'link_to': {'page': 'tickets', 'id': w.ticket_id}})
    # ---- project time (not for a contract-only report)
    projects = []
    if not p.contract:
        pids = A.visible_project_ids(request)
        pq = Project.objects.filter(pk__in=pids)
        if p.project:
            pq = pq.filter(pk=p.project)
        if p.customer:
            pq = pq.filter(pk__in=ProjectProfile.objects.filter(customer_id=p.customer).values('project_id'))
        projects = list(pq.order_by('title'))
        ts = TimeSheet.objects.filter(project__in=projects, date__gte=p.frm, date__lte=p.to).select_related('project', 'task', 'employee') \
            .order_by('date', 'id')
        if p.employee:
            ts = ts.filter(employee_id=p.employee)
        ts = list(ts)
        try:
            from ProjectControl.models import TimesheetExtra
            extras = {x.timesheet_id: x for x in TimesheetExtra.objects.filter(timesheet_id__in=[t.pk for t in ts])}
        except Exception:
            extras = {}
        cust = dict(ProjectProfile.objects.filter(project_id__in=[x.pk for x in projects]).values_list('project_id', 'customer__name'))
        for t in ts:
            x = extras.get(t.pk)
            entries.append({'date': t.date.isoformat(), 'customer': cust.get(t.project_id) or 'Internal', 'kind': KIND_LABEL['project'],
                            'kind_key': 'project', 'reference': t.project.title, 'item': t.task.title if t.task_id else '',
                            'employee': S.emp_name(t.employee), 'hours': float(_hours_of(t, x)), 'onsite': '',
                            'description': t.description or '',
                            'status': (x.get_approval_status_display() if x is not None else t.get_status_display()),
                            'link_to': {'page': 'projects', 'id': t.project_id}})
    entries.sort(key=lambda r: (r['date'], r['kind']))
    # ---- tickets handled in the period
    tickets = []
    from django.db.models import Q
    for t in tq.filter(Q(created_at__date__gte=p.frm, created_at__date__lte=p.to) | Q(resolved_at__date__gte=p.frm, resolved_at__date__lte=p.to) |
                       Q(pk__in=logs.values('ticket_id'))).select_related('customer', 'assigned_to').order_by('created_at').distinct():
        tickets.append({'number': t.number, 'subject': t.subject, 'customer': t.customer.name if t.customer_id else 'Internal',
                        'stage': t.get_stage_display(), 'priority': t.get_priority_display(),
                        'created': t.created_at.date().isoformat(), 'resolved': t.resolved_at.date().isoformat() if t.resolved_at else '',
                        'engineer': S.emp_name(t.assigned_to) if t.assigned_to_id else '', 'billing': t.get_billable_status_display(),
                        'resolution': (t.resolution_note or '')[:300]})
    # ---- project tasks (state today)
    tasks = []
    hours_by_task = defaultdict(float)
    for e in entries:
        if e['kind_key'] == 'project' and e['item']:
            hours_by_task[(e['reference'], e['item'])] += e['hours']
    for pr in projects:
        for t in Task.objects.filter(project=pr).order_by('start_date', 'sequence', 'id').prefetch_related('task_members'):
            tasks.append({'project': pr.title, 'task': t.title, 'status': t.get_status_display(),
                          'start': t.start_date.isoformat() if t.start_date else '', 'due': t.end_date.isoformat() if t.end_date else '',
                          'people': ', '.join(S.emp_name(m) for m in t.task_members.all()),
                          'hours_in_period': round(hours_by_task.get((pr.title, t.title), 0), 2),
                          'overdue': bool(t.end_date and t.end_date < S.local_today() and t.status != 'completed')})
    summary = defaultdict(float)
    for e in entries:
        summary[e['kind_key']] += e['hours']
    title_bits = []
    if p.customer:
        title_bits.append(Customer.objects.filter(pk=p.customer).values_list('name', flat=True).first() or '')
    if p.contract:
        title_bits.append(AMCContract.objects.filter(pk=p.contract).values_list('number', flat=True).first() or '')
    if p.project:
        title_bits.append(next((x.title for x in projects if str(x.pk) == str(p.project)), ''))
    amc = None
    if p.contract:
        c = AMCContract.objects.filter(pk=p.contract).first()
        if c:
            u = S.contract_usage(c)
            amc = {'number': c.number, 'included': 'Unlimited' if c.unlimited_hours else float(u['included_hours']),
                   'used': float(u['used_hours']), 'remaining': None if u['remaining_hours'] is None else float(u['remaining_hours']),
                   'period': f'{c.start_date:%d %b %Y} – {c.end_date:%d %b %Y}'}
    return {'entries': entries, 'tickets': tickets, 'tasks': tasks, 'amc': amc,
            'summary': {k: round(v, 2) for k, v in summary.items()}, 'total_hours': round(sum(summary.values()), 2),
            'subject': ' · '.join(b for b in title_bits if b) or 'All customers',
            'from': p.frm.isoformat(), 'to': p.to.isoformat(), 'company': _company()}


def _company():
    try:
        from django.db import connection
        t = getattr(connection, 'tenant', None)
        for k in ('name', 'company_name', 'title'):
            if t is not None and getattr(t, k, None):
                return str(getattr(t, k))
    except Exception:
        pass
    return ''


# ------------------------------------------------------------------ files
ENTRY_COLS = [('date', 'Date'), ('customer', 'Customer'), ('kind', 'Type'), ('reference', 'Ticket / project'), ('item', 'Subject / task'),
              ('employee', 'Employee'), ('hours', 'Hours'), ('onsite', 'Where'), ('description', 'What was done'), ('status', 'Status')]
TICKET_COLS = [('number', 'Ticket'), ('subject', 'Subject'), ('customer', 'Customer'), ('priority', 'Priority'), ('stage', 'Stage'),
               ('created', 'Created'), ('resolved', 'Resolved'), ('engineer', 'Engineer'), ('billing', 'Billing'), ('resolution', 'Resolution')]
TASK_COLS = [('project', 'Project'), ('task', 'Task'), ('status', 'Status'), ('start', 'Start'), ('due', 'Due'), ('people', 'People'),
             ('hours_in_period', 'Hours in period')]


def build_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1E1B4B')

    def sheet(ws, cols, rows, widths=None):
        ws.append([c[1] for c in cols])
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append([r.get(k, '') for k, _ in cols])
        for i, (k, label) in enumerate(cols, 1):
            longest = max([len(str(r.get(k, ''))) for r in rows[:300]] + [len(label)])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(10, longest + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'

    ws = wb.active
    ws.title = 'Summary'
    ws.append(['Work history', data['subject']])
    ws.append(['Period', f"{data['from']} to {data['to']}"])
    ws.append([])
    ws.append(['Type', 'Hours'])
    for k in ('amc', 'chargeable', 'support', 'project'):
        if data['summary'].get(k):
            ws.append([KIND_LABEL[k], data['summary'][k]])
    ws.append(['Total', data['total_hours']])
    if data['amc']:
        a = data['amc']
        ws.append([])
        ws.append(['AMC', a['number'], a['period']])
        ws.append(['Hours included / used / left', a['included'], a['used'], a['remaining'] if a['remaining'] is not None else ''])
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 40
    sheet(wb.create_sheet('Time entries'), ENTRY_COLS, data['entries'])
    sheet(wb.create_sheet('Tickets'), TICKET_COLS, data['tickets'])
    if data['tasks']:
        sheet(wb.create_sheet('Project tasks'), TASK_COLS, data['tasks'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Work history – {data['subject']}")
    ss = getSampleStyleSheet()
    small = ParagraphStyle('s', parent=ss['BodyText'], fontSize=7.5, leading=9.5)
    h1 = ParagraphStyle('h1', parent=ss['Title'], fontSize=17, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle('h2', parent=ss['Heading2'], fontSize=11.5, textColor=colors.HexColor('#4338ca'), spaceBefore=10, spaceAfter=4)
    esc = lambda v: (str(v if v is not None else '')).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def table(cols, rows, widths):
        body = [[Paragraph(f'<b>{esc(c[1])}</b>', small) for c in cols]]
        body += [[Paragraph(esc(r.get(k, '')), small) for k, _ in cols] for r in rows]
        t = Table(body, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e0e7ff')), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                               ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d4d7e3')),
                               ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f8fc')])]))
        return t

    el = [Paragraph(f"Work history – {esc(data['subject'])}", h1),
          Paragraph(f"{esc(data['company'])}{' · ' if data['company'] else ''}Period {esc(data['from'])} to {esc(data['to'])}", ss['BodyText'])]
    sums = [[KIND_LABEL[k], f"{data['summary'][k]:.2f} h"] for k in ('amc', 'chargeable', 'support', 'project') if data['summary'].get(k)]
    sums.append(['Total', f"{data['total_hours']:.2f} h"])
    if data['amc']:
        a = data['amc']
        left = '-' if a['remaining'] is None else f"{a['remaining']:.2f}"
        sums.append([f"AMC {a['number']} ({a['period']})", f"included {a['included']} · used {a['used']:.2f} · left {left}"])
    st = Table(sums, colWidths=[90 * mm, 90 * mm])
    st.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 9), ('LINEBELOW', (0, -1), (-1, -1), 0.5, colors.grey),
                            ('FONTNAME', (0, -1 if not data['amc'] else -2), (-1, -1 if not data['amc'] else -2), 'Helvetica-Bold')]))
    el += [Spacer(1, 6), st]
    el += [Paragraph('Time spent', h2)]
    el += [table(ENTRY_COLS, data['entries'], [18, 30, 24, 32, 42, 26, 12, 14, 56, 22])] if data['entries'] else [Paragraph('No time logged in the period.', small)]
    el += [Paragraph('Tickets', h2)]
    el += [table(TICKET_COLS, data['tickets'], [18, 50, 30, 16, 26, 18, 18, 26, 22, 49])] if data['tickets'] else [Paragraph('No tickets in the period.', small)]
    if data['tasks']:
        el += [Paragraph('Project tasks', h2), table(TASK_COLS, data['tasks'], [50, 75, 24, 20, 20, 52, 22])]
    doc.build(el)
    return buf.getvalue()
