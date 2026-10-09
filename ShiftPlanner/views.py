"""
v1.12.0 – Shift planner API (mounted at /shift-planner/api/, every call with ?schema=<company>).

Rights (the central AccessControl layer is switched off for these views because the tables keep plain ids; the same
rules are applied here):
  * HR / admin ("manage"): company admin, or add_/change_rosterperiod, or the existing add_/change_employeeshiftschedule
    – they work only with employees and rosters of their branches (User branch access);
  * shift master: everybody of the company reads it; add_/change_/delete_shift change it;
  * roster approval: company admin, the approver named on the roster, or approve_rosterperiod (not the submitter);
    publishing: admin, the named approver, publish_rosterperiod or approve_rosterperiod;
  * reporting managers see their team's schedules and approve their team's swap / change / cancellation requests
    and open-shift claims;
  * employees (ESS) see only their own schedule, availability, requests and notices; they can claim open shifts of
    their branch / department and answer swap requests sent to them.
"""
import csv
import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation

from django.db import connection, transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import BasePermission
from rest_framework.response import Response
from rest_framework.views import APIView

from . import notify
from . import services as S
from .models import Availability, OpenShift, OpenShiftClaim, RosterEntry, RosterPeriod, ShiftNotice, ShiftRequest, ShiftRule
from .resolver import Book, ZERO, crosses_midnight, describe, holidays
from .services import PlanError, Rights, person

MAX_PERIOD_DAYS = 62


# ------------------------------------------------------------------ helpers
def _schema(request):
    if connection.schema_name != 'public':
        return connection.schema_name
    return request.GET.get('schema')


class Member(BasePermission):
    """Logged in and a user of the company in ?schema=."""
    message = 'Please log in.'

    def has_permission(self, request, view):
        u = request.user
        if not u or not u.is_authenticated:
            self.message = 'Please log in.'
            return False
        sch = _schema(request)
        if not sch or sch == 'public':
            self.message = 'Choose a company.'
            return False
        if u.is_superuser or u.tenants.filter(schema_name=sch).exists():
            return True
        self.message = 'You do not have access to this company.'
        return False


class Base(APIView):
    permission_classes = [Member]
    parser_classes = [JSONParser, MultiPartParser, FormParser]
    zeo_access = False   # rights are checked in the views (see the module docstring)
    zeo_scope = False

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        self.r = Rights(request)

    def handle_exception(self, exc):
        if isinstance(exc, PlanError):
            body = {'detail': exc.message}
            if isinstance(exc.problems, dict):
                body['errors'] = exc.problems
            elif exc.problems:
                body['problems'] = exc.problems
            return Response(body, status=exc.status)
        return super().handle_exception(exc)


def deny(msg='You may not do this.'):
    raise PlanError(msg, status=403)


def pdate(v, field='date', required=True):
    if v in (None, ''):
        if required:
            raise PlanError(f'Enter the {field.replace("_", " ")} (YYYY-MM-DD).', problems={field: 'Required.'})
        return None
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise PlanError(f'The {field.replace("_", " ")} “{v}” is not a date. Use YYYY-MM-DD.', problems={field: 'Not a date.'})


def pint(v, field, required=False):
    if v in (None, '', 'null'):
        if required:
            raise PlanError(f'Choose the {field.replace("_", " ")}.', problems={field: 'Required.'})
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        raise PlanError(f'The {field.replace("_", " ")} must be a number.', problems={field: 'Not a number.'})


def pbool(v, default=False):
    if v in (None, ''):
        return default
    return str(v).lower() in ('1', 'true', 'yes', 'on', 'y')


def fl(v):
    return None if v is None else float(v)


def hhmm(t):
    return t.strftime('%H:%M') if t else None


def emp_or_404(i):
    e = S.emp_by_id(i)
    if e is None:
        raise PlanError('Employee not found.', status=404)
    return e


def names():
    from django.apps import apps
    M = apps.get_model
    return {'b': dict(M('OrganisationManager', 'brnch_mstr').objects.values_list('id', 'branch_name')),
            'd': dict(M('OrganisationManager', 'dept_master').objects.values_list('id', 'dept_name')),
            'u': dict(M('UserManagement', 'CustomUser').objects.values_list('id', 'username'))}


def csv_response(filename, header, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for r in rows:
        w.writerow(r)
    resp = HttpResponse(buf.getvalue(), content_type='text/csv')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


# ------------------------------------------------------------------ 1. shift master (Shift + ShiftRule)
RULE_NUM = {'grace_in_minutes': 'int', 'grace_out_minutes': 'int', 'ot_after_minutes': 'int', 'min_hours': 'dec', 'max_hours': 'dec',
            'half_day_hours': 'dec', 'shift_allowance': 'dec', 'night_allowance': 'dec'}


def shift_row(s, rule, n=None, usage=None):
    d = describe(s, rule, date.today(), 'master') if s.start_time and s.end_time else None
    row = {'id': s.pk, 'name': s.name, 'start_time': hhmm(s.start_time), 'end_time': hhmm(s.end_time),
           'break_minutes': int(s.break_duration.total_seconds() // 60) if s.break_duration else 0,
           'cross_midnight': crosses_midnight(s), 'hours': fl(d['hours']) if d else 0.0,
           'code': rule.code if rule else '', 'colour': rule.colour if rule else '#5b4ff5', 'active': rule.active if rule else True,
           'grace_in_minutes': rule.grace_in_minutes if rule else 0, 'grace_out_minutes': rule.grace_out_minutes if rule else 0,
           'min_hours': fl(rule.min_hours) if rule else None, 'max_hours': fl(rule.max_hours) if rule else None,
           'half_day_hours': fl(rule.half_day_hours) if rule else None,
           'night_shift': (rule.night_shift if (rule and not rule.night_auto) else crosses_midnight(s)),
           'night_auto': rule.night_auto if rule else True, 'ot_after_minutes': rule.ot_after_minutes if rule else 0,
           'shift_allowance': fl(rule.shift_allowance) if rule else 0.0, 'night_allowance': fl(rule.night_allowance) if rule else 0.0,
           'branch_ids': list(rule.branch_ids or []) if rule else [], 'notes': rule.notes if rule else '', 'has_rules': rule is not None}
    if n is not None:
        row['branches'] = [n['b'].get(int(b), str(b)) for b in row['branch_ids']]
    if usage is not None:
        row['in_use'] = usage
    return row


def _ptime(v, field, errors):
    if v in (None, ''):
        return None
    s = str(v).strip()
    for fmt in ('%H:%M', '%H:%M:%S', '%I:%M %p'):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    errors[field] = 'Use HH:MM, e.g. 09:00.'
    return None


def shift_input(data, shift=None, rule=None, r=None):
    """(shift values, rule values) from the form; raises PlanError with errors per field."""
    from calendars.models import Shift
    errors = {}
    sv, rv = {}, {}
    if shift is None or 'name' in data:
        name = (data.get('name') or '').strip()
        if not name:
            errors['name'] = 'Enter the shift name.'
        elif len(name) > 50:
            errors['name'] = 'Use at most 50 characters.'
        elif Shift.objects.filter(name__iexact=name).exclude(pk=getattr(shift, 'pk', None)).exists():
            errors['name'] = 'A shift with this name already exists.'
        sv['name'] = name
    st = _ptime(data.get('start_time'), 'start_time', errors) if 'start_time' in data or shift is None else shift.start_time
    en = _ptime(data.get('end_time'), 'end_time', errors) if 'end_time' in data or shift is None else shift.end_time
    if (st is None) != (en is None) and not errors.get('start_time') and not errors.get('end_time'):
        errors['end_time' if st else 'start_time'] = 'Enter both start and end time (or neither for a day-off shift).'
    if st and en and st == en:
        errors['end_time'] = 'The end time must differ from the start time.'
    sv['start_time'], sv['end_time'] = st, en
    if 'break_minutes' in data or shift is None:
        try:
            bm = int(data.get('break_minutes') or 0)
            if bm < 0:
                raise ValueError
        except (TypeError, ValueError):
            errors['break_minutes'] = 'Enter the break in whole minutes (0 or more).'
            bm = 0
        sv['break_duration'] = timedelta(minutes=bm)
    brk = sv.get('break_duration', getattr(shift, 'break_duration', None) or timedelta(0))
    length = None
    if st and en:
        a, b = datetime.combine(date.today(), st), datetime.combine(date.today(), en)
        if b <= a:
            b += timedelta(days=1)
        length = Decimal((b - a).total_seconds()) / Decimal(3600)
        if brk and brk >= (b - a) and 'break_minutes' not in errors:
            errors['break_minutes'] = 'The break must be shorter than the shift.'
    # rules
    if 'code' in data or rule is None:
        code = (data.get('code') or '').strip().upper()
        if len(code) > 20:
            errors['code'] = 'Use at most 20 characters.'
        elif code and ShiftRule.objects.filter(code__iexact=code).exclude(shift_id=getattr(shift, 'pk', -1)).exists():
            errors['code'] = 'Another shift already uses this code.'
        rv['code'] = code
    if 'colour' in data:
        c = (data.get('colour') or '').strip()
        if c and not (c.startswith('#') and len(c) in (4, 7) and all(x in '0123456789abcdefABCDEF' for x in c[1:])):
            errors['colour'] = 'Use a colour like #5b4ff5.'
        rv['colour'] = c or '#5b4ff5'
    for f, kind in RULE_NUM.items():
        if f not in data:
            continue
        v = data.get(f)
        if v in (None, ''):
            rv[f] = None if kind == 'dec' and f.endswith('hours') else (ZERO if kind == 'dec' else 0)
            continue
        try:
            val = int(v) if kind == 'int' else Decimal(str(v))
            if val < 0:
                raise ValueError
            if kind == 'dec' and f.endswith('hours') and val > 24:
                errors[f] = 'Use 24 hours or less.'
                continue
            rv[f] = val
        except (TypeError, ValueError, InvalidOperation):
            errors[f] = 'Enter a number of 0 or more.'
    mn = rv.get('min_hours', getattr(rule, 'min_hours', None))
    mx = rv.get('max_hours', getattr(rule, 'max_hours', None))
    hd = rv.get('half_day_hours', getattr(rule, 'half_day_hours', None))
    if mn is not None and mx is not None and mn > mx and 'max_hours' not in errors:
        errors['max_hours'] = 'Maximum hours must be at least the minimum hours.'
    if hd is not None and mn is not None and hd > mn and 'half_day_hours' not in errors:
        errors['half_day_hours'] = 'Half-day hours must not be more than the minimum hours.'
    if 'active' in data:
        rv['active'] = pbool(data.get('active'), True)
    if 'night_auto' in data:
        rv['night_auto'] = pbool(data.get('night_auto'), True)
    if 'night_shift' in data:
        rv['night_shift'] = pbool(data.get('night_shift'))
        if 'night_auto' not in data:
            rv['night_auto'] = False   # set by hand
    if 'notes' in data:
        rv['notes'] = (data.get('notes') or '')[:255]
    if 'branch_ids' in data:
        raw = data.get('branch_ids') or []
        if isinstance(raw, str):
            raw = [x for x in raw.replace(';', ',').split(',') if x.strip()]
        try:
            ids = sorted({int(x) for x in raw})
        except (TypeError, ValueError):
            ids, errors['branch_ids'] = [], 'Choose branches from the list.'
        if r is not None and r.branches is not None:
            if not ids and rule is not None and not rule.branch_ids:
                pass   # a shift of all branches stays one of all branches
            elif not ids:
                errors['branch_ids'] = 'Choose at least one of your branches.'
            elif any(i not in r.branches for i in ids):
                errors['branch_ids'] = 'You can only choose your own branches.'
        rv['branch_ids'] = ids
    if errors:
        raise PlanError('Please correct the marked fields.', problems=errors)
    return sv, rv, length


def _apply_rule(rule, rv, shift):
    for k, v in rv.items():
        setattr(rule, k, v)
    if rule.night_auto:
        rule.night_shift = crosses_midnight(shift)
    rule.save()
    return rule


def _shift_visible(r, rule):
    if r.branches is None or rule is None or not rule.branch_ids:
        return True
    return any(int(b) in r.branches for b in rule.branch_ids)


class ShiftMasterView(Base):
    """GET all shifts with their rules (?active=1 only active); POST a new shift with its rules."""

    def get(self, request):
        from calendars.models import Shift, ShiftPattern
        rules = {x.shift_id: x for x in ShiftRule.objects.all()}
        n = names()
        used = set(RosterEntry.objects.values_list('shift_id', flat=True).distinct())
        rows = []
        for s in Shift.objects.order_by('name'):
            rule = rules.get(s.pk)
            if not _shift_visible(self.r, rule):
                continue
            if (request.GET.get('active') == '1' or not (self.r.manage or self.r.has('view_shift'))) and rule is not None and not rule.active:
                continue
            rows.append(shift_row(s, rule, n, usage=s.pk in used))
        return Response(rows)

    def post(self, request):
        from calendars.models import Shift
        if not self.r.has('add_shift'):
            deny('You may not add shifts.')
        sv, rv, _ = shift_input(request.data, r=self.r)
        if 'branch_ids' not in rv and self.r.branches is not None:
            rv['branch_ids'] = sorted(self.r.branches)   # branch HR: new shifts belong to their branches
        with transaction.atomic():
            s = Shift(**sv)
            s.created_by = request.user
            s.save()
            rule = _apply_rule(ShiftRule(shift_id=s.pk), rv, s)
        return Response(shift_row(s, rule, names()), status=201)


class ShiftMasterDetail(Base):
    def _get(self, pk):
        from calendars.models import Shift
        s = Shift.objects.filter(pk=pk).first()
        if s is None:
            raise PlanError('Shift not found.', status=404)
        rule = ShiftRule.objects.filter(shift_id=s.pk).first()
        if not _shift_visible(self.r, rule):
            raise PlanError('Shift not found.', status=404)
        return s, rule

    def get(self, request, pk):
        s, rule = self._get(pk)
        return Response(shift_row(s, rule, names()))

    def patch(self, request, pk):
        if not self.r.has('change_shift'):
            deny('You may not change shifts.')
        s, rule = self._get(pk)
        if rule is not None and rule.branch_ids and self.r.branches is not None and any(int(b) not in self.r.branches for b in rule.branch_ids):
            deny('This shift is shared with branches you do not manage.')
        sv, rv, _ = shift_input(request.data, shift=s, rule=rule, r=self.r)
        with transaction.atomic():
            for k, v in sv.items():
                setattr(s, k, v)
            s.save()
            rule = _apply_rule(rule or ShiftRule(shift_id=s.pk), rv, s)
        return Response(shift_row(s, rule, names()))

    put = patch

    def delete(self, request, pk):
        from calendars.models import Attendance, EmployeeShiftSchedule, ShiftOverride, ShiftPattern
        if not self.r.has('delete_shift'):
            deny('You may not delete shifts.')
        s, rule = self._get(pk)
        used = []
        if RosterEntry.objects.filter(shift_id=s.pk).exists() or OpenShift.objects.filter(shift_id=s.pk).exists():
            used.append('rosters')
        if ShiftPattern.objects.filter(Q(rotating_shift=s) | Q(monday_shift=s) | Q(tuesday_shift=s) | Q(wednesday_shift=s) | Q(thursday_shift=s)
                                       | Q(friday_shift=s) | Q(saturday_shift=s) | Q(sunday_shift=s)).exists():
            used.append('shift patterns')
        if Attendance.objects.filter(shift=s).exists():
            used.append('attendance')
        if ShiftOverride.objects.filter(override_shift=s).exists():
            used.append('shift overrides')
        if used:
            raise PlanError(f'This shift is used in {", ".join(used)}. Switch it off (Active = No) instead of deleting it.', status=409)
        with transaction.atomic():
            ShiftRule.objects.filter(shift_id=s.pk).delete()
            s.delete()
        return Response(status=204)


# ------------------------------------------------------------------ 2. roster periods
def period_row(p, n=None, r=None, counts=None):
    n = n or names()
    row = {'id': p.id, 'name': p.name, 'branch_id': p.branch_id, 'branch': n['b'].get(p.branch_id, ''),
           'department_id': p.department_id, 'department': n['d'].get(p.department_id, '') if p.department_id else 'All departments',
           'date_from': p.date_from, 'date_to': p.date_to, 'status': p.status, 'status_label': p.get_status_display(),
           'approver_id': p.approver_id, 'approver': n['u'].get(p.approver_id, '') if p.approver_id else '',
           'default_shift_id': p.default_shift_id, 'min_rest_hours': fl(p.min_rest_hours), 'max_week_hours': fl(p.max_week_hours),
           'notes': p.notes, 'decision_note': p.decision_note, 'submitted_at': p.submitted_at, 'approved_at': p.approved_at,
           'approved_by': n['u'].get(p.approved_by_id, ''), 'published_at': p.published_at, 'published_by': n['u'].get(p.published_by_id, ''),
           'created_by': n['u'].get(p.created_by_id, ''), 'created_at': p.created_at}
    if counts is not None:
        row.update(counts)
    if r is not None:
        row['can'] = {'edit': r.can_manage_period(p) and p.status in S.EDITABLE, 'submit': r.can_manage_period(p) and p.status in ('draft', 'rejected'),
                      'approve': p.status == 'submitted' and r.can_approve_period(p), 'publish': p.status == 'approved' and r.can_publish(p),
                      'reopen': p.status in ('submitted', 'approved', 'published') and r.can_manage_period(p),
                      'delete': r.can_manage_period(p) and p.status in ('draft', 'rejected')}
    return row


def period_input(data, p=None):
    from calendars.models import Shift
    errors = {}
    v = {}
    if p is None or 'name' in data:
        v['name'] = (data.get('name') or '').strip()[:120]
        if not v['name']:
            errors['name'] = 'Enter a name, e.g. “Sharjah – October week 2”.'
    if p is None or 'branch' in data or 'branch_id' in data:
        b = data.get('branch', data.get('branch_id'))
        try:
            v['branch_id'] = int(b)
        except (TypeError, ValueError):
            errors['branch'] = 'Choose the branch.'
    if 'department' in data or 'department_id' in data:
        dv = data.get('department', data.get('department_id'))
        v['department_id'] = int(dv) if str(dv or '').isdigit() else None
    for f in ('date_from', 'date_to'):
        if p is None or f in data:
            try:
                v[f] = pdate(data.get(f), f)
            except PlanError as e:
                errors[f] = e.message
    a = v.get('date_from', getattr(p, 'date_from', None))
    b = v.get('date_to', getattr(p, 'date_to', None))
    if a and b and 'date_to' not in errors:
        if b < a:
            errors['date_to'] = 'The end date must be on or after the start date.'
        elif (b - a).days + 1 > MAX_PERIOD_DAYS:
            errors['date_to'] = f'A roster can cover at most {MAX_PERIOD_DAYS} days. Split it into weeks or months.'
    if 'approver' in data or 'approver_id' in data:
        av = data.get('approver', data.get('approver_id'))
        v['approver_id'] = int(av) if str(av or '').isdigit() else None
    if 'default_shift' in data or 'default_shift_id' in data:
        sv = data.get('default_shift', data.get('default_shift_id'))
        v['default_shift_id'] = int(sv) if str(sv or '').isdigit() else None
        if v['default_shift_id'] and not Shift.objects.filter(pk=v['default_shift_id']).exists():
            errors['default_shift'] = 'Choose a shift from the list.'
    for f, lo, hi in (('min_rest_hours', 0, 24), ('max_week_hours', 1, 168)):
        if f in data:
            try:
                x = Decimal(str(data.get(f)))
                if x < lo or x > hi:
                    raise ValueError
                v[f] = x
            except (TypeError, ValueError, InvalidOperation):
                errors[f] = f'Enter a number between {lo} and {hi}.'
    if 'notes' in data:
        v['notes'] = data.get('notes') or ''
    if errors:
        raise PlanError('Please correct the marked fields.', problems=errors)
    return v


def _period_visible(r, p):
    if r.can_manage_period(p) or r.admin or p.approver_id == r.user.pk:
        return True
    if r.has('approve_rosterperiod', 'publish_rosterperiod') and r.branch_ok(p.branch_id):
        return True
    return False


def get_period(r, pk, manage=False):
    p = RosterPeriod.objects.filter(pk=pk).first()
    if p is None or not (_period_visible(r, p) or (not manage and r.team and _team_in_period(r, p))):
        raise PlanError('Roster not found.', status=404)
    if manage and not r.can_manage_period(p):
        deny('You may not change rosters of this branch.')
    return p


def _team_in_period(r, p):
    return any(e.pk in r.team for e in S.period_employees(p))


class PeriodsView(Base):
    def get(self, request):
        qs = RosterPeriod.objects.all()
        st = request.GET.get('status')
        if st:
            qs = qs.filter(status=st)
        n = names()
        rows = []
        for p in qs[:300]:
            if _period_visible(self.r, p) or (self.r.team and p.status == 'published' and _team_in_period(self.r, p)):
                cnt = {'entries': p.entries.count(), 'changed_after_publish': p.entries.filter(changed_after_publish=True).count()}
                rows.append(period_row(p, n, self.r, cnt))
        return Response(rows)

    def post(self, request):
        if not self.r.manage:
            deny('You may not create rosters.')
        v = period_input(request.data)
        if not self.r.branch_ok(v['branch_id']):
            deny('You can only create rosters for your own branches.')
        p = RosterPeriod.objects.create(created_by_id=request.user.pk, **v)
        out = period_row(p, r=self.r)
        if pbool(request.data.get('generate')):
            out['generated'] = S.generate(p, request.user)
        return Response(out, status=201)


class PeriodDetail(Base):
    def get(self, request, pk):
        p = get_period(self.r, pk)
        return Response(period_row(p, r=self.r, counts={'entries': p.entries.count()}))

    def patch(self, request, pk):
        p = get_period(self.r, pk, manage=True)
        S.check_editable(p)
        v = period_input(request.data, p)
        if 'branch_id' in v and v['branch_id'] != p.branch_id:
            if p.entries.exists():
                raise PlanError('The branch cannot change once the roster has shifts. Create a new roster instead.')
            if not self.r.branch_ok(v['branch_id']):
                deny('You can only use your own branches.')
        if ('date_from' in v or 'date_to' in v) and p.entries.exclude(date__gte=v.get('date_from', p.date_from), date__lte=v.get('date_to', p.date_to)).exists():
            raise PlanError('Some planned days fall outside the new dates. Clear them first.')
        for k, x in v.items():
            setattr(p, k, x)
        p.save()
        return Response(period_row(p, r=self.r))

    def delete(self, request, pk):
        p = get_period(self.r, pk, manage=True)
        if p.status not in ('draft', 'rejected'):
            raise PlanError('Only draft rosters can be deleted. Send it back to draft first.', status=409)
        p.delete()
        return Response(status=204)


class PeriodAction(Base):
    """POST periods/<id>/<action>/ : generate, check, submit, recall, approve, reject, publish, reopen, copy-week."""

    def post(self, request, pk, action):
        p = get_period(self.r, pk)
        d = request.data
        u = request.user
        now = timezone.now()
        if action == 'generate':
            if not self.r.can_manage_period(p):
                deny('You may not change rosters of this branch.')
            res = S.generate(p, u, overwrite=pbool(d.get('overwrite')))
            return Response({'result': res, 'problems': S.check_period(p)})
        if action == 'check':
            return Response({'problems': S.check_period(p)})
        if action == 'copy-week':
            if not self.r.can_manage_period(p):
                deny('You may not change rosters of this branch.')
            t = pdate(d.get('target_start'), 'target_start')
            s = pdate(d.get('source_start'), 'source_start', required=False)
            n = S.copy_week(p, t, s, u)
            return Response({'copied': n, 'problems': S.check_period(p)})
        if action == 'submit':
            if not self.r.can_manage_period(p):
                deny('You may not submit rosters of this branch.')
            if p.status not in ('draft', 'rejected'):
                raise PlanError(f'This roster is {p.get_status_display().lower()} already.', status=409)
            if not p.entries.exists():
                raise PlanError('The roster is empty. Generate it or plan some shifts first.')
            probs = S.check_period(p)
            errs = [x for x in probs if x['level'] == 'error']
            if errs:
                raise PlanError(f'Fix {len(errs)} problem(s) before submitting.', problems=errs)
            p.status, p.submitted_by_id, p.submitted_at, p.decision_note = 'submitted', u.pk, now, ''
            p.save()
            self._tell_approvers(p)
            return Response(period_row(p, r=self.r) | {'problems': probs})
        if action == 'recall':
            if p.status != 'submitted' or not self.r.can_manage_period(p):
                raise PlanError('Only a submitted roster can be called back by its planners.', status=409)
            p.status = 'draft'
            p.save()
            return Response(period_row(p, r=self.r))
        if action in ('approve', 'reject'):
            if p.status != 'submitted':
                raise PlanError('Only a submitted roster can be approved or sent back.', status=409)
            if not self.r.can_approve_period(p):
                deny('You may not approve this roster.')
            note = (d.get('note') or '').strip()
            if action == 'reject':
                if not note:
                    raise PlanError('Say why the roster is sent back, so the planner can fix it.', problems={'note': 'Required.'})
                p.status, p.decision_note = 'rejected', note[:255]
                p.save()
                self._tell_user(p.submitted_by_id, 'Roster sent back', f'The roster “{p.name}” was sent back: {note}')
                return Response(period_row(p, r=self.r))
            p.status, p.approved_by_id, p.approved_at, p.decision_note = 'approved', u.pk, now, note[:255]
            p.save()
            self._tell_user(p.submitted_by_id, 'Roster approved', f'The roster “{p.name}” was approved and can be published.')
            if pbool(d.get('publish')) and self.r.can_publish(p):
                return Response(self._publish(p, u, now))
            return Response(period_row(p, r=self.r))
        if action == 'publish':
            if p.status != 'approved':
                raise PlanError('Approve the roster before publishing it.', status=409)
            if not self.r.can_publish(p):
                deny('You may not publish this roster.')
            return Response(self._publish(p, u, now))
        if action == 'reopen':
            if not self.r.can_manage_period(p) or p.status not in ('submitted', 'approved', 'published'):
                raise PlanError('Only submitted, approved or published rosters can be reopened.', status=409)
            was = p.status
            p.status = 'draft'
            p.save()
            if was == 'published':
                for emp in S.E().objects.filter(pk__in=p.entries.values('employee_id')):
                    notify.send(employee=emp, kind='unpublished', title='Schedule withdrawn',
                                message=f'Your schedule {p.date_from:%d %b} – {p.date_to:%d %b %Y} was taken back for changes. You will be told when it is published again.')
            return Response(period_row(p, r=self.r))
        raise PlanError('Unknown action.', status=404)

    def _publish(self, p, u, now):
        probs = S.check_period(p)
        errs = [x for x in probs if x['level'] == 'error']
        if errs:
            raise PlanError(f'Fix {len(errs)} problem(s) before publishing.', problems=errs)
        p.status, p.published_by_id, p.published_at = 'published', u.pk, now
        p.save()
        p.entries.update(changed_after_publish=False)
        shifts = S.shifts_by_id()
        sent = 0
        by_emp = {}
        for e in p.entries.order_by('date'):
            by_emp.setdefault(e.employee_id, []).append(e)
        for emp in S.E().objects.filter(pk__in=list(by_emp)):
            work = [e for e in by_emp[emp.pk] if not e.off and e.shift_id]
            lines = '; '.join(f'{e.date:%a %d %b} {S.shift_label(shifts.get(e.shift_id))}' for e in work[:14])
            notify.send(employee=emp, kind='published', title='Your schedule is published',
                        message=f'Your schedule {p.date_from:%d %b} – {p.date_to:%d %b %Y}: {len(work)} shift day(s). {lines}')
            sent += 1
        return period_row(p) | {'notified': sent, 'problems': probs}

    def _tell_user(self, uid, title, msg):
        from UserManagement.models import CustomUser
        u = CustomUser.objects.filter(pk=uid).first() if uid else None
        if u is not None:
            notify.send(user=u, kind='roster', title=title, message=msg)

    def _tell_approvers(self, p):
        if p.approver_id:
            self._tell_user(p.approver_id, 'Roster waiting for approval', f'The roster “{p.name}” ({p.date_from:%d %b} – {p.date_to:%d %b}) waits for your approval.')


# ------------------------------------------------------------------ roster grid + cells
def grid_payload(p, r, a=None, b=None):
    a = max(a or p.date_from, p.date_from)
    b = min(b or p.date_to, p.date_to)
    emps = S.period_employees(p)
    if not r.can_manage_period(p) and not r.can_approve_period(p) and not r.can_publish(p):
        emps = [e for e in emps if e.pk in r.team]   # reporting manager: own team only
    ids = [e.pk for e in emps]
    shifts = S.shifts_by_id()
    rules = {x.shift_id: x for x in ShiftRule.objects.all()}
    entries = {(x.employee_id, x.date): x for x in RosterEntry.objects.filter(period=p, date__gte=a, date__lte=b, employee_id__in=ids)}
    leaves = S.leave_map(ids, a, b)
    avail = S.availability_index(ids)
    probs = {}
    for x in S.check_period(p, set(ids)):
        probs.setdefault((x['employee_id'], x['date']), []).append({'kind': x['kind'], 'level': x['level'], 'message': x['message']})
    days = list(S.days(a, b))
    hol_cache = {}
    rows = []
    for e in emps:
        hk = (e.emp_branch_id_id, e.emp_dept_id_id)
        if hk not in hol_cache:
            hol_cache[hk] = holidays(e, a, b)
        cells = {}
        hours = ZERO
        for d in days:
            x = entries.get((e.pk, d))
            sh = shifts.get(x.shift_id) if x is not None and x.shift_id and not x.off else None
            info = describe(sh, rules.get(sh.pk) if sh else None, d, 'roster') if sh else None
            if info:
                hours += info['hours']
            lv = leaves.get((e.pk, d))
            cells[d.isoformat()] = {
                'entry_id': x.pk if x else None, 'shift_id': sh.pk if sh else None, 'off': bool(x and x.off), 'empty': x is None,
                'code': (rules.get(sh.pk).code if sh and rules.get(sh.pk) else '') or (sh.name[:3].upper() if sh else ('OFF' if x and x.off else '')),
                'label': S.shift_label(sh, rules) if sh else ('Day off' if x and x.off else ''),
                'colour': (rules.get(sh.pk).colour if sh and rules.get(sh.pk) else '#5b4ff5') if sh else '',
                'start': hhmm(sh.start_time) if sh else None, 'end': hhmm(sh.end_time) if sh else None,
                'night': bool(info and info['night_shift']), 'source': x.source if x else '', 'note': x.note if x else '',
                'changed': bool(x and x.changed_after_publish), 'leave': lv, 'holiday': d in hol_cache[hk],
                'availability': S.availability_flag(avail.get(e.pk, []), d, sh.pk if sh else None) if avail.get(e.pk) else '',
                'problems': probs.get((e.pk, d), []),
            }
        week_probs = [y for (eid, dd), lst in probs.items() if eid == e.pk and dd not in [dx for dx in days] for y in lst]
        rows.append({'employee_id': e.pk, 'code': e.emp_code, 'name': person(e), 'department': getattr(e.emp_dept_id, 'dept_name', '') if e.emp_dept_id_id else '',
                     'designation': getattr(e.emp_desgntn_id, 'desgntn_job_title', '') if e.emp_desgntn_id_id else '',
                     'hours': fl(hours), 'cells': cells, 'notes': week_probs})
    return {'period': period_row(p, r=r), 'from': a, 'to': b,
            'days': [{'date': d, 'weekday': d.strftime('%a'), 'day': d.day} for d in days],
            'shifts': [shift_row(s, rules.get(s.pk)) for s in shifts.values() if not (rules.get(s.pk) and not rules[s.pk].active)],
            'rows': rows}


class GridView(Base):
    def get(self, request, pk):
        p = get_period(self.r, pk)
        a = pdate(request.GET.get('from'), 'from', required=False)
        b = pdate(request.GET.get('to'), 'to', required=False)
        return Response(grid_payload(p, self.r, a, b))


class CellsView(Base):
    """POST {cells: [{employee, date, shift (id or null), off (bool), note}]} – set roster cells (empty shift + off false clears)."""

    def post(self, request, pk):
        p = get_period(self.r, pk, manage=True)
        S.check_editable(p)
        cells = request.data.get('cells')
        if not isinstance(cells, list) or not cells:
            raise PlanError('Send the cells to change.')
        scope = {e.pk: e for e in S.period_employees(p)}
        shifts = S.shifts_by_id()
        bad = []
        with transaction.atomic():
            for i, c in enumerate(cells):
                try:
                    emp = scope.get(pint(c.get('employee'), 'employee', True))
                    if emp is None:
                        raise PlanError('This employee is not part of the roster (branch / department).')
                    d = pdate(c.get('date'), 'date')
                    sid = pint(c.get('shift'), 'shift')
                    if sid is not None and sid not in shifts:
                        raise PlanError('Choose a shift from the list.')
                    S.upsert_entry(p, emp, d, sid, pbool(c.get('off')), 'manual', request.user, note=c.get('note'))
                except PlanError as e:
                    bad.append({'row': i + 1, 'employee': c.get('employee'), 'date': c.get('date'), 'message': e.message})
            if bad:
                transaction.set_rollback(True)
        if bad:
            raise PlanError('Nothing was saved. Fix the cells listed.', problems=bad)
        ids = {pint(c.get('employee'), 'employee') for c in cells}
        return Response({'saved': len(cells), 'problems': S.check_period(p, ids)})


# ------------------------------------------------------------------ 3. availability
def avail_row(a, n=None):
    return {'id': a.id, 'employee_id': a.employee_id, 'employee': person(S.emp_by_id(a.employee_id)), 'kind': a.kind, 'kind_label': a.get_kind_display(),
            'weekday': a.weekday, 'weekday_label': ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday'][a.weekday] if a.weekday is not None else '',
            'date_from': a.date_from, 'date_to': a.date_to, 'shift_id': a.shift_id, 'note': a.note}


def _target_emp(r, data_emp, what='this'):
    """The employee a request is for: own by default; HR / manager may name one of theirs."""
    if data_emp in (None, ''):
        if r.emp is None:
            raise PlanError('Your login is not linked to an employee. Choose the employee.', problems={'employee': 'Required.'})
        return r.emp
    emp = emp_or_404(data_emp)
    if r.emp is not None and emp.pk == r.emp.pk:
        return emp
    if r.manages_emp(emp) or r.is_manager_of(emp):
        return emp
    deny(f'You can only do {what} for yourself.')


def avail_input(data, errors=None):
    from calendars.models import Shift
    errors = {} if errors is None else errors
    v = {}
    kind = data.get('kind') or 'unavailable'
    if kind not in dict(Availability.KINDS):
        errors['kind'] = 'Choose available, unavailable or preferred.'
    v['kind'] = kind
    wd = data.get('weekday')
    v['weekday'] = None
    if wd not in (None, ''):
        try:
            v['weekday'] = int(wd)
            if not 0 <= v['weekday'] <= 6:
                raise ValueError
        except (TypeError, ValueError):
            errors['weekday'] = 'Choose a weekday.'
    try:
        v['date_from'] = pdate(data.get('date_from'), 'date_from', required=False)
        v['date_to'] = pdate(data.get('date_to'), 'date_to', required=False)
    except PlanError as e:
        errors['date_from'] = e.message
    if v.get('weekday') is None and not v.get('date_from') and 'weekday' not in errors:
        errors['weekday'] = 'Choose a weekday or a date range.'
    if v.get('date_from') and v.get('date_to') and v['date_to'] < v['date_from']:
        errors['date_to'] = 'The end date must be on or after the start date.'
    sid = data.get('shift')
    v['shift_id'] = int(sid) if str(sid or '').isdigit() else None
    if v['shift_id'] and not Shift.objects.filter(pk=v['shift_id']).exists():
        errors['shift'] = 'Choose a shift from the list.'
    v['note'] = (data.get('note') or '')[:255]
    if errors:
        raise PlanError('Please correct the marked fields.', problems=errors)
    return v


class AvailabilityView(Base):
    def get(self, request):
        qs = Availability.objects.all()
        eid = request.GET.get('employee')
        if eid:
            emp = emp_or_404(eid)
            if not self.r.can_see_emp(emp):
                deny('You can only see your own availability.')
            qs = qs.filter(employee_id=emp.pk)
        else:
            allowed = set(self.r.team) | ({self.r.emp.pk} if self.r.emp else set())
            if self.r.manage:
                emps = S.E().objects.all() if self.r.branches is None else S.E().objects.filter(emp_branch_id__in=self.r.branches)
                allowed |= set(emps.values_list('pk', flat=True))
            qs = qs.filter(employee_id__in=allowed)
        return Response([avail_row(a) for a in qs[:1000]])

    def post(self, request):
        emp = _target_emp(self.r, request.data.get('employee'), 'availability')
        v = avail_input(request.data)
        a = Availability.objects.create(employee_id=emp.pk, created_by_id=request.user.pk, **v)
        return Response(avail_row(a), status=201)


class AvailabilityDetail(Base):
    def _get(self, pk):
        a = Availability.objects.filter(pk=pk).first()
        if a is None:
            raise PlanError('Not found.', status=404)
        emp = S.emp_by_id(a.employee_id)
        if not ((self.r.emp and self.r.emp.pk == a.employee_id) or self.r.manages_emp(emp) or self.r.is_manager_of(emp)):
            raise PlanError('Not found.', status=404)
        return a

    def patch(self, request, pk):
        a = self._get(pk)
        merged = {**avail_row(a), 'shift': a.shift_id, **request.data.dict()} if hasattr(request.data, 'dict') else {**avail_row(a), 'shift': a.shift_id, **request.data}
        v = avail_input(merged)
        for k, x in v.items():
            setattr(a, k, x)
        a.save()
        return Response(avail_row(a))

    def delete(self, request, pk):
        self._get(pk).delete()
        return Response(status=204)


# ------------------------------------------------------------------ 4. open shifts
def open_row(o, r, shifts, n, mine=None):
    claims = list(o.claims.all())
    approved = sum(1 for c in claims if c.status == 'approved')
    row = {'id': o.id, 'date': o.date, 'shift_id': o.shift_id, 'shift': S.shift_label(shifts.get(o.shift_id)),
           'branch_id': o.branch_id, 'branch': n['b'].get(o.branch_id, ''), 'department_id': o.department_id,
           'department': n['d'].get(o.department_id, '') if o.department_id else 'Any department', 'slots': o.slots,
           'filled': approved, 'left': max(o.slots - approved, 0), 'note': o.note, 'status': o.status, 'status_label': o.get_status_display()}
    can_all = r.manage and r.branch_ok(o.branch_id)
    visible = [c for c in claims if can_all or (r.emp and c.employee_id == r.emp.pk) or c.employee_id in r.team]
    row['claims'] = [{'id': c.id, 'employee_id': c.employee_id, 'employee': person(S.emp_by_id(c.employee_id)), 'status': c.status,
                      'status_label': c.get_status_display(), 'note': c.note, 'decision_note': c.decision_note, 'created_at': c.created_at,
                      'can_decide': c.status == 'pending' and r.can_decide_for(S.emp_by_id(c.employee_id))} for c in visible]
    if r.emp:
        m = next((c for c in claims if c.employee_id == r.emp.pk and c.status in ('pending', 'approved')), None)
        row['my_claim'] = {'id': m.id, 'status': m.status, 'status_label': m.get_status_display()} if m else None
        row['can_claim'] = (o.status == 'open' and m is None and r.emp.emp_branch_id_id == o.branch_id
                            and (not o.department_id or o.department_id == r.emp.emp_dept_id_id))
    row['can_manage'] = can_all
    return row


def _eligible(o, emp):
    return emp.emp_branch_id_id == o.branch_id and (not o.department_id or o.department_id == emp.emp_dept_id_id)


class OpenShiftsView(Base):
    def get(self, request):
        qs = OpenShift.objects.prefetch_related('claims')
        st = request.GET.get('status')
        if st:
            qs = qs.filter(status=st)
        if request.GET.get('from'):
            qs = qs.filter(date__gte=pdate(request.GET.get('from'), 'from'))
        shifts, n = S.shifts_by_id(), names()
        rows = []
        for o in qs[:500]:
            if (self.r.manage and self.r.branch_ok(o.branch_id)) or (self.r.emp and _eligible(o, self.r.emp)) \
                    or any(c.employee_id in self.r.team for c in o.claims.all()):
                rows.append(open_row(o, self.r, shifts, n))
        return Response(rows)

    def post(self, request):
        from calendars.models import Shift
        d = request.data
        errors = {}
        try:
            day = pdate(d.get('date'), 'date')
            if day < date.today():
                errors['date'] = 'Choose today or a later day.'
        except PlanError as e:
            errors['date'], day = e.message, None
        sid = d.get('shift')
        if not str(sid or '').isdigit() or not Shift.objects.filter(pk=sid).exists():
            errors['shift'] = 'Choose a shift from the list.'
        bid = d.get('branch')
        if not str(bid or '').isdigit():
            errors['branch'] = 'Choose the branch.'
        try:
            slots = int(d.get('slots')) if d.get('slots') not in (None, '') else 1
            if slots < 1 or slots > 100:
                raise ValueError
        except (TypeError, ValueError):
            errors['slots'], slots = 'Enter how many people are needed (1–100).', 1
        if errors:
            raise PlanError('Please correct the marked fields.', problems=errors)
        if not self.r.manage or not self.r.branch_ok(int(bid)):
            deny('You may not post open shifts for this branch.')
        dep = int(d['department']) if str(d.get('department') or '').isdigit() else None
        o = OpenShift.objects.create(date=day, shift_id=int(sid), branch_id=int(bid), department_id=dep, slots=slots,
                                     note=(d.get('note') or '')[:255], created_by_id=request.user.pk)
        sh = S.shifts_by_id().get(o.shift_id)
        told = 0
        if pbool(d.get('notify'), True):
            qs = S.E().objects.filter(emp_branch_id_id=o.branch_id).exclude(is_active=False)
            if dep:
                qs = qs.filter(emp_dept_id_id=dep)
            for emp in qs[:300]:
                if emp.users_id:
                    notify.send(employee=emp, kind='open_shift', title='Open shift available',
                                message=f'An open shift {S.shift_label(sh)} on {o.date:%a %d %b %Y} needs {o.slots} person(s). Claim it in My schedule.')
                    told += 1
        return Response(open_row(o, self.r, S.shifts_by_id(), names()) | {'notified': told}, status=201)


class OpenShiftAction(Base):
    """POST open-shifts/<id>/claim/ (employee) or /cancel/ (HR)."""

    def post(self, request, pk, action):
        o = OpenShift.objects.filter(pk=pk).first()
        if o is None:
            raise PlanError('Open shift not found.', status=404)
        if action == 'cancel':
            if not (self.r.manage and self.r.branch_ok(o.branch_id)):
                deny('You may not cancel this open shift.')
            if o.status != 'open':
                raise PlanError(f'This open shift is {o.get_status_display().lower()}.', status=409)
            o.status = 'cancelled'
            o.save()
            for c in o.claims.filter(status='pending'):
                S.stamp(c, request.user, 'rejected', 'The open shift was cancelled.')
                c.save()
                notify.send(employee=S.emp_by_id(c.employee_id), kind='open_shift', title='Open shift cancelled',
                            message=f'The open shift on {o.date:%d %b %Y} you claimed was cancelled.')
            return Response(open_row(o, self.r, S.shifts_by_id(), names()))
        if action != 'claim':
            raise PlanError('Unknown action.', status=404)
        emp = self.r.emp
        if emp is None:
            raise PlanError('Your login is not linked to an employee.', status=403)
        if o.status != 'open':
            raise PlanError('This open shift is no longer open.', status=409)
        if not _eligible(o, emp):
            deny('This open shift is for another branch or department.')
        if o.claims.filter(employee_id=emp.pk, status__in=('pending', 'approved')).exists():
            raise PlanError('You have already claimed this shift.', status=409)
        S.refuse_if_leave(emp, o.date)
        sid, off = S.current_shift(emp, o.date)
        if sid and not off:
            raise PlanError(f'You already work {S.shift_label(S.shifts_by_id().get(sid))} on {o.date:%d %b %Y}. Ask for a shift change instead.', status=409)
        c = OpenShiftClaim.objects.create(open_shift=o, employee_id=emp.pk, note=(request.data.get('note') or '')[:255])
        mgr = emp.emp_reporting_manager
        if mgr is not None:
            notify.send(user=mgr, kind='open_shift', title='Open shift claimed',
                        message=f'{person(emp)} claimed the open shift on {o.date:%a %d %b %Y}. Approve or reject it in Shift planner → Requests.')
        return Response({'id': c.id, 'status': c.status, 'status_label': c.get_status_display()}, status=201)


class ClaimAction(Base):
    """POST claims/<id>/<approve|reject|withdraw>/"""

    def post(self, request, pk, action):
        c = OpenShiftClaim.objects.select_related('open_shift').filter(pk=pk).first()
        if c is None:
            raise PlanError('Claim not found.', status=404)
        emp = S.emp_by_id(c.employee_id)
        o = c.open_shift
        if c.status != 'pending':
            raise PlanError(f'This claim is {c.get_status_display().lower()} already.', status=409)
        if action == 'withdraw':
            if not (self.r.emp and self.r.emp.pk == c.employee_id):
                deny('Only the employee can withdraw the claim.')
            c.status = 'withdrawn'
            c.save()
            return Response({'status': c.status})
        if action not in ('approve', 'reject'):
            raise PlanError('Unknown action.', status=404)
        if not self.r.can_decide_for(emp):
            deny('You may not decide on this claim.')
        note = (request.data.get('note') or '')[:255]
        if action == 'reject':
            S.stamp(c, request.user, 'rejected', note)
            c.save()
            notify.send(employee=emp, kind='open_shift', title='Open shift claim rejected',
                        message=f'Your claim for the open shift on {o.date:%d %b %Y} was rejected.' + (f' {note}' if note else ''))
            return Response({'status': c.status})
        if o.status != 'open':
            raise PlanError('This open shift is no longer open.', status=409)
        S.refuse_if_leave(emp, o.date)
        with transaction.atomic():
            S.set_shift(emp, o.date, o.shift_id, 'open_shift', request.user, note='Open shift')
            S.stamp(c, request.user, 'approved', note)
            c.save()
            if o.claims.filter(status='approved').count() >= o.slots:
                o.status = 'filled'
                o.save()
                for other in o.claims.filter(status='pending'):
                    S.stamp(other, request.user, 'rejected', 'All places were filled.')
                    other.save()
                    notify.send(employee=S.emp_by_id(other.employee_id), kind='open_shift', title='Open shift filled',
                                message=f'The open shift on {o.date:%d %b %Y} was filled by others.')
        sh = S.shifts_by_id().get(o.shift_id)
        notify.send(employee=emp, kind='open_shift', title='Open shift approved',
                    message=f'You work {S.shift_label(sh)} on {o.date:%a %d %b %Y}.')
        return Response({'status': c.status, 'open_shift_status': o.status})


# ------------------------------------------------------------------ 5/6. swap, change and cancellation requests
def req_row(q, r, shifts, n):
    emp, other = S.emp_by_id(q.employee_id), S.emp_by_id(q.swap_employee_id) if q.swap_employee_id else None
    row = {'id': q.id, 'kind': q.kind, 'kind_label': q.get_kind_display(), 'employee_id': q.employee_id, 'employee': person(emp),
           'date': q.date, 'from_shift_id': q.from_shift_id, 'from_shift': S.shift_label(shifts.get(q.from_shift_id)) if q.from_shift_id else 'Off',
           'to_shift_id': q.to_shift_id, 'to_shift': S.shift_label(shifts.get(q.to_shift_id)) if q.to_shift_id else '',
           'swap_employee_id': q.swap_employee_id, 'swap_employee': person(other), 'swap_date': q.swap_date,
           'swap_shift': (S.shift_label(shifts.get(q.swap_shift_id)) if q.swap_shift_id else 'Off') if q.kind == 'swap' else '',
           'reason': q.reason, 'status': q.status, 'status_label': q.get_status_display(), 'decision_note': q.decision_note,
           'decided_by': n['u'].get(q.decided_by_id, ''), 'decided_at': q.decided_at, 'created_at': q.created_at}
    me = r.emp.pk if r.emp else None
    row['can'] = {'accept': q.status == 'peer' and me is not None and me == q.swap_employee_id,
                  'approve': q.status == 'pending' and r.can_decide_for(emp),
                  'withdraw': q.status in ('peer', 'pending') and (me == q.employee_id or q.created_by_id == r.user.pk)}
    return row


def _visible_req(r, q):
    me = r.emp.pk if r.emp else None
    if me in (q.employee_id, q.swap_employee_id) or q.created_by_id == r.user.pk:
        return True
    emp = S.emp_by_id(q.employee_id)
    return r.manages_emp(emp) or r.is_manager_of(emp)


class RequestsView(Base):
    def get(self, request):
        qs = ShiftRequest.objects.all()
        scope = request.GET.get('scope') or 'all'
        me = self.r.emp.pk if self.r.emp else -1
        if scope == 'mine':
            qs = qs.filter(Q(employee_id=me) | Q(swap_employee_id=me))
        elif scope == 'team':
            qs = qs.filter(employee_id__in=list(self.r.team))
        st = request.GET.get('status')
        if st:
            qs = qs.filter(status__in=st.split(','))
        if request.GET.get('kind'):
            qs = qs.filter(kind=request.GET['kind'])
        shifts, n = S.shifts_by_id(), names()
        return Response([req_row(q, self.r, shifts, n) for q in qs[:500] if _visible_req(self.r, q)])

    def post(self, request):
        d = request.data
        kind = d.get('kind')
        if kind not in dict(ShiftRequest.KINDS):
            raise PlanError('Choose swap, change or cancel.', problems={'kind': 'Required.'})
        emp = _target_emp(self.r, d.get('employee'), 'shift requests')
        day = pdate(d.get('date'), 'date')
        if day < date.today():
            raise PlanError('You can only change today or later days.', problems={'date': 'In the past.'})
        if ShiftRequest.objects.filter(employee_id=emp.pk, date=day, status__in=('peer', 'pending')).exists():
            raise PlanError('There is already an open request for this day. Withdraw it first.', status=409)
        cur_id, cur_off = S.current_shift(emp, day)
        shifts = S.shifts_by_id()
        q = ShiftRequest(kind=kind, employee_id=emp.pk, date=day, from_shift_id=None if cur_off else cur_id,
                         reason=(d.get('reason') or '')[:255], created_by_id=request.user.pk)
        if kind == 'change':
            to = pint(d.get('to_shift'), 'to_shift', True)
            if to not in shifts:
                raise PlanError('Choose the shift you want.', problems={'to_shift': 'Choose a shift.'})
            rule = ShiftRule.objects.filter(shift_id=to).first()
            if rule is not None and not rule.active:
                raise PlanError('That shift is switched off. Choose another one.', problems={'to_shift': 'Inactive.'})
            if q.from_shift_id == to:
                raise PlanError('You already work this shift on that day.', problems={'to_shift': 'Same shift.'})
            S.refuse_if_leave(emp, day)
            q.to_shift_id, q.status = to, 'pending'
        elif kind == 'cancel':
            if q.from_shift_id is None:
                raise PlanError('There is no shift to cancel on that day.', status=409)
            if not q.reason:
                raise PlanError('Say why the shift should be cancelled.', problems={'reason': 'Required.'})
            q.status = 'pending'
        else:  # swap
            other = emp_or_404(pint(d.get('swap_employee'), 'swap_employee', True))
            if other.pk == emp.pk:
                raise PlanError('Choose a colleague, not yourself.', problems={'swap_employee': 'Same employee.'})
            if other.emp_branch_id_id != emp.emp_branch_id_id:
                raise PlanError('You can only swap with a colleague of your branch.', problems={'swap_employee': 'Other branch.'})
            sday = pdate(d.get('swap_date'), 'swap_date', required=False) or day
            if sday < date.today():
                raise PlanError("The colleague's day is in the past.", problems={'swap_date': 'In the past.'})
            for x, dd in ((emp, day), (other, sday), (emp, sday), (other, day)):
                S.refuse_if_leave(x, dd)
            o_id, o_off = S.current_shift(other, sday)
            q.swap_employee_id, q.swap_date, q.swap_shift_id = other.pk, sday, None if o_off else o_id
            if sday == day and q.swap_shift_id == q.from_shift_id:
                raise PlanError('You both work the same shift that day – there is nothing to swap.', status=409)
            if sday == day and q.swap_shift_id is None and q.from_shift_id is None:
                raise PlanError('You are both off that day – there is nothing to swap.', status=409)
            q.status = 'peer'
        q.save()
        shift_txt = S.shift_label(shifts.get(q.from_shift_id)) if q.from_shift_id else 'Off'
        if q.kind == 'swap':
            notify.send(employee=S.emp_by_id(q.swap_employee_id), kind='swap', title='Shift swap request',
                        message=f'{person(emp)} asks to swap: their {day:%a %d %b} ({shift_txt}) for your {q.swap_date:%a %d %b} '
                                f'({S.shift_label(shifts.get(q.swap_shift_id)) if q.swap_shift_id else "Off"}). Accept or decline in My schedule.')
        else:
            self._tell_approver(emp, q)
        return Response(req_row(q, self.r, shifts, names()), status=201)

    @staticmethod
    def _tell_approver(emp, q):
        mgr = emp.emp_reporting_manager
        if mgr is not None:
            notify.send(user=mgr, kind='request', title=f'{q.get_kind_display()} waiting for approval',
                        message=f'{person(emp)} asks for a {q.get_kind_display().lower()} on {q.date:%a %d %b %Y}. Approve or reject it in Shift planner → Requests.')


def apply_request(q, user):
    emp = S.emp_by_id(q.employee_id)
    if q.kind == 'change':
        S.refuse_if_leave(emp, q.date)
        S.set_shift(emp, q.date, q.to_shift_id, 'change', user, note=q.reason)
    elif q.kind == 'cancel':
        S.set_shift(emp, q.date, None, 'cancel', user, note=q.reason or 'Shift cancelled')
    else:
        other = S.emp_by_id(q.swap_employee_id)
        dates = sorted({q.date, q.swap_date or q.date})
        for x in (emp, other):
            S.refuse_if_leave(x, *dates)
        now = {(x.pk, dd): S.current_shift(x, dd) for x in (emp, other) for dd in dates}
        for dd in dates:
            a_id, a_off = now[(emp.pk, dd)]
            b_id, b_off = now[(other.pk, dd)]
            S.set_shift(emp, dd, None if b_off else b_id, 'swap', user, note=f'Swap with {other.emp_code}')
            S.set_shift(other, dd, None if a_off else a_id, 'swap', user, note=f'Swap with {emp.emp_code}')


class RequestAction(Base):
    """POST requests/<id>/<accept|decline|approve|reject|withdraw>/"""

    def post(self, request, pk, action):
        q = ShiftRequest.objects.filter(pk=pk).first()
        if q is None or not _visible_req(self.r, q):
            raise PlanError('Request not found.', status=404)
        note = (request.data.get('note') or '')[:255]
        emp = S.emp_by_id(q.employee_id)
        shifts, n = S.shifts_by_id(), names()
        me = self.r.emp.pk if self.r.emp else None
        if action in ('accept', 'decline'):
            if q.status != 'peer' or me != q.swap_employee_id:
                raise PlanError('Only the colleague can answer this swap while it waits for them.', status=403)
            q.peer_decided_at = timezone.now()
            if action == 'decline':
                q.status, q.decision_note = 'declined', note
                q.save()
                notify.send(employee=emp, kind='swap', title='Swap declined', message=f'{person(self.r.emp)} declined your swap for {q.date:%d %b %Y}.' + (f' {note}' if note else ''))
            else:
                for x, dd in ((emp, q.date), (self.r.emp, q.swap_date or q.date)):
                    S.refuse_if_leave(x, dd)
                q.status = 'pending'
                q.save()
                notify.send(employee=emp, kind='swap', title='Swap accepted', message=f'{person(self.r.emp)} accepted your swap for {q.date:%d %b %Y}. It now waits for approval.')
                RequestsView._tell_approver(emp, q)
            return Response(req_row(q, self.r, shifts, n))
        if action == 'withdraw':
            if q.status not in ('peer', 'pending') or not (me == q.employee_id or q.created_by_id == request.user.pk):
                raise PlanError('Only the requester can withdraw an open request.', status=403)
            q.status = 'withdrawn'
            q.save()
            if q.swap_employee_id:
                notify.send(employee=S.emp_by_id(q.swap_employee_id), kind='swap', title='Swap withdrawn', message=f'{person(emp)} withdrew the swap for {q.date:%d %b %Y}.')
            return Response(req_row(q, self.r, shifts, n))
        if action not in ('approve', 'reject'):
            raise PlanError('Unknown action.', status=404)
        if q.status != 'pending':
            raise PlanError(f'This request is {q.get_status_display().lower()}.', status=409)
        if not self.r.can_decide_for(emp):
            deny('You may not decide on this request.')
        if action == 'reject':
            if not note:
                raise PlanError('Say why the request is rejected.', problems={'note': 'Required.'})
            S.stamp(q, request.user, 'rejected', note)
            q.save()
        else:
            with transaction.atomic():
                apply_request(q, request.user)
                S.stamp(q, request.user, 'approved', note)
                q.save()
        word = 'approved' if action == 'approve' else 'rejected'
        for x in filter(None, [emp, S.emp_by_id(q.swap_employee_id) if q.swap_employee_id else None]):
            notify.send(employee=x, kind='request', title=f'{q.get_kind_display()} {word}',
                        message=f'The {q.get_kind_display().lower()} for {q.date:%a %d %b %Y} was {word}.' + (f' {note}' if note else ''))
        return Response(req_row(q, self.r, shifts, n))


# ------------------------------------------------------------------ 9. My schedule (ESS / mobile) and team schedule
def day_rows(emp, a, b, book=None, shifts=None):
    book = book or Book(a, b, [emp.pk])
    lv = S.leave_map([emp.pk], a, b)
    hol = holidays(emp, a, b)
    reqs = {}
    for q in ShiftRequest.objects.filter(Q(employee_id=emp.pk) | Q(swap_employee_id=emp.pk), status__in=('peer', 'pending'), date__gte=a, date__lte=b):
        reqs.setdefault(q.date, []).append({'id': q.id, 'kind': q.get_kind_display(), 'status': q.get_status_display()})
    out = []
    for d in S.days(a, b):
        x = book.get(emp, d)
        out.append({'date': d, 'weekday': d.strftime('%A'), 'shift_id': x['shift_id'] if x else None,
                    'shift': (x['name'] if x and not x['off'] else ('Day off' if x else 'Not planned')),
                    'code': x['code'] if x else '', 'colour': x['colour'] if x else '',
                    'start': x['start'].strftime('%H:%M') if x and x['start'] else None,
                    'end': x['end'].strftime('%H:%M') if x and x['end'] else None,
                    'next_day': bool(x and x['end'] and x['end'].date() > d), 'break_minutes': x['break_minutes'] if x else 0,
                    'hours': fl(x['hours']) if x else 0.0, 'off': bool(x and x['off']), 'planned': x is not None,
                    'night': bool(x and x['night_shift']), 'source': x['source'] if x else '',
                    'leave': lv.get((emp.pk, d)), 'holiday': d in hol, 'requests': reqs.get(d, [])})
    return out


def _range(request, default_days=14):
    a = pdate(request.GET.get('from'), 'from', required=False) or date.today()
    b = pdate(request.GET.get('to'), 'to', required=False) or (a + timedelta(days=default_days - 1))
    if b < a:
        raise PlanError('The end date must be on or after the start date.')
    if (b - a).days > 92:
        raise PlanError('Choose at most 3 months.')
    return a, b


class MyScheduleView(Base):
    """GET my-schedule/?from&to – the logged-in employee's published shifts with leave and holidays (default 14 days)."""

    def get(self, request):
        emp = self.r.emp
        if emp is None:
            raise PlanError('Your login is not linked to an employee.', status=403)
        a, b = _range(request)
        shifts, n = S.shifts_by_id(), names()
        opens = [open_row(o, self.r, shifts, n) for o in OpenShift.objects.filter(status='open', date__gte=date.today(), branch_id=emp.emp_branch_id_id)
                 .prefetch_related('claims') if _eligible(o, emp)]
        mine = ShiftRequest.objects.filter(Q(employee_id=emp.pk) | Q(swap_employee_id=emp.pk)).order_by('-created_at')[:30]
        return Response({'employee': {'id': emp.pk, 'code': emp.emp_code, 'name': person(emp)}, 'from': a, 'to': b,
                         'days': day_rows(emp, a, b, shifts=shifts),
                         'open_shifts': opens,
                         'requests': [req_row(q, self.r, shifts, n) for q in mine],
                         'availability': [avail_row(x) for x in Availability.objects.filter(employee_id=emp.pk)],
                         'notices': [{'id': x.id, 'kind': x.kind, 'title': x.title, 'message': x.message, 'read': x.read, 'created_at': x.created_at}
                                     for x in ShiftNotice.objects.filter(Q(employee_id=emp.pk) | Q(user_id=request.user.pk))[:20]],
                         'shifts': [{'id': s.pk, 'name': S.shift_label(s)} for s in shifts.values()
                                    if not ShiftRule.objects.filter(shift_id=s.pk, active=False).exists()]})


class TeamScheduleView(Base):
    """GET team-schedule/?from&to[&branch] – HR: employees of their branches; managers: their team; employees: colleagues of the
    same department (for picking a swap partner)."""

    def get(self, request):
        a, b = _range(request, 7)
        if self.r.manage:
            qs = S.E().objects.exclude(is_active=False)
            if self.r.branches is not None:
                qs = qs.filter(emp_branch_id__in=self.r.branches)
            if request.GET.get('branch'):
                qs = qs.filter(emp_branch_id=pint(request.GET.get('branch'), 'branch'))
            if request.GET.get('department'):
                qs = qs.filter(emp_dept_id=pint(request.GET.get('department'), 'department'))
            emps, full = list(qs[:300]), True
        elif self.r.team:
            emps, full = list(S.E().objects.filter(pk__in=self.r.team)), True
        elif self.r.emp is not None:
            emps = list(S.E().objects.filter(emp_branch_id=self.r.emp.emp_branch_id_id, emp_dept_id=self.r.emp.emp_dept_id_id).exclude(is_active=False).exclude(pk=self.r.emp.pk))
            full = False
        else:
            emps, full = [], False
        book = Book(a, b, [e.pk for e in emps])
        rows = []
        for e in emps:
            ds = day_rows(e, a, b, book)
            if not full:  # colleagues: only the shift, not their leave
                ds = [{k: v for k, v in x.items() if k in ('date', 'weekday', 'shift_id', 'shift', 'start', 'end', 'off', 'planned')} for x in ds]
            rows.append({'employee_id': e.pk, 'code': e.emp_code, 'name': person(e), 'days': ds})
        return Response({'from': a, 'to': b, 'rows': rows})


class NoticesView(Base):
    def get(self, request):
        me = self.r.emp.pk if self.r.emp else -1
        qs = ShiftNotice.objects.filter(Q(user_id=request.user.pk) | Q(employee_id=me))
        if request.GET.get('unread') == '1':
            qs = qs.filter(read=False)
        return Response([{'id': x.id, 'kind': x.kind, 'title': x.title, 'message': x.message, 'read': x.read, 'emailed': x.emailed,
                          'created_at': x.created_at} for x in qs[:100]])

    def post(self, request):
        me = self.r.emp.pk if self.r.emp else -1
        qs = ShiftNotice.objects.filter(Q(user_id=request.user.pk) | Q(employee_id=me))
        ids = request.data.get('ids')
        if ids:
            qs = qs.filter(pk__in=ids)
        return Response({'marked': qs.update(read=True)})


class ResolveView(Base):
    """GET resolve/?employee&date[&to] – the shift(s) attendance and payroll will use (shift resolver output)."""

    def get(self, request):
        emp = emp_or_404(request.GET.get('employee') or (self.r.emp.pk if self.r.emp else None))
        if not self.r.can_see_emp(emp):
            deny('You can only look at your own shifts.')
        a = pdate(request.GET.get('date'), 'date')
        b = pdate(request.GET.get('to'), 'to', required=False) or a
        if (b - a).days > 62:
            raise PlanError('Choose at most 62 days.')
        book = Book(a, b, [emp.pk])
        out = []
        for d in S.days(a, b):
            x = book.get(emp, d)
            if x is not None:
                x = {k: (fl(v) if isinstance(v, Decimal) else v) for k, v in x.items() if k != 'shift'}
            out.append({'date': d, 'result': x})
        from .resolver import scheduled_days
        return Response({'employee': person(emp), 'days': out, 'scheduled_days': scheduled_days(emp, a, b, book)})


class PayrollPreview(Base):
    """GET payroll/?employee&from&to – the shift variables payroll formulas get."""

    def get(self, request):
        emp = emp_or_404(request.GET.get('employee'))
        if not (self.r.manages_emp(emp) or (self.r.has('view_payslip', 'add_payrollrun', 'view_payrollrun') and self.r.branch_ok(emp.emp_branch_id_id))):
            deny('You may not see payroll figures of this employee.')
        from .payroll import variables
        a = pdate(request.GET.get('from'), 'from')
        b = pdate(request.GET.get('to'), 'to')
        return Response({k: fl(v) for k, v in variables(emp, a, b).items()})


# ------------------------------------------------------------------ 10. bulk upload
ROSTER_COLS = ['employee_code', 'date', 'shift', 'note']
SHIFT_COLS = ['name', 'code', 'start_time', 'end_time', 'break_minutes', 'grace_in_minutes', 'grace_out_minutes', 'min_hours', 'max_hours',
              'half_day_hours', 'night_shift', 'ot_after_minutes', 'shift_allowance', 'night_allowance', 'colour', 'active']


def read_table(f):
    """Rows (dicts with lower-case headers) of an uploaded CSV or XLSX file."""
    name = (getattr(f, 'name', '') or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h or '').strip().lower().replace(' ', '_') for h in next(it, [])]
        rows = []
        for vals in it:
            if not any(v not in (None, '') for v in vals):
                continue
            row = {}
            for h, v in zip(head, vals):
                if isinstance(v, datetime):
                    v = v.date().isoformat() if h == 'date' else v.strftime('%H:%M')
                elif isinstance(v, date):
                    v = v.isoformat()
                elif isinstance(v, time):
                    v = v.strftime('%H:%M')
                row[h] = '' if v is None else str(v).strip()
            rows.append(row)
        return rows
    if name and not name.endswith('.csv') and not name.endswith('.txt'):
        raise PlanError('Upload a .csv or .xlsx file.')
    raw = f.read()
    text = raw.decode('utf-8-sig') if isinstance(raw, bytes) else raw
    rd = csv.DictReader(io.StringIO(text))
    rd.fieldnames = [(h or '').strip().lower().replace(' ', '_') for h in (rd.fieldnames or [])]
    return [{k: (v or '').strip() for k, v in r.items() if k} for r in rd if any((v or '').strip() for v in r.values() if isinstance(v, str))]


def shift_lookup():
    from calendars.models import Shift
    out = {}
    rules = {x.shift_id: x for x in ShiftRule.objects.all()}
    for s in Shift.objects.all():
        out[s.name.strip().lower()] = s.pk
        r = rules.get(s.pk)
        if r is not None and r.code:
            out[r.code.strip().lower()] = s.pk
    return out


class UploadRosterView(Base):
    """POST upload/roster/ (multipart: file, period, commit=0|1). Preview first; commit saves only when every row is valid."""

    def post(self, request):
        p = get_period(self.r, pint(request.data.get('period'), 'period', True), manage=True)
        S.check_editable(p)
        f = request.FILES.get('file')
        if f is None:
            raise PlanError('Choose the file to upload.', problems={'file': 'Required.'})
        rows = read_table(f)
        if not rows:
            raise PlanError('The file has no rows. Use the template.')
        if 'employee_code' not in rows[0] or 'date' not in rows[0] or 'shift' not in rows[0]:
            raise PlanError('The file needs the columns employee_code, date and shift. Download the template.')
        scope = {e.emp_code.lower(): e for e in S.period_employees(p)}
        lookup = shift_lookup()
        out, seen = [], set()
        for i, row in enumerate(rows, start=2):
            errs = []
            emp = scope.get((row.get('employee_code') or '').lower())
            if emp is None:
                errs.append('Employee code not found in this roster’s branch / department.')
            try:
                d = pdate(row.get('date'), 'date')
                if d < p.date_from or d > p.date_to:
                    errs.append(f'Date outside the roster ({p.date_from} – {p.date_to}).')
            except PlanError as e:
                d = None
                errs.append(e.message)
            sv = (row.get('shift') or '').strip()
            off = sv.lower() in ('off', 'day off', 'o', 'rest')
            sid = None if off else lookup.get(sv.lower())
            if not off and sid is None:
                errs.append(f'Shift “{sv}” not found. Use a shift code, a shift name or OFF.')
            key = (row.get('employee_code', '').lower(), d)
            if key in seen and d:
                errs.append('The same employee and date appear twice.')
            seen.add(key)
            out.append({'row': i, 'employee_code': row.get('employee_code'), 'employee': person(emp) if emp else '', 'date': d,
                        'shift': 'OFF' if off else sv, 'shift_id': sid, 'off': off, 'note': row.get('note', ''), 'errors': errs, 'ok': not errs})
        bad = [x for x in out if not x['ok']]
        if not pbool(request.data.get('commit')):
            return Response({'rows': out, 'valid': len(out) - len(bad), 'errors': len(bad)})
        if bad:
            raise PlanError(f'{len(bad)} row(s) have errors. Nothing was saved.', problems=bad)
        with transaction.atomic():
            for x in out:
                S.upsert_entry(p, scope[x['employee_code'].lower()], x['date'], x['shift_id'], x['off'], 'upload', request.user, note=x['note'])
        return Response({'saved': len(out), 'problems': S.check_period(p)})


class UploadShiftsView(Base):
    """POST upload/shifts/ (multipart: file, commit=0|1) – add or update shifts with their rules by name."""

    def post(self, request):
        from calendars.models import Shift
        if not self.r.has('add_shift') or not self.r.has('change_shift'):
            deny('You may not add or change shifts.')
        f = request.FILES.get('file')
        if f is None:
            raise PlanError('Choose the file to upload.', problems={'file': 'Required.'})
        rows = read_table(f)
        if not rows or 'name' not in rows[0]:
            raise PlanError('The file needs at least the column name. Download the template.')
        out, plans = [], []
        names_seen = set()
        for i, row in enumerate(rows, start=2):
            name = (row.get('name') or '').strip()
            s = Shift.objects.filter(name__iexact=name).first() if name else None
            rule = ShiftRule.objects.filter(shift_id=s.pk).first() if s else None
            data = {k: v for k, v in row.items() if k in SHIFT_COLS and (v != '' or k in ('start_time', 'end_time'))}
            if s is not None:
                data.pop('name', None)
            try:
                if name.lower() in names_seen:
                    raise PlanError('The same shift name appears twice.')
                names_seen.add(name.lower())
                sv, rv, _ = shift_input(data, shift=s, rule=rule, r=self.r)
                out.append({'row': i, 'name': name, 'action': 'update' if s else 'add', 'errors': [], 'ok': True})
                plans.append((s, rule, sv, rv))
            except PlanError as e:
                msgs = [f'{k.replace("_", " ")}: {v}' for k, v in e.problems.items()] if isinstance(e.problems, dict) and e.problems else [e.message]
                out.append({'row': i, 'name': name, 'action': 'update' if s else 'add', 'errors': msgs, 'ok': False})
        bad = [x for x in out if not x['ok']]
        if not pbool(request.data.get('commit')):
            return Response({'rows': out, 'valid': len(out) - len(bad), 'errors': len(bad)})
        if bad:
            raise PlanError(f'{len(bad)} row(s) have errors. Nothing was saved.', problems=bad)
        with transaction.atomic():
            for s, rule, sv, rv in plans:
                if s is None:
                    s = Shift(created_by=request.user)
                for k, v in sv.items():
                    setattr(s, k, v)
                s.save()
                _apply_rule(rule or ShiftRule(shift_id=s.pk), rv, s)
        return Response({'saved': len(plans)})


class TemplateView(Base):
    def get(self, request, kind):
        if kind == 'roster':
            ex = S.E().objects.exclude(is_active=False).order_by('emp_code').first()
            code = next((r.code for r in ShiftRule.objects.exclude(code='')[:1]), 'GEN')
            return csv_response('roster_upload_template.csv', ROSTER_COLS,
                                [[ex.emp_code if ex else 'EMP0001', date.today().isoformat(), code, ''],
                                 [ex.emp_code if ex else 'EMP0001', (date.today() + timedelta(days=1)).isoformat(), 'OFF', 'Weekly off']])
        if kind == 'shifts':
            return csv_response('shift_upload_template.csv', SHIFT_COLS,
                                [['Morning 07:00-15:00', 'MOR', '07:00', '15:00', '60', '10', '10', '8', '10', '4', 'no', '30', '0', '0', '#2e7d32', 'yes'],
                                 ['Night 22:00-06:00', 'NGT', '22:00', '06:00', '30', '15', '15', '7', '9', '4', 'yes', '30', '10', '25', '#283593', 'yes']])
        raise PlanError('Unknown template.', status=404)


# ------------------------------------------------------------------ 12. reports
def roster_report_rows(r, a, b, branch=None, department=None, status='published'):
    qs = RosterEntry.objects.filter(date__gte=a, date__lte=b).select_related('period')
    if status:
        qs = qs.filter(period__status=status)
    if branch:
        qs = qs.filter(period__branch_id=branch)
    if department:
        qs = qs.filter(period__department_id=department)
    emps = {e.pk: e for e in S.E().objects.filter(pk__in=qs.values('employee_id')).select_related('emp_branch_id', 'emp_dept_id')}
    shifts = S.shifts_by_id()
    rules = {x.shift_id: x for x in ShiftRule.objects.all()}
    rows = []
    for x in qs.order_by('date', 'employee_id'):
        e = emps.get(x.employee_id)
        if e is None or not (r.manages_emp(e) or r.is_manager_of(e) or r.admin):
            continue
        sh = shifts.get(x.shift_id) if x.shift_id and not x.off else None
        info = describe(sh, rules.get(sh.pk) if sh else None, x.date, 'roster') if sh else None
        rows.append({'employee_code': e.emp_code, 'employee': person(e), 'branch': getattr(e.emp_branch_id, 'branch_name', ''),
                     'department': getattr(e.emp_dept_id, 'dept_name', '') if e.emp_dept_id_id else '', 'date': x.date,
                     'shift': sh.name if sh else 'Off', 'start': hhmm(sh.start_time) if sh else '', 'end': hhmm(sh.end_time) if sh else '',
                     'hours': fl(info['hours']) if info else 0.0, 'night': bool(info and info['night_shift']), 'roster': x.period.name,
                     'roster_status': x.period.get_status_display(), 'source': x.get_source_display(), 'changed_after_publish': x.changed_after_publish})
    return rows


class RosterReport(Base):
    def get(self, request):
        if not (self.r.manage or self.r.team):
            deny('This report is for HR and managers.')
        a, b = _range(request, 31)
        rows = roster_report_rows(self.r, a, b, pint(request.GET.get('branch'), 'branch'), pint(request.GET.get('department'), 'department'),
                                  request.GET.get('status', 'published'))
        if request.GET.get('export') == 'csv':
            cols = list(rows[0].keys()) if rows else ['employee_code', 'employee', 'date', 'shift']
            return csv_response('roster_report.csv', cols, [[x[c] for c in cols] for x in rows])
        return Response(rows)


class RequestsReport(Base):
    def get(self, request):
        if not (self.r.manage or self.r.team):
            deny('This report is for HR and managers.')
        a, b = _range(request, 31)
        qs = ShiftRequest.objects.filter(date__gte=a, date__lte=b)
        if request.GET.get('status'):
            qs = qs.filter(status=request.GET['status'])
        shifts, n = S.shifts_by_id(), names()
        rows = []
        for q in qs.order_by('date'):
            e = S.emp_by_id(q.employee_id)
            if not (self.r.manages_emp(e) or self.r.is_manager_of(e)):
                continue
            x = req_row(q, self.r, shifts, n)
            x.pop('can', None)
            rows.append(x)
        for c in OpenShiftClaim.objects.filter(open_shift__date__gte=a, open_shift__date__lte=b).select_related('open_shift'):
            e = S.emp_by_id(c.employee_id)
            if not (self.r.manages_emp(e) or self.r.is_manager_of(e)):
                continue
            rows.append({'id': c.id, 'kind': 'open_shift', 'kind_label': 'Open shift claim', 'employee_id': c.employee_id, 'employee': person(e),
                         'date': c.open_shift.date, 'from_shift': '', 'to_shift': S.shift_label(shifts.get(c.open_shift.shift_id)),
                         'swap_employee': '', 'reason': c.note, 'status': c.status, 'status_label': c.get_status_display(),
                         'decision_note': c.decision_note, 'decided_by': n['u'].get(c.decided_by_id, ''), 'decided_at': c.decided_at, 'created_at': c.created_at})
        if request.GET.get('export') == 'csv':
            cols = ['kind_label', 'employee', 'date', 'from_shift', 'to_shift', 'swap_employee', 'reason', 'status_label', 'decision_note', 'decided_by', 'decided_at', 'created_at']
            return csv_response('shift_requests_report.csv', cols, [[x.get(c, '') for c in cols] for x in rows])
        return Response(rows)


class MetaView(Base):
    """GET meta/ – what the user may do, plus lookups for the screens."""

    def get(self, request):
        from django.apps import apps
        M = apps.get_model
        br = M('OrganisationManager', 'brnch_mstr').objects.all()
        if self.r.branches is not None:
            br = br.filter(pk__in=self.r.branches) if self.r.manage else br.filter(pk__in=[self.r.emp.emp_branch_id_id] if self.r.emp else [])
        return Response({
            'rights': {'manage': self.r.manage, 'admin': self.r.admin, 'manager': bool(self.r.team), 'employee': self.r.emp is not None,
                       'add_shift': self.r.has('add_shift'), 'change_shift': self.r.has('change_shift'), 'delete_shift': self.r.has('delete_shift'),
                       'approve': self.r.has('approve_rosterperiod'), 'publish': self.r.has('publish_rosterperiod', 'approve_rosterperiod')},
            'me': {'id': self.r.emp.pk, 'name': person(self.r.emp), 'branch_id': self.r.emp.emp_branch_id_id} if self.r.emp else None,
            'branches': list(br.values('id', 'branch_name')),
            'departments': list(M('OrganisationManager', 'dept_master').objects.values('id', 'dept_name')),
            'approvers': list(M('UserManagement', 'CustomUser').objects.filter(tenants__schema_name=_schema(request)).values('id', 'username').distinct()[:300]),
        })
