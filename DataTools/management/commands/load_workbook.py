"""v1.62.0 – load a data workbook (Excel) into a company: test data, demo data or a migration.

    python manage.py load_workbook ZEO_Demo_Data.xlsx --schema gdis [--user gdis_admin] [--only "Employees,Leave requests"] [--check]

The workbook has an "Index" sheet that names, in order, the sheets to load and how:

    Order | Sheet | App | Screen | Endpoint | How | Phase | Run as | Hand import | Notes  [| Method]

  How     import  – rows go through the generic import of the list screens (/tools/api/import/): the same API, rights, branch
                    rules, duplicate checks and workflows as the screen; linked columns take the code or the name. Such a sheet
                    can also be imported by hand from its screen (Import button) – its column titles are the screen's labels.
          api     – each row is posted to the screen's own API (the "New" form of the app screens), one request per row.
                    (Index column Method: PATCH / PUT for a sheet that changes one record, e.g. /users/api/company/{company}/.)
          action  – each row is an API call: columns @method, @endpoint, @body (JSON) – approvals, posting, payroll runs …
          job     – each row runs one of the daily jobs the scheduler runs every morning (column job: maintenance-daily), so the
                    records it makes (preventive work orders of the plans that are due) exist before the next sheets.
  Phase   setup   – loaded first, in the order of the Index.
          monthly – loaded month by month: every monthly sheet's rows of a month (column @month = YYYY-MM), in Index order,
                    before the next month – the way the company would have worked.

Special columns (start with @; every other column is a field of the screen – its label or its key):
  @ref      a name for the new record; later cells refer to it as {ref:NAME} (its id).
  @skip_if  a lookup – when it finds a record the row is skipped (the record keeps the @ref): makes a second run safe.
  @as       the user (username or e-mail) who does it – the employee applying, the manager approving …
  @month    YYYY-MM for monthly sheets.
  @date     the business day of the row: the record's created / submitted time is moved to that day (history looks real);
            on an action row the times the action sets on its record (approved, resolved, issued …) move to that day.
  @method, @endpoint, @body, @expect – action rows.
  @until (YYYY-MM-DD) – the row is for a future day (the screen refuses past days): after that day it is left out and counted as "date passed".
  @example  v1.62.0 workbook: any value (e.g. "example") – the row is an example of the template and is never loaded.
Placeholders in any cell:  {ref:NAME}  {q:App.Model?field=value&field2=value}  (id of the newest match)  {user:username}  {today}  {today:+N} (N days after the load day)
                           {company} (id of the company being loaded)   {schema}
Cells that start with [ or { are JSON (lists of lines, several values).

v1.62.0 workbook – values as people type them: dates dd/mm/yyyy, d-m-yyyy, 5 Mar 2026, yyyy-mm-dd or Excel dates; numbers with
thousands commas or "AED"; yes / no; linked columns by code or name. Column titles may be the screen labels (with or without the
"Columns" sheet). Unknown columns (with the closest title), missing mandatory columns, sheets missing from the Index, the same row
twice, and the reason of every refused row (with the column and how to fix it) are reported.

--check   a full trial: every row is loaded inside one database transaction that is rolled back at the end – nothing is saved, no
          e-mail leaves, and a row that needs a record of an earlier sheet (employee of the Employees sheet) is checked properly.
--quick-check  the old check: import rows validated only, other rows not tried (seconds).
Every failed row is listed (sheet, row, column, reason, fix) and written to <workbook>.log.json; nothing stops the rest of the load.
The same engine runs the Settings › Data workbook screen and its API (DataTools/workbook_jobs.py).
"""
import datetime
import difflib
import hashlib
import json
import re
import time
import traceback
from decimal import Decimal  # noqa: F401

from django.core.management.base import BaseCommand, CommandError

SPECIAL = ('@ref', '@skip_if', '@until', '@as', '@month', '@date', '@method', '@endpoint', '@body', '@expect', '@note', '@example')
TS_FIELDS = ('created_at', 'created', 'created_on', 'updated_at', 'reported_at', 'requested_at', 'opened_at', 'submitted_at',
             'applied_on', 'request_date', 'raised_at', 'logged_at')
# v1.62.0 workbook: sheets of a workbook that are never loaded (help, lists)
SYSTEM_SHEETS = ('read me', 'readme', 'index', 'columns', 'logins', 'columns help', 'lists', 'summary', 'mapping')
MONTHS = {m: i for i, m in enumerate(('jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'), start=1)}


class Stopped(Exception):
    """v1.62.0 workbook: the load was stopped from the screen / API."""


class _Undo(Exception):
    def __init__(self, st, res):
        super().__init__(st)
        self.st, self.res = st, res


# ---------------------------------------------------------------------- v1.62.0 workbook: values as people type them
def norm_date(v):
    """dd/mm/yyyy, d-m-yyyy, d.m.yy, 5 Mar 2026, 5-Mar-26, Mar 5 2026, yyyy-mm-dd, yyyy/mm/dd, Excel serial → 'YYYY-MM-DD' (None = not a date)."""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
    s = str(v or '').strip()
    if not s:
        return None
    m = re.fullmatch(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T]00:00(?::00)?)?', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.fullmatch(r'(\d{1,2})[-/. ](\d{1,2})[-/. ](\d{2}|\d{4})', s)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if mo > 12 >= d:   # 03/25/2026 – the month first (US): turned round
                d, mo = mo, d
        else:
            m = re.fullmatch(r'(\d{1,2})(?:st|nd|rd|th)?[-/. ]+([A-Za-z]{3,9})[-/., ]+(\d{2}|\d{4})', s) or None
            if m:
                d, mo, y = int(m.group(1)), MONTHS.get(m.group(2)[:3].lower()), int(m.group(3))
            else:
                m = re.fullmatch(r'([A-Za-z]{3,9})[-/. ]+(\d{1,2})(?:st|nd|rd|th)?[-/., ]+(\d{2}|\d{4})', s)
                if not m:
                    return None
                mo, d, y = MONTHS.get(m.group(1)[:3].lower()), int(m.group(2)), int(m.group(3))
            if not mo:
                return None
        if y < 100:
            y += 2000 if y < 70 else 1900
    try:
        return datetime.date(y, mo, d).isoformat()
    except (ValueError, TypeError):
        return None


def norm_number(v):
    """'12,500.00', 'AED 1,200', '1 200', '(300)' → a number (None = not a number)."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        return v
    s = str(v or '').strip()
    if not s:
        return None
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()')
    s = re.sub(r'^(AED|Dhs?\.?|SAR|QAR|OMR|KWD|BHD|INR|USD|Rs\.?|\$)\s*', '', s, flags=re.I)
    s = re.sub(r'\s*(AED|Dhs?\.?|SAR|QAR|OMR|KWD|BHD|INR|USD)$', '', s, flags=re.I)
    if not re.fullmatch(r'-?(\d{1,3}([,  ]\d{3})+|\d+)(\.\d+)?', s):
        return None
    s = re.sub(r'[,  ]', '', s)
    n = float(s) if '.' in s else int(s)
    return -n if neg else n


def closest(name, options, cutoff=0.68):
    opts = [o for o in options if o]
    low = {str(o).lower(): o for o in opts}
    hit = difflib.get_close_matches(str(name or '').lower(), list(low), n=1, cutoff=cutoff)
    return low[hit[0]] if hit else None


def col_letter(i):
    s, i = '', i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def clean_title(h):
    """'Employee code *' / 'Employee code (required)' → 'Employee code'."""
    s = str(h or '').strip()
    s = re.sub(r'\s*\((required|mandatory|optional)\)\s*$', '', s, flags=re.I)
    return s.rstrip('* ').strip()


class Loader:
    def __init__(self, cmd, schema, user, check=False, verbose=1, trial=False, memory=None, on_row=None, should_stop=None, may_act_as=None):
        from django.conf import settings
        from rest_framework.test import APIClient
        from UserManagement.models import CustomUser
        self.cmd, self.schema, self.check, self.verbose = cmd, schema, check, verbose
        if '*' not in settings.ALLOWED_HOSTS:
            settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['localhost', 'testserver']
        self.U = CustomUser
        self.admin = self._user(user) if user else self._admin()
        self.APIClient = APIClient
        self.clients = {}
        self.refs = {}
        self.colmap = {}
        self.fields = {}
        self.log = []
        self.stats = {}
        # v1.62.0 workbook
        self.trial = trial                 # --check: everything really runs inside a transaction that is rolled back
        self.memory = memory               # rows loaded by earlier runs (DataTools.WorkbookRow) – a re-run skips them
        self.on_row = on_row               # progress callback(sheet)
        self.should_stop = should_stop     # () → True when the load was stopped
        self.may_act_as = may_act_as       # (user) → None or the reason the workbook may not act as that user
        self.block = 500
        self.digests = {}                  # (sheet, excel row) → digest of the row
        self.rowvals = {}                  # (sheet, excel row) → {key: raw value}
        self.headers = {}                  # sheet → (titles, keys)
        self.sheet_fields = {}             # sheet → {key: field description} (screen fields, when known)
        self.dups = {}                     # (sheet, excel row) → the earlier row with the same values
        self._fieldcache = {}
        self._company_id = None
        self.failed_vals = {}              # value of a failed row → (sheet, row): "not found" later says why

    # ------------------------------------------------------------------ users and calls
    def _admin(self):
        from UserManagement.models import company
        t = company.objects.filter(schema_name=self.schema).first()
        if t is None:
            raise CommandError(f'No company with the schema "{self.schema}".')
        u = self.U.objects.filter(username=f'{self.schema}_admin').first() or self.U.objects.filter(tenants=t, is_superuser=True).order_by('pk').first() \
            or self.U.objects.filter(is_superuser=True).order_by('pk').first()
        if u is None:
            raise CommandError('No administrator found – use --user.')
        return u

    def _user(self, name):
        name = str(name).strip()
        u = self.U.objects.filter(username=name).first() or self.U.objects.filter(email__iexact=name).first()
        if u is None:
            raise ValueError(f'user "{name}" not found')
        return u

    def client(self, user):
        c = self.clients.get(user.pk)
        if c is None:
            c = self.APIClient(HTTP_HOST='localhost')
            c.force_authenticate(user)
            self.clients[user.pk] = c
        return c

    def url(self, p):
        p = '/' + p.lstrip('/')
        return p + ('&' if '?' in p else '?') + 'schema=' + self.schema

    def call(self, user, method, path, data=None):
        if not self.trial:
            return self._call(user, method, path, data)
        # v1.62.0 workbook --check: the call really runs, in a savepoint of the trial transaction; a refused call is undone,
        # and what the screens do "after saving" (on_commit) runs at once so the next rows see it
        from django.db import transaction
        try:
            with transaction.atomic():
                st, res = self._call(user, method, path, data)
                if st >= 400:
                    raise _Undo(st, res)
        except _Undo as u:
            st, res = u.st, u.res
        self.drain()
        return st, res

    def _call(self, user, method, path, data=None):
        c = self.client(user)
        fn = getattr(c, method.lower())
        r = fn(self.url(path), data=data, format='json') if data is not None else fn(self.url(path))
        body = None
        ct = r.get('Content-Type', '') if hasattr(r, 'get') else ''
        if 'json' in ct:
            try:
                body = r.json()
            except Exception:
                body = None
        return r.status_code, body

    def drain(self):
        """v1.62.0 workbook trial: run the on_commit callbacks now (they never run in a transaction that is rolled back)."""
        from django.db import connection, transaction
        for _ in range(5):
            cbs = list(getattr(connection, 'run_on_commit', []) or [])
            if not cbs:
                return
            connection.run_on_commit = []
            for item in cbs:
                func = item[1] if isinstance(item, (list, tuple)) else item
                try:
                    with transaction.atomic():
                        func()
                except Exception:   # noqa: BLE001 – a side effect (mail, webhook) of a trial row
                    pass

    def stop_check(self):
        if self.should_stop and self.should_stop():
            raise Stopped()

    # ------------------------------------------------------------------ placeholders
    def resolve(self, v):
        if not isinstance(v, str) or '{' not in v:
            return v
        # innermost placeholders first, so {q:…{ref:X}…} and {q:…} inside longer text both resolve
        pat = re.compile(r'\{(ref|q|user|today|company|schema)(?::([^{}]*))?\}')
        s = v.strip()
        while True:
            ms = list(pat.finditer(s))
            if not ms:
                return s
            if len(ms) == 1 and ms[0].span() == (0, len(s)):
                return self._ph(ms[0].group(1), ms[0].group(2) or '')
            s = pat.sub(lambda m: str(self._ph(m.group(1), m.group(2) or '')), s)

    def _ph(self, kind, arg):
        if kind == 'today':   # v1.62.0 {today:+3} / {today:-1}: days from the load day (requests that must be for today or later)
            return (datetime.date.today() + datetime.timedelta(days=int(arg or 0))).isoformat()
        if kind == 'company':   # v1.62.0 workbook: the company being loaded (company details sheet)
            if self._company_id is None:
                from UserManagement.models import company
                self._company_id = company.objects.filter(schema_name=self.schema).values_list('pk', flat=True).first()
            return self._company_id
        if kind == 'schema':
            return self.schema
        if kind == 'ref':
            if arg not in self.refs:
                raise ValueError(f'unknown reference {{ref:{arg}}}')
            return self.refs[arg]
        if kind == 'user':
            return self._user(arg).pk
        if kind == 'q':
            return self.lookup(arg, required=True)
        return ''

    def lookup(self, spec, required=False):
        from django.apps import apps
        from django_tenants.utils import schema_context
        spec = self.resolve(spec) if '{' in spec else spec
        model, _, query = spec.partition('?')
        app, _, name = model.partition('.')
        M = apps.get_model(app, name)
        flt = {}
        for part in filter(None, query.split('&')):
            k, _, val = part.partition('=')
            val = self.resolve(val)
            if isinstance(val, str) and val.lower() in ('true', 'false'):
                val = val.lower() == 'true'
            elif k.endswith('__in') and isinstance(val, str):
                val = [x.strip() for x in val.split(',') if x.strip()]
            elif isinstance(val, str) and val.lower() == 'null':
                k, val = k + '__isnull', True
            flt[k] = val
        if M._meta.label == 'auth.Group' and len(flt) == 1 and next(iter(flt)) in ('name', 'name__iexact'):
            # v1.62.0 workbook: a permission group by name = the group THIS company uses for that name (its own copy of a
            # standard group, its own group) – group names are server-wide, another company's group would be refused
            from ZeoApps import groupscope
            with schema_context(self.schema):
                g = groupscope.resolve(next(iter(flt.values())))
            pk = g.pk if g is not None else None
        else:
            with schema_context(self.schema):
                pk = M.objects.filter(**flt).order_by('-pk').values_list('pk', flat=True).first()
        if pk is None and required:
            raise ValueError(f'nothing found for {{q:{spec}}}')
        return pk

    # ------------------------------------------------------------------ cells
    def cell(self, v):
        if v is None:
            return None
        if isinstance(v, datetime.datetime):
            return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
        if isinstance(v, datetime.date):
            return v.isoformat()
        if isinstance(v, float) and v.is_integer():
            return int(v)
        if isinstance(v, str):
            s = v.strip()
            if s[:1] in '[{' and s[-1:] in ']}' and not re.fullmatch(r'\{(ref|q|user|today|company|schema)(:[^{}]*)?\}', s):
                try:
                    return self.deep(json.loads(s))
                except ValueError:
                    pass
            return self.resolve(s)
        return v

    def deep(self, o):
        if isinstance(o, dict):
            return {k: self.deep(x) for k, x in o.items()}
        if isinstance(o, list):
            return [self.deep(x) for x in o]
        return self.resolve(o) if isinstance(o, str) else o

    def value(self, sheet, key, v):
        """v1.62.0 workbook: a cell as the screen wants it – dates, numbers with commas, yes / no – when the field type is known
        (screen fields); a cell of an api sheet that is clearly a date (25/10/2026) becomes yyyy-mm-dd."""
        if v is None or (isinstance(v, str) and not v.strip()):
            return v
        f = (self.sheet_fields.get(sheet) or {}).get(key)
        t = (f or {}).get('type')
        if isinstance(v, str) and '{' in v:
            return self.cell(v)
        if t == 'date':
            d = norm_date(v)
            return d if d else self.cell(v)
        if t == 'datetime':
            if isinstance(v, datetime.datetime):
                return v.isoformat()
            d = norm_date(v)
            return d if d else self.cell(v)
        if t == 'number':
            n = norm_number(v)
            return n if n is not None else self.cell(v)
        if t == 'yes/no' and isinstance(v, str):
            s = v.strip().lower()
            if s in ('yes', 'y', 'true', '1', 'active', 'نعم'):
                return True
            if s in ('no', 'n', 'false', '0', 'inactive', 'لا'):
                return False
        if f is None and isinstance(v, str) and not v.strip().startswith(('[', '{')):
            d = norm_date(v) if re.fullmatch(r'\d{1,2}[-/.]\d{1,2}[-/.]\d{4}', v.strip()) else None
            if d:
                return d
        return self.cell(v)

    # ------------------------------------------------------------------ screen fields (v1.62.0 workbook)
    def fields_of(self, endpoint):
        """{key: {label, required, type, choices …}} of an importable screen (None when not known)."""
        ep = '/' + str(endpoint or '').split('?')[0].strip('/') + '/'
        if ep in self._fieldcache:
            return self._fieldcache[ep]
        out = None
        try:
            from django_tenants.utils import schema_context
            from rest_framework.request import Request
            from rest_framework.test import APIRequestFactory
            from DataTools.views import _view_for, writable_fields, extra_fields, _overrides
            from UserManagement.models import company
            func, cls, path = _view_for(ep)
            if cls is not None:
                with schema_context(self.schema):
                    raw = APIRequestFactory().get(f'{path}?schema={self.schema}')
                    raw.tenant = company.objects.filter(schema_name=self.schema).first()
                    req = Request(raw)
                    req.user = self.admin
                    fl = writable_fields(cls, req)
                    ov = _overrides(path)
                    fl = fl + list(ov.get('extra', []))   # hidden template columns stay valid columns
                    try:
                        from AccessControl.access import view_model
                        fl += extra_fields(view_model(cls()))
                    except Exception:
                        pass
                out = {f['key']: f for f in fl}
        except Exception:   # noqa: BLE001 – unknown screen: no column checks
            out = None
        self._fieldcache[ep] = out
        return out

    # ------------------------------------------------------------------ backdating
    def model_of(self, endpoint):
        from django.urls import resolve, Resolver404
        try:
            m = resolve('/' + endpoint.split('?')[0].strip('/') + '/')
        except Resolver404:
            return None
        cls = getattr(m.func, 'cls', None) or getattr(m.func, 'view_class', None)
        try:
            from AccessControl.access import view_model
            v = cls()
            mod = view_model(v)
            if mod is not None:
                return mod
        except Exception:
            pass
        try:   # app kit register: /<app>/api/r/<kind>/
            kind = m.kwargs.get('kind')
            app = (getattr(m.func, 'view_initkwargs', {}) or {}).get('app')
            if kind and app is not None:
                return app.regs[kind].model
        except Exception:
            pass
        try:   # app registers keyed by kind in the view's module (HSE: registers.REGS)
            import sys
            regs = getattr(sys.modules.get(cls.__module__), 'REGS', None)
            kind = m.kwargs.get('kind')
            if kind and isinstance(regs, dict) and kind in regs:
                return regs[kind].model
        except Exception:
            pass
        if getattr(cls, 'zeo_model', None) is not None:   # views that declare their model for the record panel
            return cls.zeo_model
        qs = getattr(cls, 'queryset', None)
        return getattr(qs, 'model', None)

    @staticmethod
    def _skip_spec(skip):
        s = str(skip).strip()
        if s.startswith('{q:') and s.endswith('}'):   # remove one level only – nested {q:…{ref:X}} keeps its inner braces
            s = s[3:-1]
        return s

    def target_of(self, endpoint):
        """(model, pk) of the record an action URL works on: …/<pk>/… of a detail or app-kit record."""
        from django.urls import resolve, Resolver404
        path = '/' + endpoint.split('?')[0].strip('/') + '/'
        try:
            kw = resolve(path).kwargs
        except Resolver404:
            return None, None
        pk = kw.get('pk') or kw.get('id') or kw.get('rid')
        if not pk:
            return None, None
        cands = [path, re.sub(r'/do/[^/]+/$', '/', path)]
        cut = path.rstrip('/').split('/')
        while cut and cut[-1] != str(pk):   # …/<pk>/assign/ → …/<pk>/
            cut.pop()
            cands.append('/'.join(cut) + '/')
        for c in cands:
            m = self.model_of(c)
            if m is not None:
                return m, pk
        return None, pk

    def stamps(self, model, pk):
        from django.db import models as M
        from django_tenants.utils import schema_context
        if model is None or not pk:
            return None
        with schema_context(self.schema):
            obj = model.objects.filter(pk=pk).first()
            if obj is None:
                return None
            return {f.name: getattr(obj, f.name) for f in model._meta.fields
                    if isinstance(f, (M.DateField, M.DateTimeField)) and f.name not in ('date_joined',)}

    def stamp_action(self, model, pk, before, day):
        """An action row with @date: the times the action just set (approved / resolved / issued … = now or today) move to that day."""
        import zoneinfo
        from django.db import models as M
        from django.utils import timezone
        from django_tenants.utils import schema_context
        try:
            day = datetime.date.fromisoformat(str(day)[:10])
        except ValueError:
            return
        after = self.stamps(model, pk)
        if not after or before is None:
            return
        now, today = timezone.now(), timezone.localdate()
        upd = {}
        with schema_context(self.schema):
            for f in model._meta.fields:
                if f.name not in after or after[f.name] == before.get(f.name) or after[f.name] is None:
                    continue
                v = after[f.name]
                if isinstance(f, M.DateTimeField) and abs((now - v).total_seconds()) < 900:
                    upd[f.name] = datetime.datetime.combine(day, datetime.time(10 + int(pk) % 7, (int(pk) * 11) % 60)).replace(tzinfo=zoneinfo.ZoneInfo('Asia/Dubai'))
                elif isinstance(f, M.DateField) and not isinstance(f, M.DateTimeField) and v == today:
                    upd[f.name] = day
            if upd:
                model.objects.filter(pk=pk).update(**upd)

    def backdate(self, model, pk, day):
        if model is None or not pk or not day:
            return
        import zoneinfo
        from django.db import models as M
        from django_tenants.utils import schema_context
        try:
            day = datetime.date.fromisoformat(str(day)[:10])
        except ValueError:
            return
        with schema_context(self.schema):
            obj = model.objects.filter(pk=pk).first()
            if obj is None:
                return
            upd = {}
            for f in model._meta.fields:
                if f.name not in TS_FIELDS or getattr(obj, f.name, None) is None:
                    continue
                if isinstance(f, M.DateTimeField):
                    upd[f.name] = datetime.datetime.combine(day, datetime.time(9 + pk % 8, (pk * 7) % 60)).replace(tzinfo=zoneinfo.ZoneInfo('Asia/Dubai'))
                elif isinstance(f, M.DateField):
                    upd[f.name] = day
            if upd:
                model.objects.filter(pk=pk).update(**upd)

    # ------------------------------------------------------------------ results
    def note(self, sheet, row, ok, msg=''):
        s = self.stats.setdefault(sheet, {'ok': 0, 'failed': 0, 'skipped': 0})
        s[ok] += 1
        if ok == 'failed':
            entry = self.explain(sheet, row, str(msg))
            for k, v in (self.rowvals.get((sheet, row)) or {}).items():   # v1.62.0 workbook: later rows that need this one
                if not str(k).startswith('@') and v not in (None, '') and len(str(v)) < 80:
                    self.failed_vals.setdefault(str(v).strip().lower(), (sheet, row))
            self.log.append(entry)
            if self.verbose and self.cmd is not None:
                self.cmd.stdout.write(f'  ! {sheet} row {row}: {str(msg)[:300]}' + (f'  [{entry["letter"]} {entry["column"]}]' if entry.get('column') else ''))
        if self.on_row:
            self.on_row(sheet)

    def problem(self, sheet, msg, fix='', column='', letter='', level='error', row=1):
        """v1.62.0 workbook: a problem of the sheet itself (name, columns) – listed with the rows, not counted as a row."""
        self.log.append({'sheet': sheet, 'row': row, 'error': msg, 'fix': fix, 'column': column, 'letter': letter, 'value': '', 'level': level, 'kind': 'sheet'})

    FIXES = (   # v1.62.0 workbook: (part of the message, the fix shown next to it)
        ('this field is required', 'Fill in this column – it is mandatory.'),
        ('may not be null', 'Fill in this column – it is mandatory.'),
        ('may not be blank', 'Fill in this column – it is mandatory.'),
        ('date has wrong format', 'Write the date as dd/mm/yyyy (e.g. 25/10/2026) or format the cell as a date.'),
        ('datetime has wrong format', 'Write the date as dd/mm/yyyy (e.g. 25/10/2026), with the time if needed (25/10/2026 09:30).'),
        ('a valid number is required', 'Write a plain number (e.g. 12500.50); thousands commas and AED are fine, letters are not.'),
        ('a valid integer is required', 'Write a whole number (no decimals, no letters).'),
        ('is not a valid choice', 'Use one of the allowed values (see the "Columns help" sheet).'),
        ('matches several records', 'Use the code (or the id) instead of the name – this name belongs to more than one record.'),
        ('not found', 'Use a code or name that exists in ZEO, or add it in its own sheet (it loads earlier – see the Index order).'),
        ('is also on row', 'The same value is twice in this sheet – keep one row or correct the value.'),
        ('is the same as row', 'Delete one of the two rows (or correct the values if both are needed).'),
        ('already exist', 'The record is already in ZEO – delete the row, or keep it: it is skipped when it has the same values as before.'),
        ('unknown reference', 'The {ref:…} name is given by an earlier row that failed or is missing – fix that row first.'),
        ('nothing found for', 'The record this row points to is not in ZEO (an earlier row failed or the value is wrong).'),
        ('user "', 'The @as user does not exist – use the user name or e-mail of a login of this company.'),
        ('may not act as', 'Use the user name of an employee login of this company (not a system administrator).'),
        ('permission', 'The user who loads it (or the @as user) lacks the right – load as a company administrator or give the right.'),
        ('no more than', 'Shorten the value.'),
        ('valid email', 'Write a valid e-mail address (name@company.com).'),
        ('unique', 'This value must be unique – another record already has it.'),
    )

    def explain(self, sheet, row, msg):
        """v1.62.0 workbook: the column a message is about (letter + title), its value, and how to fix it."""
        titles, keys = self.headers.get(sheet, ([], []))
        flds = self.sheet_fields.get(sheet) or {}
        low = msg.lower()
        best = None
        try:
            from DataTools.duplicates import pretty
        except Exception:   # noqa: BLE001
            def pretty(x):
                return x
        for i, (h, k) in enumerate(zip(titles, keys)):
            if not h or str(h).startswith('@'):
                continue
            names = {str(h), str(k), pretty(k), str(k).replace('_', ' '), str((flds.get(k) or {}).get('label') or '')}
            for n in names:
                n = n.strip().lower()
                if n and (low.startswith(n + ':') or f'; {n}:' in low or f' {n}:' in low or f'"{n}"' in low or f'“{n}”' in low or low.startswith(n + ' ')):
                    if best is None or len(n) > best[1]:
                        best = (i, len(n))
        col = best[0] if best else None
        if col is not None and best and low[:best[1] + 1] == low[:best[1]] + ':' and str(titles[col]).lower() != low[:best[1]]:
            msg = str(titles[col]) + msg[best[1]:]   # "Emp dept id: …" → the column title the person sees ("Department: …")
        val = ''
        if col is not None:
            v = (self.rowvals.get((sheet, row)) or {}).get(keys[col])
            val = '' if v is None else str(v)[:120]
        fix = ''
        for part, f in self.FIXES:
            if part in low:
                fix = f
                break
        m = re.search(r'“([^”]+)” not found', msg)
        if m and m.group(1).strip().lower() in self.failed_vals:
            fs, fr = self.failed_vals[m.group(1).strip().lower()]
            if fs != sheet or fr != row:
                fix = f'It is on row {fr} of the sheet "{fs}", which failed – fix that row first; this row then loads too.'
        if 'is not a valid choice' in low and col is not None:
            ch = (flds.get(keys[col]) or {}).get('choices') or []
            if ch:
                fix = 'Use one of: ' + ', '.join(str(c.get('label') or c.get('value')) for c in ch[:25]) + '.'
        msg = re.sub(r'^(POST|PUT|PATCH|DELETE|GET) \S+ → ', 'Refused by the screen: ', msg)
        return {'sheet': sheet, 'row': row, 'error': msg[:900], 'fix': fix, 'column': titles[col] if col is not None else '',
                'letter': col_letter(col) if col is not None else '', 'value': val, 'level': 'error'}

    @staticmethod
    def flat(res):
        """A refusal of the screen's API in plain words ("Employee code: already exists")."""
        if isinstance(res, dict) and set(res) <= {'detail'}:
            return str(res.get('detail'))
        if isinstance(res, (dict, list)):
            try:
                from DataTools.views import _flatten
                return '; '.join(_flatten(res))[:700]
            except Exception:   # noqa: BLE001
                pass
        return json.dumps(res, default=str)[:400] if res is not None else ''

    # ------------------------------------------------------------------ memory of loaded rows (v1.62.0 workbook)
    def remembered(self, sheet, rn, row):
        if self.memory is None:
            return False
        dg = self.digests.get((sheet, rn))
        hit = self.memory.get(sheet, dg) if dg else None
        if hit is None:
            return False
        rid, ref = hit
        if ref and rid:
            self.refs[ref] = rid
        elif row.get('@ref') and rid:
            self.refs[row['@ref']] = rid
        return True

    def remember(self, sheet, rn, row, rid):
        if self.memory is None or self.trial or self.check:
            return
        dg = self.digests.get((sheet, rn))
        if dg:
            self.memory.add(sheet, dg, rid, row.get('@ref') or '')

    def act_as(self, name):
        u = self._user(name)
        if self.may_act_as:
            why = self.may_act_as(u)
            if why:
                raise ValueError(why)
        return u

    # ------------------------------------------------------------------ sheets
    def keys_of(self, sheet, headers, endpoint=''):
        m = self.colmap.get(sheet, {})
        flds = self.fields_of(endpoint) if endpoint else None
        by_label = {}
        if flds:
            for k, f in flds.items():
                by_label[str(f.get('label') or '').strip().lower()] = k
                by_label[str(k).lower()] = k
        out = []
        for h in headers:
            if (h or '').startswith('@'):
                out.append(h)
            elif h in m:
                out.append(m[h])
            elif by_label.get(clean_title(h).lower()):   # v1.62.0 workbook: screen labels without a Columns sheet, "Name *"
                out.append(by_label[clean_title(h).lower()])
            else:
                out.append(h)
        return out

    def load_rows(self, spec, rows):
        """rows: [(excel row number, {key: value})]"""
        how, sheet, ep = spec['how'], spec['sheet'], spec['endpoint']
        default_as = spec.get('as')
        rows = self._drop_dups(sheet, rows)
        if how == 'import':
            return self._import(spec, rows, default_as)
        if how == 'job':   # v1.62.0 the scheduler's daily job of an app (only these – a workbook cannot run anything else)
            from django_tenants.utils import schema_context
            jobs = {'maintenance-daily': ('ZeoMaintenance.app', 'daily')}
            for rn, row in rows:
                self.stop_check()
                name_ = str(row.get('job') or '').strip()
                if name_ not in jobs:
                    self.note(sheet, rn, 'failed', f'unknown job "{name_}" – one of {sorted(jobs)}'); continue
                if self.check:   # --quick-check saves nothing: the job is not run
                    self.note(sheet, rn, 'skipped'); continue
                try:
                    import importlib
                    mod, fn = jobs[name_]
                    with schema_context(self.schema):
                        res = getattr(importlib.import_module(mod), fn)()
                    self.note(sheet, rn, 'ok' if (res or {}).get('preventive_jobs', 1) else 'skipped')
                except Exception as e:   # noqa: BLE001
                    self.note(sheet, rn, 'failed', f'{type(e).__name__}: {e}')
            return None
        for rn, row in rows:
            self.stop_check()
            try:
                if self.remembered(sheet, rn, row):
                    self.note(sheet, rn, 'skipped')
                    continue
                skip = row.get('@skip_if')
                if skip:
                    found = self.lookup(self._skip_spec(skip))
                    if found:
                        if row.get('@ref'):
                            self.refs[row['@ref']] = found
                        self.note(sheet, rn, 'skipped')
                        continue
                until = str(row.get('@until') or '').strip()[:10]
                if until and datetime.date.today().isoformat() > until:   # v1.62.0 a row for a future day (shift request …) after that day has passed
                    s_ = self.stats.setdefault(sheet, {'ok': 0, 'failed': 0, 'skipped': 0})
                    s_['past'] = s_.get('past', 0) + 1
                    if self.on_row:
                        self.on_row(sheet)
                    continue
                user = self.act_as(row['@as']) if row.get('@as') else (self.act_as(default_as) if default_as else self.admin)
                if how == 'action':
                    method = (row.get('@method') or 'POST').upper()
                    path = self.cell(row.get('@endpoint') or ep)
                    body = row.get('@body')
                    body = self.cell(body) if isinstance(body, str) and body.strip() else (body if body not in (None, '') else {})
                else:
                    method = (spec.get('method') or 'POST').upper()
                    path = self.cell(row.get('@endpoint') or ep)
                    body = {k: self.value(sheet, k, v) for k, v in row.items() if not k.startswith('@') and v not in (None, '')}
                    if spec.get('wrap'):   # v1.62.0 workbook: Index "Wrap" – the row goes inside a list key of the body
                        body = self._wrap(spec['wrap'], body)
                if self.check:
                    self.note(sheet, rn, 'ok')
                    continue
                tgt = self.target_of(str(path)) if (how == 'action' and row.get('@date')) else (None, None)
                before = self.stamps(*tgt) if tgt[0] is not None else None
                st, res = self.call(user, method, str(path), body if method not in ('GET', 'DELETE') else None)
                if before is not None and st < 400:
                    self.stamp_action(tgt[0], tgt[1], before, self.cell(row['@date']))
                expect = str(row.get('@expect') or '').strip()
                good = (str(st) in [x.strip() for x in expect.split(',')]) if expect else st < 400
                if not good:
                    self.note(sheet, rn, 'failed', f'{method} {path} → {st} {self.flat(res)}')
                    continue
                if spec.get('wrap') and isinstance(res, dict) and isinstance(res.get('errors'), list):   # "Row 1: …" of a one-row body
                    one = next((v[0] for v in body.values() if isinstance(v, list) and v and isinstance(v[0], dict)), {})
                    said = {'employee not found': ('Employee code', one.get('employee_code') or one.get('employee')),
                            'leave type not found': ('Leave type', one.get('leave_type'))}
                    res['errors'] = [(lambda t: f'{said[t][0]}: “{said[t][1]}” not found' if t in said else t)(re.sub(r'^Row \d+: ', '', str(x)))
                                     for x in res['errors']]
                # v1.62.0 workbook: a 200 answer that lists errors and loaded nothing (opening balances …) is a failure
                if isinstance(res, dict) and isinstance(res.get('errors'), list) and res['errors'] and not expect \
                        and (res.get('error_count') or len(res['errors'])) and not (res.get('loaded') or res.get('ok') or res.get('done')):
                    self.note(sheet, rn, 'failed', '; '.join(map(str, res['errors'][:10])))
                    continue
                new_id = res.get('id') if isinstance(res, dict) else None
                if row.get('@ref') and new_id:
                    self.refs[row['@ref']] = new_id
                if row.get('@date') and new_id and how != 'action':
                    self.backdate(self.model_of(str(path)), new_id, self.cell(row['@date']))
                self.remember(sheet, rn, row, new_id)
                self.note(sheet, rn, 'ok')
            except Stopped:
                raise
            except Exception as ex:   # noqa: BLE001 – one bad row never stops the load
                self.note(sheet, rn, 'failed', f'{type(ex).__name__}: {ex}' if not isinstance(ex, ValueError) else str(ex))

    @staticmethod
    def _wrap(wrap, body):
        """wrap 'rows' → {'rows': [body]}; 'rows|date,mode' → {'date': .., 'mode': .., 'rows': [rest of the row]}."""
        key, _, top = str(wrap).partition('|')
        tops = [t.strip() for t in top.split(',') if t.strip()]
        out = {t: body.pop(t) for t in tops if t in body}
        out[key.strip()] = [body]
        return out

    def _drop_dups(self, sheet, rows):
        """v1.62.0 workbook: the second of two identical rows fails (both would make the same record)."""
        out = []
        for rn, row in rows:
            first = self.dups.get((sheet, rn))
            if first:
                self.note(sheet, rn, 'failed', f'Row {rn} is the same as row {first}.')
            else:
                out.append((rn, row))
        return out

    def _import(self, spec, rows, default_as):
        sheet, ep = spec['sheet'], spec['endpoint']
        todo = []
        for rn, row in rows:
            try:
                if self.remembered(sheet, rn, row):
                    self.note(sheet, rn, 'skipped')
                    continue
                skip = row.get('@skip_if')
                if skip:
                    found = self.lookup(self._skip_spec(skip))
                    if found:
                        if row.get('@ref'):
                            self.refs[row['@ref']] = found
                        self.note(sheet, rn, 'skipped')
                        continue
                data = {k: self.value(sheet, k, v) for k, v in row.items() if not k.startswith('@') and v not in (None, '')}
                todo.append((rn, row, data))
            except Exception as ex:   # noqa: BLE001
                self.note(sheet, rn, 'failed', f'{type(ex).__name__}: {ex}' if not isinstance(ex, ValueError) else str(ex))
        # rows done by another user go one by one; the rest in blocks
        groups = {}
        for t in todo:
            who = t[1].get('@as') or default_as or ''
            groups.setdefault(who, []).append(t)
        model = self.model_of(ep)
        for who, items in groups.items():
            try:
                user = self.act_as(who) if who else self.admin
            except Exception as ex:   # noqa: BLE001
                for rn, _, _ in items:
                    self.note(sheet, rn, 'failed', str(ex))
                continue
            for i in range(0, len(items), self.block):
                self.stop_check()
                block = items[i:i + self.block]
                # --check (trial) imports for real inside the trial transaction; --quick-check only validates
                st, res = self.call(user, 'POST', '/tools/api/import/', {'endpoint': ep, 'rows': [d for _, _, d in block], 'dry_run': bool(self.check)})
                if st >= 400 or not isinstance(res, dict):
                    for rn, _, _ in block:
                        self.note(sheet, rn, 'failed', f'import {st}: {self.flat(res)}')
                    continue
                for (rn, row, _), r in zip(block, res.get('results') or []):
                    if r.get('ok'):
                        if row.get('@ref') and r.get('id'):
                            self.refs[row['@ref']] = r['id']
                        if row.get('@date') and r.get('id') and not self.check:
                            self.backdate(model, r['id'], self.cell(row['@date']))
                        self.remember(sheet, rn, row, r.get('id'))
                        self.note(sheet, rn, 'ok')
                    else:
                        msgs = []
                        for m in (r.get('errors') or ['refused']):   # "is also on row 3" (row of the block) → the Excel row
                            msgs.append(re.sub(r'(row )(\d+)', lambda mm: mm.group(1) + str(block[int(mm.group(2)) - 1][0])
                                               if 0 < int(mm.group(2)) <= len(block) else mm.group(0), str(m)))
                        self.note(sheet, rn, 'failed', '; '.join(msgs))


# ---------------------------------------------------------------------- v1.62.0 workbook: reading and running a workbook
def read_book(path, L, only=None, out=None):
    """Read the workbook into L (column map, keys, values, digests) – returns the specs and the rows of each sheet.
    Problems of the workbook itself (names, columns, the same row twice) are added to L.log / L.problem."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as ex:   # noqa: BLE001
        raise CommandError(f'The file could not be opened as an Excel workbook (.xlsx): {ex}')
    names = list(wb.sheetnames)
    index = next((n for n in names if n.strip().lower() == 'index'), None)
    if index is None:
        raise CommandError('The workbook has no "Index" sheet. Start from ZEO_Client_Data_Template_1.62.0.xlsx (Settings › Data workbook › '
                           'Download the template) – its Index lists the sheets to load, in order.')
    colsheet = next((n for n in names if n.strip().lower() == 'columns'), None)
    if colsheet:
        it = wb[colsheet].iter_rows(values_only=True)
        hdr = [str(h or '').strip().lower() for h in next(it, [])]
        for r in it:
            d = dict(zip(hdr, r))
            if d.get('sheet') and d.get('column') and d.get('key'):
                L.colmap.setdefault(str(d['sheet']).strip(), {})[str(d['column']).strip()] = str(d['key']).strip()
    it = wb[index].iter_rows(values_only=True)
    hdr = [str(h or '').strip().lower() for h in next(it, [])]
    specs, listed = [], set()
    by_low = {n.strip().lower(): n for n in names}
    for r in it:
        d = dict(zip(hdr, r))
        if not d.get('sheet'):
            continue
        sheet = str(d['sheet']).strip()
        listed.add(sheet.lower())
        if only and sheet not in only:
            continue
        real = sheet if sheet in names else by_low.get(sheet.lower())
        if real is None:
            near = closest(sheet, [n for n in names if n.strip().lower() not in SYSTEM_SHEETS])
            L.problem(sheet, f'The Index names the sheet "{sheet}" but the workbook has no sheet with that name – it is not loaded.',
                      fix=(f'Rename the sheet "{near}" to "{sheet}" (or correct the Index).' if near else f'Add a sheet called "{sheet}" or remove it from the Index.'),
                      row=0)
            if out:
                out(f'Sheet "{sheet}" is in the Index but not in the workbook – skipped.' + (f' (closest: "{near}")' if near else ''))
            continue
        specs.append({'sheet': real, 'endpoint': str(d.get('endpoint') or '').strip(), 'how': str(d.get('how') or 'import').strip().lower(),
                      'phase': str(d.get('phase') or 'setup').strip().lower(), 'as': str(d.get('run as') or '').strip() or None,
                      'method': str(d.get('method') or 'POST').strip() or 'POST', 'wrap': str(d.get('wrap') or '').strip(),
                      'mandatory': str(d.get('mandatory') or '').strip().lower().startswith(('y', 'go'))})
    for n in names:   # sheets the Index does not list
        if n.strip().lower() in SYSTEM_SHEETS or n.strip().lower() in listed or (only and n not in only):
            continue
        near = closest(n, [s for s in listed])
        L.problem(n, f'The sheet "{n}" is not in the Index – it is not loaded.',
                  fix=(f'Rename it to the sheet the Index lists ("{next((s["sheet"] for s in specs if s["sheet"].lower() == near), near)}").' if near
                       else 'Add it to the Index (with its screen, endpoint and how), or delete it.'), level='warning', row=0)
    data = {}
    for s in specs:
        sheet = s['sheet']
        rows = wb[sheet].iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else '' for h in next(rows, [])]
        while headers and not headers[-1]:
            headers.pop()
        flds = L.fields_of(s['endpoint']) if s['how'] in ('import', 'api') and s['endpoint'] else None
        if flds:
            L.sheet_fields[sheet] = flds
        keys = L.keys_of(sheet, headers, s['endpoint'] if s['how'] in ('import', 'api') else '')
        L.headers[sheet] = (headers, keys)
        if s['how'] == 'import' and flds:
            # unknown columns (the import would drop them) and mandatory columns that are missing
            labels = [str(f.get('label') or k) for k, f in flds.items()]
            for i, (h, k) in enumerate(zip(headers, keys)):
                if not h or h.startswith('@') or k in flds:
                    continue
                near = closest(clean_title(h), labels + list(flds))
                if near in flds:
                    near = str(flds[near].get('label') or near)
                L.problem(sheet, f'Unknown column "{h}" – it is not a field of this screen and would be left out.',
                          fix=(f'Did you mean "{near}"? Rename the column title.' if near else 'Rename it to a column of the "Columns help" sheet, or delete it.'),
                          column=h, letter=col_letter(i))
            have = set(keys)
            for k, f in flds.items():
                if f.get('required') and k not in have:
                    L.problem(sheet, f'The mandatory column "{f.get("label") or k}" is missing.',
                              fix=f'Add a column "{f.get("label") or k}" and fill it in on every row.', column=str(f.get('label') or k))
        if s['how'] not in ('import', 'api', 'action', 'job'):
            L.problem(sheet, f'The Index says How = "{s["how"]}" – it must be import, api, action or job.', fix='Correct the How column of the Index.', row=0)
        out_rows, seen, occ = [], {}, {}
        for n, r in enumerate(rows, start=2):
            if r is None or all(v in (None, '') for v in r):
                continue
            row = {k: v for k, v in zip(keys, r) if k}
            if row.get('@example') not in (None, ''):   # v1.62.0 workbook: example rows of the template
                st_ = L.stats.setdefault(sheet, {'ok': 0, 'failed': 0, 'skipped': 0})
                st_['examples'] = st_.get('examples', 0) + 1
                continue
            row.pop('@example', None)
            if all(v in (None, '') for k, v in row.items() if not k.startswith('@')) and not any(row.get(k) for k in ('@endpoint', '@body')):
                continue   # only special columns filled (an empty template row)
            sig = json.dumps([(k, str(v)) for k, v in sorted(row.items()) if k not in ('@note',) and v not in (None, '')], default=str)
            occ[sig] = occ.get(sig, 0) + 1
            L.digests[(sheet, n)] = hashlib.sha1(f'{sheet}\n{sig}\n{occ[sig]}'.encode()).hexdigest()[:32]
            L.rowvals[(sheet, n)] = row
            if s['how'] in ('import', 'api') and not row.get('@ref'):
                if sig in seen:
                    L.dups[(sheet, n)] = seen[sig]
                else:
                    seen[sig] = n
            out_rows.append((n, row))
        data[sheet] = out_rows
    return specs, data


def months_of(specs, data, from_month=''):
    months = sorted({str(r.get('@month'))[:7] for s in specs if s['phase'] == 'monthly' for _, r in data[s['sheet']] if r.get('@month')})
    if from_month:
        months = [m for m in months if m >= from_month]
    return months


def run_book(L, specs, data, from_month='', out=None, on_sheet=None):
    """Load the setup sheets, then the months. out(text) gets one line per sheet; on_sheet(sheet, month) before each sheet."""
    def _run(spec, rows, month=''):
        t0 = time.time()
        if on_sheet:
            on_sheet(spec['sheet'], month, len(rows))
        before = dict(L.stats.get(spec['sheet'], {'ok': 0, 'failed': 0, 'skipped': 0}))
        try:
            L.load_rows(spec, rows)
        except Stopped:
            raise
        except Exception:   # noqa: BLE001
            L.note(spec['sheet'], 0, 'failed', traceback.format_exc()[-800:])
        now = L.stats.get(spec['sheet'], {'ok': 0, 'failed': 0, 'skipped': 0})
        d = {k: now[k] - before.get(k, 0) for k in now if isinstance(now[k], int)}
        if out:
            out(f'  {spec["sheet"]}: {d.get("ok", 0)} ok, {d.get("skipped", 0)} already there, {d.get("failed", 0)} failed'
                f'{", " + str(d["past"]) + " date passed" if d.get("past") else ""} ({time.time() - t0:.0f} s)')
    for s in [x for x in specs if x['phase'] != 'monthly']:   # also on a resume: their rows are skipped (@skip_if / loaded before)
        _run(s, data[s['sheet']])
    for m in months_of(specs, data, from_month):
        if out:
            out(f'— {m}')
        for s in [x for x in specs if x['phase'] == 'monthly']:
            rows = [x for x in data[s['sheet']] if str(x[1].get('@month'))[:7] == m]
            if rows:
                _run(s, rows, m)
    # monthly rows without @month are never loaded – say so
    for s in [x for x in specs if x['phase'] == 'monthly']:
        for rn, r in data[s['sheet']]:
            if not r.get('@month') and not from_month:
                L.note(s['sheet'], rn, 'failed', 'The @month column is empty – a monthly sheet needs the month of each row (YYYY-MM, e.g. 2026-05).')


def totals(L):
    tot = {k: sum(v.get(k, 0) for v in L.stats.values()) for k in ('ok', 'failed', 'skipped', 'past', 'examples')}
    tot['problems'] = sum(1 for e in L.log if e.get('kind') == 'sheet')
    return tot


class Command(BaseCommand):
    help = 'Load a data workbook (Index sheet + one sheet per screen) into a company through the screens\' own API.'

    def add_arguments(self, p):
        p.add_argument('workbook')
        p.add_argument('--schema', required=True, help='the company schema (setup_zeo --schema …)')
        p.add_argument('--user', default='', help='the administrator who loads it (default: <schema>_admin)')
        p.add_argument('--only', default='', help='comma separated sheet names')
        p.add_argument('--from-month', default='', help='YYYY-MM: start the monthly part here (to resume)')
        p.add_argument('--check', action='store_true', help='full trial: every row is loaded in a transaction that is rolled back (nothing saved)')
        p.add_argument('--quick-check', action='store_true', help='validate import rows only, without trying the other rows')

    def handle(self, *a, **o):
        import logging
        logging.disable(logging.WARNING)
        verbosity = int(o.get('verbosity', 1))
        trial = bool(o['check']) and not o['quick_check']
        L = Loader(self, o['schema'], o['user'], bool(o['quick_check']), verbosity, trial=trial)
        only = {s.strip() for s in o['only'].split(',') if s.strip()}
        specs, data = read_book(o['workbook'], L, only, out=lambda t: self.stdout.write(self.style.WARNING(t)))
        for e in [x for x in L.log if x.get('kind') == 'sheet']:
            self.stdout.write(self.style.WARNING(f'  ! {e["sheet"]}: {e["error"]} {e.get("fix", "")}'))
        t0 = time.time()
        what = ' (check: a trial that is rolled back – nothing is saved)' if trial else (' (quick check)' if o['quick_check'] else '')
        self.stdout.write(f'Loading {o["workbook"]} into {o["schema"]} as {L.admin.username}{what}')
        out = lambda t: self.stdout.write(self.style.HTTP_INFO(t) if t.startswith('—') else t)   # noqa: E731
        if trial:
            from django.db import transaction
            from zeo import mailguard
            with mailguard.suppressed(), transaction.atomic():
                run_book(L, specs, data, o['from_month'], out=out)
                transaction.set_rollback(True)
        else:
            run_book(L, specs, data, o['from_month'], out=out)
        tot = totals(L)
        log = o['workbook'] + '.log.json'
        try:
            json.dump({'stats': L.stats, 'failed': L.log, 'refs': len(L.refs)}, open(log, 'w'), indent=1, default=str)
        except OSError:
            log = ''
        self.stdout.write(self.style.SUCCESS(f'Done in {time.time() - t0:.0f} s: {tot["ok"]} rows {"would load" if trial else "loaded"}, {tot["skipped"]} already there, '
                                             f'{tot["failed"]} failed{" – see " + log if log and (tot["failed"] or L.log) else ""}.'
                                             + (f' {tot["past"]} rows left out – their day has passed (@until).' if tot['past'] else '')
                                             + (f' {tot["examples"]} example rows ignored.' if tot['examples'] else '')))
